#!/usr/bin/env python3
"""
Mouse colony tool - breeder ages, relatedness check and pair suggestions
for the Völgyi lab colony workbook (Volgyi_mice_colony.xlsx).

Install: double-click install_windows.bat (Windows) or install_mac.command (macOS).
Manual:  python mouse_colony_tool.py [path/to/Volgyi_mice_colony.xlsx]
Needs:   Python 3.8+ with Tkinter, and openpyxl   (pip install openpyxl)

Close the workbook in Excel before saving changes from the tool.
Every save first copies the workbook into a "backups" folder next to it.
"""
import datetime as dt
import os
import re
import shutil
import sys

import openpyxl

APP_NAME = 'Mouse Colony'
VERSION = '2.3'
WORKBOOK_NAME = 'Volgyi_mice_colony.xlsx'


def _config_path():
    if sys.platform.startswith('win'):
        base = os.environ.get('APPDATA') or os.path.expanduser('~')
        return os.path.join(base, 'MouseColony', 'settings.json')
    return os.path.join(os.path.expanduser('~'), '.mouse_colony.json')


def load_config():
    import json
    try:
        with open(_config_path(), encoding='utf-8') as fh:
            return json.load(fh)
    except Exception:
        return {}


def save_config(cfg):
    import json
    try:
        os.makedirs(os.path.dirname(_config_path()), exist_ok=True)
        with open(_config_path(), 'w', encoding='utf-8') as fh:
            json.dump(cfg, fh, indent=1)
    except Exception:
        pass


def default_workbook():
    """Last workbook opened, else the one next to the program, else Documents/Mouse colony/."""
    last = load_config().get('last_workbook')
    if last and os.path.exists(last):
        return last
    here = os.path.dirname(os.path.abspath(sys.argv[0] if getattr(sys, 'frozen', False) else __file__))
    docs = os.path.join(os.path.expanduser('~'), 'Documents', 'Mouse colony')
    for p in (os.path.join(here, WORKBOOK_NAME), os.path.join(docs, WORKBOOK_NAME)):
        if os.path.exists(p):
            return p
    return None


ERROR_LOG = os.path.join(os.path.expanduser('~'), 'mouse_colony_error.log')


def log_error(text):
    try:
        with open(ERROR_LOG, 'a', encoding='utf-8') as fh:
            fh.write(f'--- {dt.datetime.now():%Y-%m-%d %H:%M:%S}\n{text}\n')
    except Exception:
        pass

# --------------------------------------------------------------------------
# Workbook layout (must match the workbook)
# --------------------------------------------------------------------------
N = 1000
BREEDER_COLS = ['ID', 'Line', 'Sex', 'Birth', 'Cage', 'Status', 'Sire ID', 'Dam ID', 'Notes']
PAIR_COLS = ['Pair cage', 'Line', 'Male ID', 'Female 1 ID', 'Female 2 ID', 'Set up', 'Status', 'Notes']
CAGE_COLS = ['Cage', 'Line', 'Line as written', 'Sex', 'Count', 'Birth']  # then 3 formula cols
CAGE_COLS2 = ['Sire ID', 'Dam ID', 'Note', 'Reserved for']                 # cols J..M
DATE_FMT = 'yyyy-mm-dd'


PROC_HEAD = ['Date', 'Cage', 'Animal ID', 'Line', 'Sex', 'Mice', 'Born', 'Procedure', 'Details', 'Person', 'Notes']
PROC_W = [12, 9, 10, 24, 6, 6, 12, 22, 34, 14, 34]
TERM_HEAD = ['Date', 'Cage', 'Animal ID', 'Line', 'Sex', 'Mice', 'Born', 'Age (days)', 'Outcome', 'Reason / use',
             'Project', 'Person', 'Sire ID', 'Dam ID', 'Genotype', 'Procedures done', 'Notes']
TERM_W = [12, 9, 10, 24, 6, 6, 12, 9, 26, 28, 18, 14, 9, 9, 26, 60, 30]
PROCEDURE_TYPES = ['AAV injection', 'Injection (other)', 'Surgery', 'Cranial window', 'Optic fiber implant',
                   'Genotyping (tail/ear)', 'Ear tag / ID', 'Tamoxifen', 'Behaviour', 'Imaging',
                   'Electrophysiology', 'Health check', 'Other']
OUTCOMES = ['Terminated - experiment', 'Terminated - tissue / perfusion', 'Culled - over age',
            'Culled - wrong genotype', 'Culled - surplus', 'Culled - health', 'Found dead', 'Transferred out']
GONE = ('Retired', 'Terminated')
EVENTS_SHEET = 'Imported events'
EVENTS_HEAD = ['Key', 'Date', 'Cage', 'Event', 'Details', 'Imported on', 'By']
EVENT_WORDS = {'kivitel': 'removed', 'kivét': 'removed', 'kivetel': 'removed', 'áthelyez': 'moved',
               'athelyez': 'moved', 'áthelyezés': 'moved', 'leválaszt': 'weaned', 'levalaszt': 'weaned',
               'leválasztás': 'weaned', 'elpusztult': 'died', 'elhullott': 'died'}
GT_CODES = ['?', 'hom', 'het', 'wt', 'mixed']
EAR_MARKS = ['none', 'L', 'R', 'LR', 'L2', 'R2', 'L2R', 'LR2', 'L2R2']
MICE_HEAD = ['Mouse ID', 'Cage', 'Ear mark', 'Line', 'Sex', 'Born', 'Sire ID', 'Dam ID', 'Genotype', 'Status',
             'Status date', 'Breeder ID', 'Notes', 'Needs genotyping']
MICE_W = [10, 9, 9, 24, 6, 12, 9, 9, 34, 11, 12, 10, 40, 11]
GT_HELP = ('+/+ = homozygous for the allele (homo +), +/- = heterozygous, -/- = no copy (homo -, wild type), '
           '? = not genotyped')
GT_LABEL = {'?': '? (not tested)', 'hom': '+/+ (homo +)', 'het': '+/- (het)', 'wt': '-/- (homo -)', 'mixed': 'mixed'}
GT_FROM_LABEL = {v: k for k, v in GT_LABEL.items()}
GT_SHORT = {'hom': '+/+', 'het': '+/-', 'wt': '-/-', 'mixed': 'mixed', '?': '?'}


def gt_code(label):
    return GT_FROM_LABEL.get(label, label if label in GT_LABEL else '?')
DEFAULT_LOCI = {'PV-Cre': 'PV-Cre', 'GCaMP3-Thy1': 'Thy1-GCaMP3', 'FL-Cx36KO': 'Cx36-flox',
                'Cx36KO + PV-Cre': 'Cx36-flox; PV-Cre', 'CX3CR1KO (-/-)': 'Cx3cr1-KO',
                'CX3CR1KO + Bl6': 'Cx3cr1-KO', 'GCaMP3-Thy1 + CarCre': 'Thy1-GCaMP3; CarCre',
                'Cas9 + PV-Cre': 'Cas9; PV-Cre', 'C57Bl6J': ''}


def parse_gt(text):
    """'Cx36-flox: hom; PV-Cre: het' -> {'Cx36-flox': 'hom', 'PV-Cre': 'het'}"""
    out = {}
    for part in (text or '').split(';'):
        if ':' in part:
            k, v = part.split(':', 1)
            if k.strip():
                out[k.strip()] = v.strip().lower() or '?'
    return out


def format_gt(d):
    return '; '.join(f'{k}: {v}' for k, v in d.items() if v and v != '?')


def _mut_prob(code):
    return {'hom': 1.0, 'het': 0.5, 'wt': 0.0}.get(code)


def mendel(code_m, code_f):
    """Expected pup genotypes for one autosomal locus, or None if a parent is unknown/mixed."""
    pm, pf = _mut_prob(code_m), _mut_prob(code_f)
    if pm is None or pf is None:
        return None
    res = {'hom': pm * pf, 'het': pm * (1 - pf) + pf * (1 - pm), 'wt': (1 - pm) * (1 - pf)}
    return {k: v for k, v in res.items() if v > 0}


def fmt_mendel(res):
    return ', '.join(f'{v * 100:g}% {GT_SHORT.get(k, k)}' for k, v in sorted(res.items(), key=lambda x: -x[1]))


def gt_pretty(text):
    """'Cx36-flox: hom; PV-Cre: het*' -> 'Cx36-flox +/+; PV-Cre +/-*' for display."""
    return re.sub(r':\s*(hom|het|wt|mixed)(\*?)', lambda m: ' ' + GT_SHORT[m.group(1)] + m.group(2), text or '')


def breeder_formulas(i):
    return [f'=IF(D{i}="","",TODAY()-D{i})',
            f'=IF(J{i}="","",ROUND(J{i}/30.44,1))',
            f'=IF(A{i}="","",IF(OR(F{i}="Retired",F{i}="Terminated"),"-",IF(D{i}="","NO DATE",IF(J{i}>=Settings!$B$2,"OVER AGE",'
            f'IF(J{i}>=Settings!$B$2-Settings!$B$3,"Replace soon",IF(J{i}<Settings!$B$4,"Too young","OK"))))))',
            f'=IF(D{i}="","",D{i}+Settings!$B$2)']


def _age_lookup(col, i):
    return (f'=IF({col}{i}="","",IFERROR(TODAY()-INDEX(Breeders!$D$2:$D${N},'
            f'MATCH({col}{i},Breeders!$A$2:$A${N},0)),"?"))')


def pair_formulas(i):
    return [_age_lookup('C', i), _age_lookup('D', i), _age_lookup('E', i),
            f'=IF(COUNT(I{i}:K{i})=0,"",MAX(I{i}:K{i}))',
            f'=IF(OR(A{i}="",G{i}<>"Active"),"",IF(L{i}="","NO DATE",IF(L{i}>=Settings!$B$2,"OVER AGE",'
            f'IF(L{i}>=Settings!$B$2-Settings!$B$3,"Replace soon","OK"))))',
            f'=IF(OR(L{i}="",G{i}<>"Active"),"",TODAY()+Settings!$B$2-L{i})']


def line_formulas(i):
    return [f'=COUNTIFS(Pairs!$B$2:$B${N},A{i},Pairs!$G$2:$G${N},"Active")',
            f'=COUNTIFS(Pairs!$B$2:$B${N},A{i},Pairs!$M$2:$M${N},"OVER AGE")',
            f'=COUNTIFS(Pairs!$B$2:$B${N},A{i},Pairs!$M$2:$M${N},"Replace soon")',
            f'=COUNTIFS(Cages!$B$2:$B${N},A{i},Cages!$D$2:$D${N},"M",Cages!$I$2:$I${N},"Yes")',
            f'=COUNTIFS(Cages!$B$2:$B${N},A{i},Cages!$D$2:$D${N},"F",Cages!$I$2:$I${N},"Yes")',
            f'=IF(AND(C{i}="Y",G{i}=0),"NO ACTIVE PAIR",IF(H{i}>0,"Pair over age",IF(I{i}>0,"Plan replacement","")))']


LINE_SEP = ' + '
DEFAULT_PARTS = {'Cx36KO + PV-Cre': 'FL-Cx36KO; PV-Cre', 'GCaMP3-Thy1 + CarCre': 'GCaMP3-Thy1; CarCre',
                 'Cas9 + PV-Cre': 'Cas9; PV-Cre'}


def cage_formulas(i):
    return [f'=IF(F{i}="","",TODAY()-F{i})',
            f'=IF(G{i}="","",ROUND(G{i}/30.44,1))',
            f'=IF(OR(F{i}="",D{i}="TP",D{i}=""),"",IF(G{i}<Settings!$B$4,"Too young",'
            f'IF(G{i}<=Settings!$B$5,"Yes","Too old")))']


def _s(v):
    if v is None:
        return ''
    return str(v).strip()


def _d(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def fmt_date(d):
    return d.isoformat() if d else ''


def fmt_age(days):
    if days is None:
        return ''
    return f'{days / 30.44:.1f} mo'


# --------------------------------------------------------------------------
# Shared Google Sheet: download as .xlsx, push only the changed cells
# --------------------------------------------------------------------------
class SyncError(Exception):
    pass


class ConflictError(SyncError):
    pass


class ReadOnlyError(Exception):
    pass


_EPOCH = dt.datetime(1899, 12, 30)


def _serial(v):
    if isinstance(v, dt.datetime):
        return round((v - _EPOCH).total_seconds() / 86400.0, 6)
    if isinstance(v, dt.date):
        return float((v - _EPOCH.date()).days)
    return v


def _norm(v):
    """Comparable form of a cell value (dates as serial numbers, '' for empty)."""
    if v is None or v == '':
        return ''
    if isinstance(v, bool):
        return v
    if isinstance(v, (dt.datetime, dt.date)):
        return _serial(v)
    if isinstance(v, (int, float)):
        return round(float(v), 6)
    return str(v)


def _is_formula(v):
    return isinstance(v, str) and v.startswith('=')


def _user_entered(v):
    """Value to send with valueInputOption=USER_ENTERED."""
    if v is None:
        return ''
    if isinstance(v, bool) or isinstance(v, (int, float)):
        return v
    if isinstance(v, dt.datetime):
        return v.strftime('%Y-%m-%d') if v.time() == dt.time() else v.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(v, dt.date):
        return v.isoformat()
    v = str(v)
    if v == '' or v.startswith('='):
        return v
    return "'" + v                     # keep text as text (e.g. '+/+', '1/B/2/e', '0012')


def _col_letter(c):
    return openpyxl.utils.get_column_letter(c)


def _sheet_matrix(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _at(mat, r, c):
    if r - 1 < len(mat) and c - 1 < len(mat[r - 1]):
        return mat[r - 1][c - 1]
    return None


def diff_workbooks(base_wb, new_wb):
    """{sheet: [(row, col, new_value), ...]} for cells that differ, and the list of sheets that are new."""
    changes, new_sheets = {}, []
    for ws in new_wb.worksheets:
        new = _sheet_matrix(ws)
        if ws.title not in base_wb.sheetnames:
            new_sheets.append(ws.title)
            base = []
        else:
            base = _sheet_matrix(base_wb[ws.title])
        rows = max(len(new), len(base))
        cols = max([len(r) for r in new + base] or [0])
        out = []
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                a, b = _at(new, r, c), _at(base, r, c)
                if _norm(a) != _norm(b):
                    out.append((r, c, a))
        if out:
            changes[ws.title] = out
    return changes, new_sheets


def group_ranges(title, cells):
    """Consecutive cells of a row become one A1 range."""
    out, run = [], []
    for r, c, v in sorted(cells):
        if run and r == run[-1][0] and c == run[-1][1] + 1:
            run.append((r, c, v))
        else:
            if run:
                out.append(run)
            run = [(r, c, v)]
    if run:
        out.append(run)
    q = title.replace("'", "''")
    return [(f"'{q}'!{_col_letter(rn[0][1])}{rn[0][0]}:{_col_letter(rn[-1][1])}{rn[0][0]}",
             [[_user_entered(v) for _, _, v in rn]]) for rn in out]


GOOGLE_SCOPES = ['https://www.googleapis.com/auth/spreadsheets',
                 'https://www.googleapis.com/auth/drive.readonly',
                 'openid', 'https://www.googleapis.com/auth/userinfo.email']


def client_config_ok(client):
    """A Google 'Desktop app' OAuth client file: {"installed": {"client_id": ..., "client_secret": ...}}."""
    return isinstance(client, dict) and isinstance(client.get('installed'), dict) and \
        bool(client['installed'].get('client_id'))


def google_sign_in(client, timeout=300):
    """Open the browser for Google sign-in. Returns the token dict (incl. 'account': the e-mail address)."""
    import json
    from google_auth_oauthlib.flow import InstalledAppFlow
    flow = InstalledAppFlow.from_client_config(client, GOOGLE_SCOPES)
    try:
        creds = flow.run_local_server(port=0, timeout_seconds=timeout, open_browser=True,
                                      success_message='Signed in to Mouse Colony. You can close this tab.')
    except Exception as e:
        raise SyncError(f'Google sign-in did not finish ({e.__class__.__name__}: {str(e)[:200]}).')
    if creds is None:
        raise SyncError('Google sign-in was not completed within 5 minutes.')
    tok = json.loads(creds.to_json())
    tok['account'] = _google_email(creds)
    return tok


def _google_email(creds):
    try:
        from google.auth.transport.requests import AuthorizedSession
        r = AuthorizedSession(creds).get('https://www.googleapis.com/oauth2/v3/userinfo', timeout=20)
        return r.json().get('email', '')
    except Exception:
        return ''


class GoogleBackend:
    """The shared colony workbook as a Google Sheet. Access either with the user's own Google sign-in
    (token = authorized-user dict from google_sign_in) or with a service-account key (key_info)."""

    def __init__(self, key_info=None, sheet_id='', title='', token=None):
        self.key_info = key_info
        self.token = token
        self.sheet_id = sheet_id
        self.title = title or 'shared colony'
        self.account = (key_info or {}).get('client_email', '') or (token or {}).get('account', '')
        self._sh = None

    @property
    def sh(self):
        """Connect on first use (so the app can start offline from the local copy)."""
        if self._sh is None:
            import gspread
            self._gspread = gspread
            if self.key_info:
                gc = gspread.service_account_from_dict(self.key_info)
            elif self.token:
                from google.oauth2.credentials import Credentials
                info = {k: v for k, v in self.token.items() if k != 'account'}
                gc = gspread.authorize(Credentials.from_authorized_user_info(info, GOOGLE_SCOPES))
            else:
                raise SyncError('Not signed in to Google: Shared database... -> Sign in with Google.')
            self._sh = gc.open_by_key(self.sheet_id)
            self.title = self._sh.title or self.title
        return self._sh

    @staticmethod
    def sheet_id_from(text):
        """Accepts a Google Sheets link or the bare ID."""
        m = re.search(r'/spreadsheets/d/([A-Za-z0-9_-]{20,})', text or '')
        return m.group(1) if m else (text or '').strip()

    def test(self):
        """Connect and return the sheet title (raises with a readable message)."""
        try:
            return self.sh.title
        except Exception as e:
            name = e.__class__.__name__
            msg = str(e)
            if isinstance(e, SyncError):
                raise
            if name == 'RefreshError' or 'invalid_grant' in msg:
                raise SyncError('The Google sign-in has expired or was revoked. Shared database... -> Sign in '
                                'with Google again. (If this happens every week, the login ID is still in '
                                '"Testing": see GOOGLE_SETUP.md, step 3.)')
            who = self.account or 'your Google account'
            if 'has not been used' in msg or 'disabled' in msg:
                raise SyncError('The Google Sheets API or Google Drive API is not enabled in the Google Cloud '
                                'project (see GOOGLE_SETUP.md, step 2).')
            if 'PERMISSION_DENIED' in msg or '403' in msg or name == 'PermissionError':
                raise SyncError(f'The sheet is not shared with {who}. The owner of the sheet: Share -> add '
                                f'that address as Editor.')
            if 'NOT_FOUND' in msg or '404' in msg or name == 'SpreadsheetNotFound':
                raise SyncError(f'No Google Sheet with this link, or it is not shared with {who}.')
            raise SyncError(f'Could not connect: {name}: {msg[:300]}')

    def export_xlsx(self):
        sh = self.sh
        return sh.export(format=self._gspread.utils.ExportFormat.EXCEL)

    def sheet_sizes(self):
        return {w.title: (w.row_count, w.col_count) for w in self.sh.worksheets()}

    def get_values(self, titles):
        if not titles:
            return {}
        q = [f"'{t.replace(chr(39), chr(39) * 2)}'" for t in titles]
        resp = self.sh.values_batch_get(q, params={'valueRenderOption': 'FORMULA',
                                                   'dateTimeRenderOption': 'SERIAL_NUMBER'})
        return {t: vr.get('values', []) for t, vr in zip(titles, resp.get('valueRanges', []))}

    def ensure_size(self, title, rows, cols):
        ws = self.sh.worksheet(title)
        if ws.row_count < rows or ws.col_count < cols:
            ws.resize(rows=max(ws.row_count, rows + 50), cols=max(ws.col_count, cols))

    def add_sheet(self, title, rows, cols):
        self.sh.add_worksheet(title=title, rows=max(rows + 100, 200), cols=max(cols, 26))

    def update(self, ranges):
        for i in range(0, len(ranges), 400):
            data = [{'range': r, 'values': v} for r, v in ranges[i:i + 400]]
            self.sh.values_batch_update(body={'valueInputOption': 'USER_ENTERED', 'data': data})

    def append_log(self, row):
        if 'ChangeLog' not in [w.title for w in self.sh.worksheets()]:
            self.sh.add_worksheet(title='ChangeLog', rows=1000, cols=6)
            self.sh.values_append("'ChangeLog'!A1", params={'valueInputOption': 'USER_ENTERED'},
                                  body={'values': [['When', 'Who', 'Computer', 'Changes']]})
        self.sh.values_append("'ChangeLog'!A1", params={'valueInputOption': 'USER_ENTERED'},
                              body={'values': [row]})


# --------------------------------------------------------------------------
# Reading cage-card photos with Google Gemini (free API key from aistudio.google.com)
# --------------------------------------------------------------------------
GEMINI_DEFAULT_MODEL = 'gemini-3.8-flash'
GEMINI_FALLBACK_MODEL = 'gemini-flash-latest'
PHOTO_EXT = ('.jpg', '.jpeg', '.png', '.heic', '.heif', '.webp')


class PhotoReadError(Exception):
    pass


CARD_PROMPT = """You read records of the Volgyi lab mouse colony (University of Pecs, Hungary) from a photo.
The photo shows either cage cards (one or several) or a page of the animal-house log book ("EGER" sheet).
Return ONLY JSON: {"cards": [ ... ], "events": [ ... ]} - one object per cage card in "cards", one object per
row of the log page in "events". Leave a list empty when the photo has none of that kind.

=== CAGE CARDS ===

The printed Hungarian card fields:
- "Ketrec kod" = cage code, written large, e.g. I/B/1/b (Roman numeral room / rack letter / shelf number / position
  letter). A handwritten "1" often looks like "/1" or "l" - write it as 1. Use the form I/B/1/b.
- "Labor", "Kutatasvezeto" = lab and PI: ignore.
- "Egerek szama" = number of mice. Earlier counts are crossed out; use the LAST one that is not crossed out.
  If you are not sure which is current, put "count" in "uncertain".
- "Nem" = sex: the female sign = "F", the male sign = "M", both signs or "par"/"tenyesz" = a breeding pair.
- "Torzs" = mouse line, as written. Map it to one of the KNOWN LINES below and put the card's own spelling in
  "line_as_written" (e.g. "Cx3CR1 x B6" -> "CX3CR1KO + Bl6"; B6/Bl6/C57 = C57Bl6J). If none fits, give the
  written text as "line" and add "line" to "uncertain".
- "Szarmazasi hely" = cage of origin (where the mice came from), a cage code -> "origin".
- "Szuletes ideje" = birth date, written yy.mm.dd (26.09.04 = 2026-09-04) -> "birth" as yyyy-mm-dd.
- "Behozatal ideje" = date they arrived in this cage, same format -> "arrived".
- "Kiserleti procedura" = experimental procedure: put any text into "notes".

Per card:
{"cage": "...", "type": "stock" | "pair" | "empty", "line": "...", "line_as_written": "...", "sex": "M"|"F",
 "count": n, "birth": "yyyy-mm-dd", "origin": "...", "arrived": "yyyy-mm-dd", "notes": "...",
 "uncertain": [field names that were hard to read or ambiguous]}
For a breeding-pair card use "type": "pair" and, if the birth dates of the parents are written,
"males": [{"birth": "yyyy-mm-dd"}], "females": [{"birth": "..."}, ...].
Leave out fields that are empty on the card. Never guess a value you cannot see: leave it out instead.
If a card is too blurry to read: {"cage": "<if readable>", "unreadable": true, "notes": "why"}.
Today is {today}; birth and arrival dates are never in the future.

=== ANIMAL-HOUSE LOG PAGE ("EGER") ===
A table with the columns: "Datum" (date, often written on two lines: 2025 / 10.28), "Allvany/oldal/sor/doboz"
(the cage code, same form as on the cards, e.g. II/A/4/d), "allat sorszama" (number and sex of the animals,
e.g. "4 male-sign" or "2 female-sign 1 male-sign"), "esemeny" (the event), "sulyossag" (severity, only for kivitel,
e.g. 4BA), "Kutato neve" (researcher), and two signature columns (ignore them).
A vertical line drawn in a cell means "same as the row above" (ditto); a row with an empty or ditto date has the
date of the row above. Read every row, top to bottom.
Events: "kivitel" = animals taken out of the animal house -> "removed". "athelyez"/"athelyezes" = animals moved to
another cage, the target cage is written under the word -> "moved". "levalaszt"/"levalasztas" = weaning, the new
cages are written under the word -> "weaned". Animals found dead ("elpusztult", "elhullott") -> "died".
Anything else -> "other".
Per row:
{"date": "yyyy-mm-dd", "cage": "...", "event": "removed" | "moved" | "weaned" | "died" | "other",
 "event_as_written": "kivitel", "males": n, "females": n,
 "to": [{"cage": "...", "males": n, "females": n}],   (only for moved / weaned: where the animals went;
        when several target cages are listed, give each one the animals that went there - usually males to the
        first and females to the second; if you cannot tell, put "to" in "uncertain")
 "severity": "4BA", "person": "Szarka G", "notes": "...", "uncertain": [field names]}

KNOWN LINES: {lines}
KNOWN CAGE CODES: {cages}
"""


def chat_instructions(col, today=None):
    """Instructions to attach in the Gemini (or any AI) chat together with card photos."""
    today = today or dt.date.today()
    head = ('MOUSE COLONY - INSTRUCTIONS FOR READING CAGE CARDS AND ANIMAL-HOUSE LOG PAGES\n'
            '(attach this file together with the photos)\n\n'
            'Task: read every cage card and every log row in the attached photos and answer with ONLY the JSON '
            'below in a single code block, nothing else. The JSON must look exactly like this:\n\n'
            '{\n  "format": "mouse-colony-cards",\n  "version": 1,\n  "read_on": "<today, yyyy-mm-dd>",\n'
            '  "source": "<number> photos read in the Gemini chat",\n  "cards": [ ... ],\n  "events": [ ... ]\n}\n\n'
            'Rules:\n\n')
    body = card_prompt(col, today).replace('Return ONLY JSON: {"cards": [ ... ], "events": [ ... ]} - one',
                                           'One')
    body = body.replace(f'Today is {today.isoformat()};', 'Use the real date of today;')
    tail = ('\nAlso:\n- The same card can appear in two photos (e.g. at the edge of another photo): list it only once.\n'
            '- Dates may also be written in full (2026.10.02); always output yyyy-mm-dd.\n'
            '- After the JSON, nothing else. No explanations.\n')
    return head + body + tail


def cards_from_text(text):
    """A cage-card file pasted as text (also from a chat answer with ``` fences or words around it)."""
    import json
    t = (text or '').strip()
    if '```' in t:
        parts = t.split('```')
        t = max(parts[1::2] or [t], key=len)
        if t.lstrip().lower().startswith('json'):
            t = t.lstrip()[4:]
    i, j = t.find('{'), t.rfind('}')
    if i < 0 or j < i:
        raise PhotoReadError('No cage-card data found in the pasted text.')
    try:
        obj = json.loads(t[i:j + 1])
    except Exception as e:
        raise PhotoReadError(f'The pasted text is not valid JSON ({e}).')
    if isinstance(obj, dict) and ('cards' in obj or 'events' in obj):
        obj.setdefault('format', 'mouse-colony-cards')
        obj.setdefault('cards', [])
        obj.setdefault('events', [])
        obj.setdefault('read_on', dt.date.today().isoformat())
        return obj
    raise PhotoReadError('The pasted text has no "cards" or "events" list.')


def _photo_bytes(path, max_side=2048):
    """Photo as (bytes, mime). Downscaled JPEG when Pillow is available (phone photos are large)."""
    ext = os.path.splitext(path)[1].lower()
    try:
        from PIL import Image, ImageOps
        try:
            import pillow_heif                     # optional: iPhone HEIC photos
            pillow_heif.register_heif_opener()
        except Exception:
            pass
        import io
        im = ImageOps.exif_transpose(Image.open(path)).convert('RGB')
        im.thumbnail((max_side, max_side))
        buf = io.BytesIO()
        im.save(buf, 'JPEG', quality=88)
        return buf.getvalue(), 'image/jpeg'
    except Exception:
        mime = {'.jpg': 'image/jpeg', '.jpeg': 'image/jpeg', '.png': 'image/png', '.webp': 'image/webp',
                '.heic': 'image/heic', '.heif': 'image/heif'}.get(ext, 'image/jpeg')
        with open(path, 'rb') as fh:
            data = fh.read()
        if len(data) > 18 * 1024 * 1024:
            raise PhotoReadError(f'{os.path.basename(path)} is too large ({len(data) // 2**20} MB).')
        return data, mime


def card_prompt(col, today=None):
    lines = sorted(col.lines) if col is not None else []
    cages = sorted({c['cage'] for c in col.cages} | {p['cage'] for p in col.pairs}) if col is not None else []
    return (CARD_PROMPT.replace('{today}', (today or dt.date.today()).isoformat())
            .replace('{lines}', '; '.join(lines) or '(none)')
            .replace('{cages}', ', '.join(col.card_name(c) for c in cages[:400]) or '(none)'))


def _gemini_call(key, model, prompt, data, mime, timeout=120):
    import json, base64, urllib.request, urllib.error
    url = f'https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent'
    body = {'contents': [{'role': 'user', 'parts': [
                {'text': prompt},
                {'inline_data': {'mime_type': mime, 'data': base64.b64encode(data).decode('ascii')}}]}],
            'generationConfig': {'responseMimeType': 'application/json', 'temperature': 0}}
    req = urllib.request.Request(url, data=json.dumps(body).encode('utf-8'), method='POST',
                                 headers={'Content-Type': 'application/json', 'x-goog-api-key': key})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            resp = json.loads(r.read().decode('utf-8'))
    except urllib.error.HTTPError as e:
        try:
            err = json.loads(e.read().decode('utf-8')).get('error', {})
        except Exception:
            err = {}
        msg, status = err.get('message', '') or str(e), err.get('status', '')
        if e.code == 404 or status == 'NOT_FOUND':
            raise PhotoReadError(f'MODEL:{msg}')
        if e.code in (500, 502, 503, 504) or status in ('UNAVAILABLE', 'INTERNAL', 'DEADLINE_EXCEEDED'):
            raise PhotoReadError(f'BUSY:Gemini is busy right now (error {e.code}: {msg[:160]}). Try again in a few '
                                 f'minutes.')
        if e.code == 429 or status == 'RESOURCE_EXHAUSTED':
            raise PhotoReadError('The free Gemini limit is used up for now (per minute or per day). Try again '
                                 'later, or send the photos to Claude as before.')
        if e.code in (400, 401, 403) and ('API key' in msg or 'API_KEY' in msg or status in ('UNAUTHENTICATED',)):
            raise PhotoReadError('Google does not accept the Gemini API key. Check it in Card photo settings '
                                 '(make a new one at aistudio.google.com -> API keys).')
        if 'location' in msg.lower() or 'region' in msg.lower() or 'country' in msg.lower():
            raise PhotoReadError(f'Gemini is not available here on this plan: {msg}')
        raise PhotoReadError(f'Gemini error {e.code}: {msg[:300]}')
    except urllib.error.URLError as e:
        raise PhotoReadError(f'Cannot reach Google (no internet?): {e.reason}')
    try:
        cand = resp['candidates'][0]
        text = ''.join(p.get('text', '') for p in cand['content']['parts'])
    except Exception:
        reason = (resp.get('promptFeedback') or {}).get('blockReason') or \
            ((resp.get('candidates') or [{}])[0].get('finishReason'))
        raise PhotoReadError(f'Gemini returned no answer ({reason or "empty"}).')
    return text


def _parse_records_json(text):
    """Gemini's answer -> (cards, events)."""
    import json
    t = text.strip()
    if t.startswith('```'):
        t = t.strip('`')
        t = t[t.find('{'):]
    try:
        obj = json.loads(t)
    except Exception:
        i, j = t.find('{'), t.rfind('}')
        try:
            obj = json.loads(t[i:j + 1])
        except Exception:
            raise PhotoReadError('Gemini did not answer in the expected format.')
    cards = obj.get('cards', []) if isinstance(obj, dict) else obj if isinstance(obj, list) else None
    events = obj.get('events', []) if isinstance(obj, dict) else []
    if not isinstance(cards, list) or not isinstance(events, list):
        raise PhotoReadError('Gemini did not answer in the expected format.')
    return [c for c in cards if isinstance(c, dict)], [e for e in events if isinstance(e, dict)]


def read_card_photo(path, key, model=None, col=None, today=None, prompt=None):
    """Read one photo (cage cards or a log page); returns {'cards': [...], 'events': [...]}."""
    if not key:
        raise PhotoReadError('No Gemini API key yet: Card photo settings... (free key from aistudio.google.com).')
    data, mime = _photo_bytes(path)
    prompt = prompt or card_prompt(col, today)
    import time
    model = model or GEMINI_DEFAULT_MODEL
    # busy (503): wait and retry, then try the fallback model; unknown model: go to the fallback straight away
    plan = [(model, 0), (model, 4), (model, 12)] + ([(GEMINI_FALLBACK_MODEL, 2), (GEMINI_FALLBACK_MODEL, 10)]
                                                    if model != GEMINI_FALLBACK_MODEL else [])
    text, last = None, None
    for i, (m_, wait) in enumerate(plan):
        if last is not None and str(last).startswith('MODEL:') and m_ == model:
            continue
        time.sleep(wait)
        try:
            text = _gemini_call(key, m_, prompt, data, mime)
            break
        except PhotoReadError as e:
            if not (str(e).startswith('BUSY:') or str(e).startswith('MODEL:')):
                raise
            last = e
    if text is None:
        raise PhotoReadError(str(last).replace('BUSY:', '').replace('MODEL:', 'Unknown Gemini model: '))
    cards, events = _parse_records_json(text)
    for c in cards + events:
        c['photo'] = os.path.basename(path)
        if isinstance(c.get('uncertain'), str):
            c['uncertain'] = [c['uncertain']]
    return {'cards': cards, 'events': events}


def _row_empty(mat, r):
    return r - 1 >= len(mat) or all(_norm(v) == '' for v in mat[r - 1])


def _last_row(mat):
    for r in range(len(mat), 0, -1):
        if not _row_empty(mat, r):
            return r
    return 0


def _shift_formula(v, old, new):
    """Same-row references in a formula of a moved row (Settings!$B$2 and other sheets are left alone)."""
    if not _is_formula(v) or old == new:
        return v
    return re.sub(r'(?<![!\w$])(\$?[A-Z]{1,3}\$?)' + str(old) + r'(?!\d)', lambda m: m.group(1) + str(new), v)


ID_SHEETS = {'Mice': ('M', 4), 'Breeders': ('B', 3)}


def rebase_appends(changes, base_wb, current):
    """Two people added rows at the same place (e.g. both logged a procedure): move OUR new rows below
    theirs, and give our new mice / breeders the next free ID if theirs took it. Returns a short summary."""
    notes = []
    for t, cells in changes.items():
        if t not in current or t not in base_wb.sheetnames:
            continue
        base, cur = _sheet_matrix(base_wb[t]), current[t]
        rows = sorted({r for r, _, _ in cells})
        appended = [r for r in rows if r > 1 and _row_empty(base, r)]
        taken = [r for r in appended if not _row_empty(cur, r)]
        if not taken:
            continue
        first = min(taken)
        move = [r for r in appended if r >= first]
        start = max([_last_row(cur)] + [r for r in rows if r not in move]) + 1
        mp = {old: start + i for i, old in enumerate(move)}
        changes[t] = [(mp.get(r, r), c, _shift_formula(v, r, mp[r]) if r in mp else v) for r, c, v in cells]
        notes.append(f'{t}: {len(move)} new row(s) placed after a colleague\'s')
    # IDs of our new mice / breeders already used by the colleague's rows -> next free ones
    idmap = {}
    for t, (pre, width) in ID_SHEETS.items():
        if t not in changes or t not in current:
            continue
        base = _sheet_matrix(base_wb[t]) if t in base_wb.sheetnames else []
        theirs = {str(row[0]).strip() for row in current[t][1:] if row and str(row[0] or '').strip()}
        mine = [(r, v) for r, c, v in changes[t] if c == 1 and isinstance(v, str) and v.startswith(pre)
                and _row_empty(base, r)]
        nums = [int(x[len(pre):]) for x in theirs | {v for _, v in mine}
                if x.startswith(pre) and x[len(pre):].isdigit()]
        nxt = max(nums or [0]) + 1
        for _, v in mine:
            if v in theirs:
                idmap[v] = f'{pre}{nxt:0{width}d}'
                nxt += 1
    if idmap:
        pat = re.compile(r'\b(' + '|'.join(re.escape(k) for k in idmap) + r')\b')
        for t in changes:
            changes[t] = [(r, c, pat.sub(lambda m: idmap[m.group(1)], v) if isinstance(v, str) else v)
                          for r, c, v in changes[t]]
        notes.append('new IDs ' + ', '.join(f'{a}->{b}' for a, b in idmap.items()))
    return '; '.join(notes)


def push_changes(remote, base_wb, new_wb, user='', force=False):
    """Write the cells that differ between base_wb (as downloaded) and new_wb to the Google Sheet.
    Refuses (ConflictError) if someone else changed one of those cells since the download."""
    changes, new_sheets = diff_workbooks(base_wb, new_wb)
    if not changes:
        return 0
    sizes = remote.sheet_sizes()
    for t in new_sheets:
        if t in sizes and not force:
            raise ConflictError(f'Someone else just created the sheet "{t}".')
    existing = [t for t in changes if t in sizes and t not in new_sheets]
    rebased = ''
    if not force:
        current = remote.get_values(existing)
        rebased = rebase_appends(changes, base_wb, current)
        clashes = []
        for t in existing:
            base = _sheet_matrix(base_wb[t])
            for r, c, _ in changes[t]:
                b, cur = _at(base, r, c), _at(current.get(t, []), r, c)
                if _is_formula(b) or _is_formula(cur):
                    continue
                if _norm(b) != _norm(cur):
                    clashes.append(f'{t}!{_col_letter(c)}{r}')
        if clashes:
            raise ConflictError('Someone else changed the same cells meanwhile ('
                                + ', '.join(clashes[:6]) + (' ...' if len(clashes) > 6 else '') + ').')
    ranges = []
    for t, cells in changes.items():
        need_r = max(r for r, _, _ in cells)
        need_c = max(c for _, c, _ in cells)
        if t in new_sheets and t not in sizes:
            remote.add_sheet(t, need_r, need_c)
        else:
            remote.ensure_size(t, need_r, need_c)
        ranges += group_ranges(t, cells)
    remote.update(ranges)
    try:
        import socket
        summary = ', '.join(f'{t}: {len(c)} cells' for t, c in changes.items()) + (
            f' ({rebased})' if rebased else '')
        who = (user or '?') + (f' ({remote.account})' if getattr(remote, 'account', '') and
                               not (getattr(remote, 'key_info', None)) else '')
        remote.append_log([dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S'), who, socket.gethostname(),
                           summary])
    except Exception:
        pass
    return sum(len(c) for c in changes.values())


# --------------------------------------------------------------------------
# Data model
# --------------------------------------------------------------------------
class Colony:
    def __init__(self, path=None, today=None, remote=None, backup_dir=None, user='', io_hook=None):
        """path: local workbook. remote: a GoogleBackend (the shared Google Sheet; a local copy is kept)."""
        self.remote = remote
        self.user = user
        self.io_hook = io_hook       # GUI: runs slow network calls in the background (fn, message) -> result
        self.on_change = None        # GUI: called after every local save in shared mode (schedules the upload)
        self._local_only = False
        self._version = 0            # bumped on every local save (tells a finished upload whether we changed more)
        self.pending = 0             # shared mode: local changes not yet uploaded
        self.readonly = False
        self.offline_reason = ''
        self.backup_dir = backup_dir
        self.today = today or dt.date.today()
        if remote:
            os.makedirs(self.backup_folder(), exist_ok=True)
            self.path = os.path.join(self.backup_folder(), 'shared_colony_latest.xlsx')
            self.base_path = os.path.join(self.backup_folder(), 'shared_colony_base.xlsx')
            self.state_path = os.path.join(self.backup_folder(), 'shared_colony_state.json')
            self._read_state()
        else:
            self.path = path
        self._started = False
        self.load()
        self._started = True

    # ---------------- local backups ----------------
    def backup_folder(self):
        if self.backup_dir:
            return self.backup_dir
        if self.remote:
            return os.path.join(os.path.expanduser('~'), 'Documents', 'Mouse colony', 'backups')
        return os.path.join(os.path.dirname(os.path.abspath(self.path)), 'backups')

    def _backup_name(self):
        if self.remote:
            safe = re.sub(r'[^A-Za-z0-9_-]+', '_', getattr(self.remote, 'title', 'shared') or 'shared')
            return f'shared_{safe}'
        return os.path.splitext(os.path.basename(self.path))[0]

    def make_backup(self):
        """Copy the current workbook into the backup folder and thin out old copies."""
        if not self.path or not os.path.exists(self.path):
            return None
        folder = self.backup_folder()
        os.makedirs(folder, exist_ok=True)
        stamp = dt.datetime.now().strftime('%Y%m%d-%H%M%S')
        dst = os.path.join(folder, f'{self._backup_name()}_{stamp}.xlsx')
        shutil.copy2(self.path, dst)
        self.prune_backups()
        return dst

    def list_backups(self):
        folder = self.backup_folder()
        out = []
        if not os.path.isdir(folder):
            return out
        pat = re.compile(re.escape(self._backup_name()) + r'_(\d{8}-\d{6})\.xlsx$')
        for f in os.listdir(folder):
            m = pat.match(f)
            if m:
                out.append((dt.datetime.strptime(m.group(1), '%Y%m%d-%H%M%S'), os.path.join(folder, f)))
        return sorted(out, reverse=True)

    def prune_backups(self, now=None):
        """Keep every copy of the last 7 days, the newest copy per day for 90 days, then one per month."""
        now = now or dt.datetime.now()
        keep, seen_day, seen_month = set(), set(), set()
        for when, path in self.list_backups():          # newest first
            age = (now - when).days
            if age <= 7:
                keep.add(path)
            elif age <= 90:
                if when.date() not in seen_day:
                    seen_day.add(when.date())
                    keep.add(path)
            else:
                if (when.year, when.month) not in seen_month:
                    seen_month.add((when.year, when.month))
                    keep.add(path)
        for _, path in self.list_backups():
            if path not in keep:
                try:
                    os.remove(path)
                except OSError:
                    pass

    def restore_backup(self, backup_path):
        """Make a backup copy the current version (the current state is backed up first)."""
        if self.readonly:
            raise ReadOnlyError('The shared sheet cannot be reached - restoring needs a connection.')
        self.make_backup()
        if self.remote:
            self.sync_now()                    # upload anything unsent and get the latest first
            base = openpyxl.load_workbook(self.path)
            new = openpyxl.load_workbook(backup_path)
            self._io(lambda: push_changes(self.remote, base, new, self.user, force=True),
                     'Restoring the shared sheet...')
        else:
            shutil.copy2(backup_path, self.path)
        self.load()
        self.make_backup()

    # ---------------- shared sheet: local first, upload in the background ----------------
    def _read_state(self):
        import json
        try:
            with open(self.state_path, encoding='utf-8') as fh:
                st = json.load(fh)
        except Exception:
            st = {}
        sid = getattr(self.remote, 'sheet_id', '')
        ok = st.get('sheet_id') == sid and os.path.exists(self.path) and os.path.exists(self.base_path)
        if st.get('pending') and not ok and os.path.exists(self.path):
            # unsent changes belong to another sheet: keep them as a file, don't upload them here
            shutil.copy2(self.path, os.path.join(self.backup_folder(),
                                                 f'unsent_changes_{dt.datetime.now():%Y%m%d-%H%M%S}.xlsx'))
        self.pending = int(st.get('pending') or 0) if ok else 0

    def _write_state(self):
        import json
        try:
            with open(self.state_path, 'w', encoding='utf-8') as fh:
                json.dump({'sheet_id': getattr(self.remote, 'sheet_id', ''), 'pending': self.pending,
                           'saved': dt.datetime.now().isoformat(timespec='seconds')}, fh)
        except OSError:
            pass

    def prepare_sync(self, force=False):
        """Main thread: snapshot of the local copy for a background upload (+ download of the latest)."""
        snap = self.path[:-5] + '_upload.xlsx'
        if self.pending:
            shutil.copy2(self.path, snap)
        return dict(snap=snap, version=self._version, pending=self.pending, force=force)

    def do_sync(self, job):
        """Background thread (no GUI, no Colony state changes): upload the local changes, then fetch the sheet."""
        pushed = 0
        if job['pending']:
            base = openpyxl.load_workbook(self.base_path)
            new = openpyxl.load_workbook(job['snap'])
            pushed = push_changes(self.remote, base, new, self.user, force=job.get('force', False))
        data = self.remote.export_xlsx()
        return dict(job, data=data, pushed=pushed)

    def finish_sync(self, res):
        """Main thread: take over the result of do_sync. Returns True if the shown data changed."""
        self.readonly, self.offline_reason = False, ''
        if self._version == res['version']:
            tmp = self.path + '.part'
            with open(tmp, 'wb') as fh:
                fh.write(res['data'])
            os.replace(tmp, self.path)
            shutil.copy2(self.path, self.base_path)
            self.pending = 0
            self._write_state()
            self._local_only = True
            try:
                self.load()
            finally:
                self._local_only = False
            if res['pending']:
                self.make_backup()
            return True
        # more local changes arrived during the upload: what was uploaded becomes the new base, the rest waits
        if res['pending']:
            shutil.copy2(res['snap'], self.base_path)
            self.pending = max(1, self.pending - res['pending'])
            self._write_state()
        return False

    def sync_failed(self, err):
        if not isinstance(err, ConflictError):
            self.readonly = True
            self.offline_reason = str(err) or err.__class__.__name__

    def sync_now(self, force=False):
        """Upload and refresh in one go (through io_hook when there is one). Raises on failure."""
        job = self.prepare_sync(force)
        try:
            res = self._io(lambda: self.do_sync(job), 'Saving to the shared Google Sheet...' if job['pending']
                           else 'Getting the latest version of the shared sheet...')
        except Exception as e:
            self.sync_failed(e)
            raise
        return self.finish_sync(res)

    def take_theirs(self):
        """Conflict: drop the unsent local changes (kept as a file) and use the shared sheet."""
        keep = os.path.join(self.backup_folder(), f'unsent_changes_{dt.datetime.now():%Y%m%d-%H%M%S}.xlsx')
        shutil.copy2(self.path, keep)
        self.pending = 0
        self._write_state()
        self.load()
        return keep

    def _io(self, fn, msg=''):
        hook = getattr(self, 'io_hook', None)
        return hook(fn, msg) if hook else fn()

    def _download(self):
        """Fetch the shared sheet into the local copy; fall back to the last copy read-only."""
        try:
            data = self._io(self.remote.export_xlsx, 'Getting the latest version of the shared sheet...')
            tmp = self.path + '.part'
            with open(tmp, 'wb') as fh:
                fh.write(data)
            os.replace(tmp, self.path)
            shutil.copy2(self.path, self.base_path)
            self.readonly = False
            self.offline_reason = ''
        except Exception as e:
            if os.path.exists(self.path):
                self.readonly = True
                self.offline_reason = str(e) or e.__class__.__name__
            else:
                raise SyncError(f'Could not download the shared sheet and there is no local copy yet: {e}')

    # ---------------- loading ----------------
    def load(self):
        outer = not getattr(self, '_in_load', False)
        self._in_load = True
        try:
            self._load()
        finally:
            if outer:
                self._in_load = False

    def _load(self):
        if self.remote and not getattr(self, '_local_only', False) and not (
                self.pending and os.path.exists(self.path)):
            self._download()
            if not getattr(self, '_started', True):
                self.make_backup()              # a local copy at every start
        wb = openpyxl.load_workbook(self.path)  # formulas, not values: we compute ages ourselves
        self.settings = self._read_settings(wb['Settings'])
        self.lines = {}
        has_parts = bool(wb['Lines'].cell(1, 14).value)
        for r in wb['Lines'].iter_rows(min_row=2, max_col=14, values_only=True):
            if _s(r[0]):
                partners = [p.strip() for p in _s(r[1]).split(';') if p.strip()] or [_s(r[0])]
                loci_txt = _s(r[12]) if wb['Lines'].cell(1, 13).value else DEFAULT_LOCI.get(_s(r[0]), '')
                parts_txt = _s(r[13]) if has_parts else DEFAULT_PARTS.get(_s(r[0]), '')
                self.lines[_s(r[0])] = dict(partners=partners, keep=_s(r[2]).upper() == 'Y', background=_s(r[3]),
                                            notes=_s(r[5]),
                                            parts=[x.strip() for x in parts_txt.split(';') if x.strip()],
                                            loci=[x.strip() for x in loci_txt.split(';') if x.strip()])
        self.breeders = {}
        self.breeder_rows = {}
        for i, r in enumerate(wb['Breeders'].iter_rows(min_row=2, max_col=14, values_only=True), 2):
            bid = _s(r[0])
            if not bid:
                continue
            self.breeders[bid] = dict(id=bid, line=_s(r[1]), sex=_s(r[2]).upper(), birth=_d(r[3]),
                                      cage=_s(r[4]), status=_s(r[5]) or 'Breeding',
                                      sire=_s(r[6]), dam=_s(r[7]), notes=_s(r[8]), gt=parse_gt(_s(r[13])))
            self.breeder_rows[bid] = i
        self.pairs = []
        for i, r in enumerate(wb['Pairs'].iter_rows(min_row=2, max_col=8, values_only=True), 2):
            if not _s(r[0]):
                continue
            self.pairs.append(dict(row=i, cage=_s(r[0]), line=_s(r[1]), m=_s(r[2]), f1=_s(r[3]), f2=_s(r[4]),
                                   setup=_d(r[5]), status=_s(r[6]) or 'Active', notes=_s(r[7])))
        self.cages = []
        ws = wb['Cages']
        for i in range(2, ws.max_row + 1):
            name = _s(ws.cell(i, 1).value)
            if not name:
                continue
            cnt = ws.cell(i, 5).value
            self.cages.append(dict(row=i, cage=name, line=_s(ws.cell(i, 2).value), sex=_s(ws.cell(i, 4).value).upper(),
                                   count=int(cnt) if isinstance(cnt, (int, float)) else None,
                                   birth=_d(ws.cell(i, 6).value), sire=_s(ws.cell(i, 10).value),
                                   dam=_s(ws.cell(i, 11).value), note=_s(ws.cell(i, 12).value),
                                   reserved=_s(ws.cell(i, 13).value), gt=parse_gt(_s(ws.cell(i, 14).value))))
        self.event_keys = {}
        if EVENTS_SHEET in wb.sheetnames:
            for r in wb[EVENTS_SHEET].iter_rows(min_row=2, max_col=7, values_only=True):
                if _s(r[0]):
                    self.event_keys[_s(r[0])] = _d(r[5]) or _s(r[5])
        self.procedures = []
        if 'Procedures' in wb.sheetnames:
            for i, r in enumerate(wb['Procedures'].iter_rows(min_row=2, max_col=11, values_only=True), 2):
                if not any(r):
                    continue
                self.procedures.append(dict(row=i, date=_d(r[0]), cage=_s(r[1]), animal=_s(r[2]), line=_s(r[3]),
                                            sex=_s(r[4]), n=r[5], birth=_d(r[6]), proc=_s(r[7]),
                                            details=_s(r[8]), person=_s(r[9]), notes=_s(r[10])))
        self.terminated = []
        if 'Terminated' in wb.sheetnames:
            for i, r in enumerate(wb['Terminated'].iter_rows(min_row=2, max_col=17, values_only=True), 2):
                if not any(r):
                    continue
                self.terminated.append(dict(row=i, date=_d(r[0]), cage=_s(r[1]), animal=_s(r[2]), line=_s(r[3]),
                                            sex=_s(r[4]), n=r[5] if isinstance(r[5], (int, float)) else 0,
                                            birth=_d(r[6]), age=r[7], outcome=_s(r[8]), reason=_s(r[9]),
                                            project=_s(r[10]), person=_s(r[11]), sire=_s(r[12]),
                                            dam=_s(r[13]), gt=_s(r[14]), procs=_s(r[15]), notes=_s(r[16])))
        self.mice = {}
        if 'Mice' in wb.sheetnames:
            for i, r in enumerate(wb['Mice'].iter_rows(min_row=2, max_col=14, values_only=True), 2):
                mid = _s(r[0])
                if not mid:
                    continue
                self.mice[mid] = dict(row=i, id=mid, cage=_s(r[1]), mark=_s(r[2]) or 'none', line=_s(r[3]),
                                      sex=_s(r[4]).upper(), birth=_d(r[5]), sire=_s(r[6]), dam=_s(r[7]),
                                      gt=parse_gt(_s(r[8])), status=_s(r[9]) or 'Alive', since=_d(r[10]),
                                      breeder=_s(r[11]), notes=_s(r[12]), todo=_s(r[13]).upper() in ('Y', 'YES'))
        self._kin = {}
        if 'Mice' not in wb.sheetnames:          # first start of v1.3: give every stock mouse an ID + ear mark
            self._save(self._open_for_write())
            return
        if not wb['Lines'].cell(1, 14).value:    # v1.6: "Made of" column for mixed lines
            self._save(self._open_for_write())
            return
        if not wb['Mice'].cell(1, 14).value:     # v1.7: "Needs genotyping" flag
            self._save(self._open_for_write())
            return
        if 'MaxAge' in str(wb['Breeders']['L2'].value or '') and not self.readonly:   # v1.8: plain references
            self._save(self._open_for_write())
            return
        if not getattr(self, '_deduping', False):
            self._deduping = True
            try:
                while self._merge_auto_duplicates():
                    pass
            finally:
                self._deduping = False

    @staticmethod
    def _read_settings(ws):
        vals = {}
        names = {2: 'MaxAge', 3: 'WarnDays', 4: 'MinAge', 5: 'StartMax'}
        defaults = dict(MaxAge=365, WarnDays=60, MinAge=49, StartMax=180)
        for r, n in names.items():
            v = ws.cell(r, 2).value
            vals[n] = int(v) if isinstance(v, (int, float)) else defaults[n]
        return vals

    # ---------------- ages ----------------
    def age(self, birth):
        return (self.today - birth).days if birth else None

    def age_flag(self, birth, status='Breeding'):
        if status in GONE:
            return '-'
        a = self.age(birth)
        s = self.settings
        if a is None:
            return 'NO DATE'
        if a >= s['MaxAge']:
            return 'OVER AGE'
        if a >= s['MaxAge'] - s['WarnDays']:
            return 'Replace soon'
        if a < s['MinAge']:
            return 'Too young'
        return 'OK'

    def days_left(self, birth):
        a = self.age(birth)
        return None if a is None else self.settings['MaxAge'] - a

    def pair_members(self, p):
        return [x for x in (p['m'], p['f1'], p['f2']) if x]

    def pair_oldest(self, p):
        ages = [self.age(self.breeders[x]['birth']) for x in self.pair_members(p)
                if x in self.breeders and self.breeders[x]['birth']]
        return max(ages) if ages else None

    def pair_flag(self, p):
        if p['status'] != 'Active':
            return ''
        a = self.pair_oldest(p)
        if a is None:
            return 'NO DATE'
        s = self.settings
        if a >= s['MaxAge']:
            return 'OVER AGE'
        if a >= s['MaxAge'] - s['WarnDays']:
            return 'Replace soon'
        return 'OK'

    def pair_replace_by(self, p):
        a = self.pair_oldest(p)
        return None if a is None else self.today + dt.timedelta(days=self.settings['MaxAge'] - a)

    # ---------------- pedigree ----------------
    # A "unit" is either a breeder ID or a stock cage written as "cage:IB8c".
    # Animals in one stock cage are littermates and share sire/dam.
    def parents(self, unit):
        if unit.startswith('mouse:'):
            m = self.mice.get(unit[6:])
            if m and m['breeder'] in self.breeders:
                b = self.breeders[m['breeder']]
                return b['sire'], b['dam']
            return (m['sire'], m['dam']) if m else ('', '')
        if unit.startswith('cage:'):
            c = self.cage(unit[5:])
            return (c['sire'], c['dam']) if c else ('', '')
        b = self.breeders.get(unit)
        return (b['sire'], b['dam']) if b else ('', '')

    def _depth(self, unit, _stack=()):
        if not unit or unit in _stack:
            return 0
        s, d = self.parents(unit)
        if not s and not d:
            return 0
        return 1 + max(self._depth(s, _stack + (unit,)), self._depth(d, _stack + (unit,)))

    def inbreeding(self, unit):
        s, d = self.parents(unit)
        return self.kinship(s, d) if s and d else 0.0

    def kinship(self, a, b):
        """Coefficient of kinship (probability two alleles drawn one from each are IBD),
        relative to the recorded founders. Offspring of a x b has F = kinship(a, b)."""
        if not a or not b:
            return 0.0
        key = (a, b) if a <= b else (b, a)
        if key in self._kin:
            return self._kin[key]
        self._kin[key] = 0.0  # guards against loops in bad data
        if a == b:
            val = 0.5 * (1 + self.inbreeding(a))
        else:
            # expand the younger unit (an ancestor always has a smaller depth)
            if self._depth(a) < self._depth(b):
                a, b = b, a
            s, d = self.parents(a)
            if not s and not d:
                val = 0.0
            else:
                val = 0.5 * (self.kinship(s, b) + self.kinship(d, b))
        self._kin[key] = val
        return val

    def ancestors(self, unit, depth=6):
        out = {}
        frontier = [(unit, 0)]
        while frontier:
            u, g = frontier.pop()
            if g >= depth:
                continue
            for p in self.parents(u):
                if p:
                    out.setdefault(p, g + 1)
                    frontier.append((p, g + 1))
        return out

    def relationship(self, a, b):
        sa, da = self.parents(a)
        sb, db = self.parents(b)
        if a == b:
            return 'same animal'
        if a in (sb, db) or b in (sa, da):
            return 'parent and offspring'
        if sa and da and sa == sb and da == db:
            return 'full siblings'
        if (sa and sa == sb) or (da and da == db):
            return 'half siblings'
        common = set(self.ancestors(a)) & set(self.ancestors(b))
        if common:
            return 'share ancestor(s): ' + ', '.join(sorted(common)[:6])
        return 'no shared ancestors recorded'

    @staticmethod
    def risk(f):
        if f >= 0.25:
            return 'HIGH - do not pair'
        if f >= 0.125:
            return 'Elevated - avoid'
        if f >= 0.0625:
            return 'Moderate - ok only if no alternative'
        if f > 0:
            return 'Low'
        return 'None recorded'

    # ---------------- lookups ----------------
    def cage(self, name):
        for c in self.cages:
            if c['cage'] == name:
                return c
        return None

    def unit_info(self, unit):
        """Line, sex, birth, label for a breeder ID, 'mouse:M0001' or 'cage:X'."""
        if unit.startswith('mouse:'):
            m = self.mice.get(unit[6:])
            if not m:
                return None
            return dict(line=m['line'], sex=m['sex'], birth=m['birth'], status=m['status'] if m['status'] !=
                        'Alive' else 'Stock', gt=m['gt'],
                        label=f"{m['id']} ({self.card_name(m['cage'])} ear {m['mark']}, {m['sex']}, {m['line']})")
        if unit.startswith('cage:'):
            c = self.cage(unit[5:])
            if not c:
                return None
            return dict(line=c['line'], sex=c['sex'], birth=c['birth'], status='Stock', gt=c['gt'],
                        label=f"cage {c['cage']} ({c['count'] or '?'} {c['sex']}, {c['line']})")
        b = self.breeders.get(unit)
        if not b:
            return None
        return dict(line=b['line'], sex=b['sex'], birth=b['birth'], status=b['status'], gt=b.get('gt') or {},
                    label=f"{b['id']} ({b['sex']}, {b['line']}, cage {b['cage']})")

    def candidate_units(self, sex, lines=None, include_breeding=False):
        """Animals that could go into a new pair: Reserve breeders and young stock cages."""
        s = self.settings
        out = []
        for b in self.breeders.values():
            if b['sex'] != sex or (lines and b['line'] not in lines):
                continue
            if b['status'] == 'Reserve' or (include_breeding and b['status'] == 'Breeding'):
                a = self.age(b['birth'])
                if a is None or a < s['MaxAge'] - s['WarnDays']:
                    out.append(b['id'])
        for c in self.cages:
            if c['sex'] != sex or not c['birth'] or not c['count']:
                continue
            if lines and c['line'] not in lines:
                continue
            a = self.age(c['birth'])
            if s['MinAge'] <= a <= s['StartMax']:
                mice = [m for m in self.mice_in(c['cage']) if not m['breeder']]
                out += ['mouse:' + m['id'] for m in mice] if mice else ['cage:' + c['cage']]
        return out

    # ---------------- lines ----------------
    def line_parts(self, line):
        """'Cx36KO + PV-Cre' -> ['Cx36KO', 'PV-Cre'] (only splits into parts that are lines themselves,
        otherwise the name is kept whole)."""
        if line in self.lines and self.lines[line].get('parts'):
            return list(self.lines[line]['parts'])
        parts = [x.strip() for x in (line or '').split('+') if x.strip()]
        return parts if len(parts) > 1 else [line]

    def combined_name(self, lines):
        parts = []
        for ln in lines:
            for x in (self.line_parts(ln) if ln not in self.lines or LINE_SEP in ln
                      or self.lines[ln].get('parts') else [ln]):
                if x not in parts:
                    parts.append(x)
        # an existing line with the same parts (any order) is reused
        key = sorted(p.lower() for p in parts)
        for ln in self.lines:
            if sorted(p.lower() for p in self.line_parts(ln)) == key:
                return ln
        return LINE_SEP.join(parts)

    def combined_loci(self, lines):
        out = []
        for ln in lines:
            for part in [ln] + self.line_parts(ln):
                for x in self.lines.get(part, {}).get('loci', []) or [y.strip() for y in
                                                                     DEFAULT_LOCI.get(part, '').split(';')
                                                                     if y.strip()]:
                    if x not in out:
                        out.append(x)
        return out

    def add_line(self, name, loci=(), partners=(), keep=False, background='', notes='', parts=()):
        """New row on the Lines sheet (formulas included)."""
        name = name.strip()
        if not name:
            raise ValueError('The line needs a name.')
        if name in self.lines:
            raise ValueError(f'Line "{name}" already exists.')
        from openpyxl.styles import Font, PatternFill
        wb = self._open_for_write()
        ws = wb['Lines']
        r = self._first_free_row(ws)
        self._copy_style(ws, r - 1, r, 13)
        vals = [name, '; '.join(partners) or None, 'Y' if keep else None, background or None, None,
                notes or None] + line_formulas(r) + ['; '.join(loci) or None, '; '.join(parts) or None]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c)
            cell.value = v
            cell.font = Font(name='Arial', size=10)
            if 7 <= c <= 12:
                cell.font = Font(name='Arial', size=10)
                cell.fill = PatternFill('solid', fgColor='F2F2F2')
        self._save(wb)
        return name

    def update_line(self, name, loci=None, partners=None, keep=None, background=None, notes=None):
        wb = self._open_for_write()
        ws = wb['Lines']
        for r in range(2, ws.max_row + 1):
            if _s(ws.cell(r, 1).value) == name:
                for c, v in ((2, '; '.join(partners) if partners is not None else None),
                             (3, ('Y' if keep else '') if keep is not None else None), (4, background),
                             (6, notes), (13, '; '.join(loci) if loci is not None else None)):
                    if v is not None:
                        ws.cell(r, c).value = v or None
        self._save(wb)

    def merge_lines(self, old, into):
        """Rename line `old` to `into` everywhere and remove its Lines row."""
        if old == into or old not in self.lines or into not in self.lines:
            raise ValueError('Pick two different existing lines.')
        wb = self._open_for_write()
        for sheet, col in (('Breeders', 2), ('Pairs', 2), ('Cages', 2), ('Mice', 4), ('Procedures', 4),
                           ('Terminated', 4)):
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            for r in range(2, ws.max_row + 1):
                if _s(ws.cell(r, col).value) == old:
                    ws.cell(r, col).value = into
        ws = wb['Lines']
        drop = None
        for r in range(2, ws.max_row + 1):
            name = _s(ws.cell(r, 1).value)
            if name == old:
                drop = r
            elif ws.cell(r, 2).value and old in _s(ws.cell(r, 2).value):
                parts = [x.strip() for x in _s(ws.cell(r, 2).value).split(';') if x.strip()]
                ws.cell(r, 2).value = '; '.join(dict.fromkeys(into if x == old else x for x in parts))
        if drop:
            ws.delete_rows(drop)
            for r in range(2, ws.max_row + 1):        # formulas refer to their own row - rewrite them
                if ws.cell(r, 1).value:
                    for c, f in enumerate(line_formulas(r), 7):
                        ws.cell(r, c).value = f
        self._save(wb)

    def _merge_auto_duplicates(self):
        """A mixed line the app created ('Cross of ...') that has the same parts as an older line is merged
        into the older one."""
        names = list(self.lines)
        for i, a in enumerate(names):
            for b in names[i + 1:]:
                pa = sorted(x.lower() for x in self.line_parts(a))
                pb = sorted(x.lower() for x in self.line_parts(b))
                if len(pa) > 1 and pa == pb and self.lines[b].get('notes', '').startswith('Cross of'):
                    self.merge_lines(b, a)
                    return True
        return False

    def ensure_combined_line(self, lines):
        """Name of the line for a cross of these lines; created (with the union of alleles) if new."""
        lines = [x for x in lines if x]
        if len(lines) == 1:
            return lines[0]
        name = self.combined_name(lines)
        if name not in self.lines:
            parts = []
            for x in lines:
                for y in self.line_parts(x):
                    if y not in parts:
                        parts.append(y)
            self.add_line(name, loci=self.combined_loci(lines), partners=[name] + [x for x in lines], parts=parts,
                          background=' / '.join(dict.fromkeys(
                              self.lines.get(x, {}).get('background', '') for x in lines if x)).strip(' /'),
                          notes='Cross of ' + ' and '.join(lines))
        return name

    def partner_lines(self, line):
        return self.lines.get(line, {}).get('partners', [line])

    # ---------------- analysis ----------------
    def loci_for(self, *lines):
        out = []
        for ln in lines:
            for x in self.lines.get(ln, {}).get('loci', []):
                if x not in out:
                    out.append(x)
        return out

    def gt_of(self, unit, locus, _depth=0):
        inf = self.unit_info(unit) if unit else None
        if not inf:
            return '?'
        if inf['line'] == 'C57Bl6J':
            return 'wt'
        v = inf['gt'].get(locus)
        if v and v != '?':
            return v
        if locus not in self.loci_for(inf['line']):
            return 'wt'                      # the line does not carry this allele
        sire, dam = self.parents(unit)       # untested, but certain from the parents (e.g. wt x hom -> het)
        if sire and dam and _depth < 6:
            res = mendel(self.gt_of(sire, locus, _depth + 1), self.gt_of(dam, locus, _depth + 1))
            if res and len(res) == 1:
                return next(iter(res))
        return '?'

    def expected_pups(self, male, female, loci=None):
        """{locus: {'hom': p, ...} or None}"""
        mi, fi = self.unit_info(male), self.unit_info(female)
        if not mi or not fi:
            return {}
        loci = loci or self.loci_for(mi['line'], fi['line'])
        return {lc: mendel(self.gt_of(male, lc), self.gt_of(female, lc)) for lc in loci}

    def pups_text(self, male, female):
        parts = []
        for lc, res in self.expected_pups(male, female).items():
            parts.append(f"{lc}: {fmt_mendel(res) if res else '? (parent not genotyped)'}")
        return '; '.join(parts)

    def expected_gt_text(self, sire, dam):
        """For a stock cage with known parents but no genotype yet."""
        if not sire or not dam or sire not in self.breeders or dam not in self.breeders:
            return ''
        res = self.expected_pups(sire, dam)
        known = {k: v for k, v in res.items() if v}
        if not known:
            return ''
        return 'exp. ' + '; '.join(f'{k} ' + (GT_SHORT[next(iter(v))] if len(v) == 1 else fmt_mendel(v))
                                   for k, v in known.items())

    def gt_display(self, gt, sire='', dam='', line=''):
        txt = format_gt(gt)
        if txt:
            return txt
        if line == 'C57Bl6J':
            return 'wt'
        return self.expected_gt_text(sire, dam)

    def check_pair(self, male, females):
        """Returns list of (level, message). level: 'bad' | 'warn' | 'ok' | 'info'."""
        msgs = []
        units = [male] + [f for f in females if f]
        infos = {}
        for u in units:
            inf = self.unit_info(u)
            if not inf:
                msgs.append(('bad', f'{u}: not found'))
                return msgs
            infos[u] = inf
        s = self.settings
        if infos[male]['sex'] != 'M':
            msgs.append(('bad', f"{infos[male]['label']} is not a male"))
        for f in females:
            if f and infos[f]['sex'] != 'F':
                msgs.append(('bad', f"{infos[f]['label']} is not a female"))
        for u in units:
            inf = infos[u]
            a = self.age(inf['birth'])
            if inf['status'] == 'Terminated':
                msgs.append(('bad', f"{inf['label']} is marked Terminated"))
            elif inf['status'] == 'Retired':
                msgs.append(('warn', f"{inf['label']} is marked Retired"))
            if a is None:
                msgs.append(('warn', f"{inf['label']}: no birth date"))
            elif a >= s['MaxAge']:
                msgs.append(('bad', f"{inf['label']}: {fmt_age(a)} - over the {s['MaxAge']}-day limit"))
            elif a >= s['MaxAge'] - s['WarnDays']:
                msgs.append(('warn', f"{inf['label']}: {fmt_age(a)} - reaches the limit in {s['MaxAge'] - a} days"))
            elif a < s['MinAge']:
                msgs.append(('warn', f"{inf['label']}: only {a} days old (min {s['MinAge']})"))
            else:
                msgs.append(('ok', f"{inf['label']}: {fmt_age(a)}, can breed for {s['MaxAge'] - a} more days"))
        for f in females:
            if not f:
                continue
            k = self.kinship(male, f)
            rel = self.relationship(male, f)
            level = 'bad' if k >= 0.125 else 'warn' if k > 0 else 'ok'
            msgs.append((level, f"{male} x {f}: {rel}. Offspring inbreeding F = {k:.3f} ({self.risk(k)})"))
        for f in females:
            if not f:
                continue
            for lc, res in self.expected_pups(male, f).items():
                gm, gf = self.gt_of(male, lc), self.gt_of(f, lc)
                if res is None:
                    who = [n for n, g in (('male', gm), ('female', gf)) if _mut_prob(g) is None]
                    msgs.append(('warn', f"{lc}: {' and '.join(who)} genotype unknown "
                                         f"({GT_SHORT.get(gm, gm)} x {GT_SHORT.get(gf, gf)}) - genotype before pairing"))
                else:
                    useless = res.get('wt', 0) == 1
                    msgs.append(('warn' if useless else 'info',
                                 f"{lc}: male {GT_SHORT.get(gm, gm)} x female {GT_SHORT.get(gf, gf)} -> pups "
                                 f"{fmt_mendel(res)}"
                                 + (' (no carriers!)' if useless else '')))
        lines = {infos[u]['line'] for u in units}
        if len(lines) > 1:
            msgs.append(('info', 'Different lines: ' + ', '.join(sorted(lines)) +
                         '. Fine for a planned cross or Bl6 outcross; check it is intended.'))
        no_parents = [u for u in units if not any(self.parents(u))]
        if no_parents:
            msgs.append(('info', 'No parents recorded for ' + ', '.join(no_parents) +
                         ' - relatedness to the others cannot be checked.'))
        return msgs

    def suggest(self, line, limit=15):
        partners = self.partner_lines(line)
        males = self.candidate_units('M', partners)
        females = self.candidate_units('F', partners)
        def key(u):
            if u.startswith('mouse:'):
                mm = self.mice[u[6:]]
                return (mm['cage'], format_gt(self.mouse_gt(mm['id'])))
            return (u,)
        dedup = {}
        for u in males + females:
            dedup.setdefault(key(u), u)
        males = [u for u in males if dedup[key(u)] == u]
        females = [u for u in females if dedup[key(u)] == u]
        out = []
        for m in males:
            for f in females:
                k = self.kinship(m, f)
                mi, fi = self.unit_info(m), self.unit_info(f)
                # at least one partner should be from the line itself
                if line not in (mi['line'], fi['line']):
                    continue
                left = min(x for x in (self.days_left(mi['birth']), self.days_left(fi['birth'])) if x is not None)
                out.append(dict(m=m, f=f, kin=k, left=left, rel=self.relationship(m, f),
                                pups=self.pups_text(m, f)))
        out.sort(key=lambda x: (round(x['kin'], 4), -x['left']))
        return out[:limit]

    def due(self, horizon_days):
        out = []
        for p in self.pairs:
            if p['status'] != 'Active':
                continue
            a = self.pair_oldest(p)
            if a is None or a + horizon_days >= self.settings['MaxAge']:
                out.append(p)
        out.sort(key=lambda p: self.pair_replace_by(p) or dt.date.min)
        return out

    def line_status(self):
        rows = []
        for line, info in self.lines.items():
            active = [p for p in self.pairs if p['line'] == line and p['status'] == 'Active']
            over = [p for p in active if self.pair_flag(p) == 'OVER AGE']
            soon = [p for p in active if self.pair_flag(p) == 'Replace soon']
            nm = len(self.candidate_units('M', self.partner_lines(line)))
            nf = len(self.candidate_units('F', self.partner_lines(line)))
            if info['keep'] and not active:
                st = 'NO ACTIVE PAIR'
            elif over:
                st = 'Pair over age'
            elif soon:
                st = 'Plan replacement'
            else:
                st = ''
            rows.append(dict(line=line, keep=info['keep'], active=len(active), over=len(over), soon=len(soon),
                             cand_m=nm, cand_f=nf, status=st))
        return rows

    def next_id(self):
        nums = [int(b[1:]) for b in self.breeders if b[:1] == 'B' and b[1:].isdigit()]
        return f'B{(max(nums) + 1) if nums else 1:03d}'

    def empty_cages(self):
        return [c['cage'] for c in self.cages if not c['line'] and not c['sex'] and not c['count']]

    @staticmethod
    def card_name(cage):
        """'IIA4e' -> 'II/A/4/e' (as written on the cage cards)."""
        m = re.match(r'^(I{1,3}|IV|V)([A-Z])(\d+)([a-z])$', cage or '')
        return '/'.join(m.groups()) if m else cage

    @staticmethod
    def rack(cage):
        m = re.match(r'^(I{1,3}|IV|V)([A-Z])', cage or '')
        return f'{m.group(1)}/{m.group(2)}' if m else '?'

    def all_mice(self):
        """One row per cage (sheet order): stock cages and breeding cages, with ages and status."""
        active = {p['cage']: p for p in self.pairs if p['status'] == 'Active'}
        s = self.settings
        out = []
        for c in self.cages:
            p = active.get(c['cage'])
            row = dict(cage=c['cage'], card=self.card_name(c['cage']), rack=self.rack(c['cage']), line=c['line'],
                       sex=c['sex'], count=c['count'] or 0, birth=c['birth'], sire=c['sire'], dam=c['dam'],
                       note=c['note'], pair=None, males=0, females=0, mice=self.mice_in(c['cage']),
                       gt=self.cage_gt_text(c['cage']) or self.gt_display(c['gt'], c['sire'], c['dam'], c['line']))
            if row['mice'] and not p:
                row['count'] = len(row['mice'])
            if p or c['sex'] == 'TP':
                members = [self.breeders[x] for x in self.pair_members(p) if x in self.breeders] if p else []
                row['pair'] = p
                row['males'] = sum(1 for b in members if b['sex'] == 'M')
                row['females'] = sum(1 for b in members if b['sex'] == 'F')
                row['count'] = len(members) or row['count']
                row['gt'] = ' | '.join(f"{b['id']} " + (('wt' if b['line'] == 'C57Bl6J' else '')
                                                         or format_gt(b.get('gt') or {}) or '?') for b in members)
                row['line'] = (p['line'] if p else '') or c['line']
                row['age'] = self.pair_oldest(p) if p else None
                row['animals'] = ', '.join(
                    f"{b['id']} {'♂' if b['sex'] == 'M' else '♀'} {fmt_age(self.age(b['birth']))}" for b in members)
                row['status'] = self.pair_flag(p) if p else 'Pair not on Pairs sheet'
                row['tag'] = {'OK': 'pair'}.get(row['status'], row['status'])
                if p:
                    fs = [f for f in (p['f1'], p['f2']) if f]
                    k = max([self.kinship(p['m'], f) for f in fs] or [0]) if p['m'] else 0
                    if k > 0:
                        row['status'] += f' | related F={k:.3f}'
            elif not c['line'] and not row['count']:
                row.update(age=None, animals='', status='empty', tag='empty')
            else:
                a = self.age(c['birth'])
                row['age'] = a
                row['males'] = row['count'] if c['sex'] == 'M' else 0
                row['females'] = row['count'] if c['sex'] == 'F' else 0
                row['animals'] = f"{row['count']} {'♂' if c['sex'] == 'M' else '♀' if c['sex'] == 'F' else '?'}"
                if row['mice']:
                    row['animals'] += '  (' + ', '.join(m['mark'] for m in row['mice']) + ')'
                elif row['count']:
                    row['animals'] += '  (no mouse IDs)'
                if c['note'] and 'terminat' in c['note'].lower():
                    row['status'], row['tag'] = 'Terminate (see note)', 'Terminate'
                elif a is None:
                    row['status'], row['tag'] = 'NO DATE', 'NO DATE'
                elif a < s['MinAge']:
                    row['status'], row['tag'] = 'Too young to breed', 'Too young'
                elif a <= s['StartMax']:
                    row['status'], row['tag'] = 'Breeder candidate', 'Candidate'
                else:
                    row['status'], row['tag'] = 'Stock (too old to start breeding)', 'Stock'
            out.append(row)
        return out

    def save_cage(self, name, line, sex, add_n, birth, sire, dam, note, gt=None):
        """Edit a stock cage: line/sex/birth/parents apply to every mouse in it; add_n new mice are created
        (gt = genotype for the new mice)."""
        wb = self._open_for_write()
        existing = self.mice_in(name)
        self._set_cage(wb['Cages'], name, line=line, sex=sex, birth=birth, sire=sire, dam=dam, note=note)
        for m in existing:
            m.update(line=line, sex=sex, birth=_d(birth) if birth else None, sire=sire, dam=dam)
            self._write_mouse(wb['Mice'], m)
        if add_n:
            self._add_mice(wb, name, add_n, line, sex, _d(birth) if birth else None, sire, dam, gt=gt or {})
        self._sync_count(wb, name)
        self._save(wb)

    def clear_cage(self, name):
        wb = self._open_for_write()
        self._set_cage(wb['Cages'], name, clear=True)
        self._save(wb)

    # ---------------- procedures + termination history ----------------
    def active_pair(self, cage):
        return next((p for p in self.pairs if p['cage'] == cage and p['status'] == 'Active'), None)

    def people(self):
        return sorted({x['person'] for x in self.procedures + self.terminated if x['person']})

    def procedure_types(self):
        extra = sorted({p['proc'] for p in self.procedures if p['proc'] and p['proc'] not in PROCEDURE_TYPES})
        return PROCEDURE_TYPES + extra

    @staticmethod
    def _ids(text):
        return [x.strip() for x in (text or '').split(',') if x.strip()]

    def procedures_for_mouse(self, mid):
        m = self.mice[mid]
        out = [p for p in self.procedures if mid in self._ids(p['animal'])
               or (not p['animal'] and p['cage'] == m['cage'] and p['birth'] == m['birth'] and p['sex'] != 'pair')]
        if m['breeder']:
            out += [p for p in self.procedures_for(animal=m['breeder']) if p not in out]
        return sorted(out, key=lambda p: p['date'] or dt.date.min)

    def procedures_for(self, cage=None, birth=None, animal=None):
        """Procedures on one breeder (its own + those logged on its whole breeding cage), or on one stock cohort
        (same cage and same birth date, so a re-used cage does not inherit old records)."""
        out = []
        for p in self.procedures:
            if animal:
                b = self.breeders.get(animal, {})
                if animal in self._ids(p['animal']) or (not p['animal'] and p['sex'] == 'pair'
                                                        and p['cage'] == b.get('cage')):
                    out.append(p)
            elif cage and p['cage'] == cage and not p['animal'] and p['birth'] == birth:
                out.append(p)
        if animal:                         # the breeder's life as a stock mouse before it was taken
            for m in self.mice.values():
                if m['breeder'] == animal:
                    out += [p for p in self.procedures if m['id'] in self._ids(p['animal']) and p not in out]
        return sorted(out, key=lambda p: p['date'] or dt.date.min)

    @staticmethod
    def proc_summary(procs):
        parts = []
        for p in procs:
            s = f"{fmt_date(p['date'])} {p['proc']}"
            if p['details']:
                s += f" ({p['details']})"
            if p['person']:
                s += f" [{p['person']}]"
            parts.append(s)
        return '; '.join(parts)

    @staticmethod
    def _ensure_history_sheets(wb):
        """Older workbooks: add the Procedures / Terminated sheets and the 'Terminated' breeder status."""
        from openpyxl.styles import Font, PatternFill, Alignment
        for name, head, widths in (('Procedures', PROC_HEAD, PROC_W), ('Terminated', TERM_HEAD, TERM_W)):
            if name in wb.sheetnames:
                continue
            idx = wb.sheetnames.index('Litters') + 1 if 'Litters' in wb.sheetnames else len(wb.sheetnames)
            if name == 'Terminated' and 'Procedures' in wb.sheetnames:
                idx = wb.sheetnames.index('Procedures') + 1
            ws = wb.create_sheet(name, index=idx)
            for c, (h, w) in enumerate(zip(head, widths), 1):
                cell = ws.cell(1, c, h)
                cell.font = Font(name='Arial', bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='1F3A5F' if name == 'Procedures' else '7F3F3F')
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
            ws.row_dimensions[1].height = 30
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = f'A1:{openpyxl.utils.get_column_letter(len(head))}{N}'
            ws.sheet_properties.tabColor = '1F3A5F' if name == 'Procedures' else '7F3F3F'
        # genotype columns (Breeders N, Cages N) and the loci per line (Lines M)
        for name, col, head, width, color in (('Breeders', 14, 'Genotype', 34, '1F3A5F'),
                                              ('Cages', 14, 'Genotype', 34, '1F3A5F'),
                                              ('Lines', 13, 'Loci tracked (; separated)', 30, '1F3A5F'),
                                              ('Lines', 14, 'Made of (lines, ; separated)', 30, '1F3A5F')):
            ws = wb[name]
            if not ws.cell(1, col).value:
                cell = ws.cell(1, col, head)
                cell.font = Font(name='Arial', bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor=color)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
                if name == 'Lines':
                    defaults = DEFAULT_LOCI if col == 13 else DEFAULT_PARTS
                    for r in range(2, ws.max_row + 1):
                        ln = _s(ws.cell(r, 1).value)
                        if ln and defaults.get(ln):
                            ws.cell(r, col, defaults[ln]).font = Font(name='Arial', size=10)
                elif name == 'Cages':
                    for r in range(2, ws.max_row + 1):
                        ln = _s(ws.cell(r, 2).value)
                        if ln == 'CX3CR1KO (-/-)' and _s(ws.cell(r, 4).value) in ('M', 'F'):
                            ws.cell(r, col, 'Cx3cr1-KO: hom').font = Font(name='Arial', size=10)
                else:
                    for r in range(2, ws.max_row + 1):
                        if _s(ws.cell(r, 2).value) == 'CX3CR1KO (-/-)':
                            ws.cell(r, col, 'Cx3cr1-KO: hom').font = Font(name='Arial', size=10)
                ws.auto_filter.ref = None
                last = openpyxl.utils.get_column_letter(col)
                ws.auto_filter.ref = f'A1:{last}{N}'
        # named cells (MaxAge...) -> plain references, so the formulas also work in Google Sheets
        names = {'MaxAge': 'Settings!$B$2', 'WarnDays': 'Settings!$B$3', 'MinAge': 'Settings!$B$4',
                 'StartMax': 'Settings!$B$5'}
        for sheet in ('Breeders', 'Pairs', 'Cages'):
            if sheet not in wb.sheetnames:
                continue
            for row in wb[sheet].iter_rows():
                for cell in row:
                    v = cell.value
                    if isinstance(v, str) and v.startswith('=') and any(k in v for k in names):
                        for k, ref in names.items():
                            v = re.sub(r'(?<![\w!$])' + k + r'(?![\w])', ref.replace('\\', ''), v)
                        cell.value = v
        br = wb['Breeders']
        if 'Terminated' not in str(br['L2'].value or 'Terminated'):
            for r in range(2, br.max_row + 1):
                if br.cell(r, 1).value:
                    br.cell(r, 12).value = breeder_formulas(r)[2]
        for dv in br.data_validations.dataValidation:
            if dv.formula1 and 'Breeding,Reserve' in dv.formula1 and 'Terminated' not in dv.formula1:
                dv.formula1 = '"Breeding,Reserve,Retired,Terminated"'

    @staticmethod
    def _append(ws, values, date_cols, ncols):
        from openpyxl.styles import Font, Alignment
        r = Colony._first_free_row(ws)
        for c, v in enumerate(values, 1):
            cell = ws.cell(r, c, v if v != '' else None)
            cell.font = Font(name='Arial', size=10)
            if c in date_cols:
                cell.number_format = DATE_FMT
        ws.cell(r, ncols).alignment = Alignment(wrap_text=False)
        return r

    def log_procedure(self, date, proc, details='', person='', notes='', cage=None, animal=None, n=None,
                      mice=None):
        """Log something done to chosen mice, a whole breeding cage, or one breeder."""
        wb = self._open_for_write()
        if mice:
            m0 = self.mice[mice[0]]
            self._append(wb['Procedures'], [date, m0['cage'], ', '.join(mice), m0['line'], m0['sex'], len(mice),
                                            m0['birth'], proc, details, person, notes], (1, 7), 11)
            self._save(wb)
            return
        if animal:
            b = self.breeders[animal]
            vals = [date, b['cage'], animal, b['line'], b['sex'], 1, b['birth']]
        else:
            c = self.cage(cage)
            pair = self.active_pair(cage)
            if pair:
                vals = [date, cage, '', pair['line'], 'pair', len(self.pair_members(pair)), None]
            else:
                vals = [date, cage, '', c['line'], c['sex'], n or c['count'], c['birth']]
        self._append(wb['Procedures'], vals + [proc, details, person, notes], (1, 7), 11)
        self._save(wb)

    def _term_row(self, wb, date, cage, animal, line, sex, n, birth, sire, dam, procs, outcome, reason, project,
                  person, notes, gt=None):
        age = (date - birth).days if (birth and date) else None
        self._append(wb['Terminated'], [date, cage, animal, line, sex, n, birth, age, outcome, reason, project,
                                        person, sire, dam, format_gt(gt or {}), self.proc_summary(procs), notes],
                     (1, 7), 17)

    def terminate_mice(self, ids, date, outcome, reason='', project='', person='', notes=''):
        """One history row per mouse; the mouse (and its breeder row, if it was one) becomes Terminated."""
        wb = self._open_for_write()
        cages = set()
        for mid in ids:
            m = self.mice[mid]
            gt = {k: v.rstrip('*') for k, v in self.mouse_gt(mid).items() if v != '?'}
            label = f"{mid} ({m['mark']})" + (f" / {m['breeder']}" if m['breeder'] else '')
            self._term_row(wb, date, m['cage'], label, m['line'], m['sex'], 1, m['birth'], m['sire'], m['dam'],
                           self.procedures_for_mouse(mid), outcome, reason, project, person, notes, gt=gt)
            m['status'], m['since'] = 'Terminated', date
            self._write_mouse(wb['Mice'], m)
            if m['breeder'] in self.breeders:
                b = self.breeders[m['breeder']]
                b['status'] = 'Terminated'
                b['notes'] = (b['notes'] + '; ' if b['notes'] else '') + f'Terminated {date} ({outcome})'
                self._write_breeder(wb['Breeders'], b)
            cages.add(m['cage'])
        for c in cages:
            self._sync_count(wb, c)
        self._save(wb)

    def terminate_cage(self, cage, n, date, outcome, reason='', project='', person='', notes=''):
        """Remove n animals from a stock cage into the termination history."""
        mice = self.mice_in(cage)
        if mice:
            if n < 1 or n > len(mice):
                raise ValueError(f'{self.card_name(cage)} has {len(mice)} mice.')
            return self.terminate_mice([m['id'] for m in mice[:n]], date, outcome, reason, project, person, notes)
        c = self.cage(cage)
        if self.active_pair(cage):
            raise ValueError(f'{self.card_name(cage)} is a breeding cage - terminate its animals instead.')
        if not c or not c['count']:
            raise ValueError(f'{self.card_name(cage)} has no animals.')
        if n < 1 or n > c['count']:
            raise ValueError(f'{self.card_name(cage)} has {c["count"]} animals.')
        wb = self._open_for_write()
        self._ensure_history_sheets(wb)
        # former breeders living in a stock cage (their IDs are in the cage note)
        ids = [x for x in re.findall(r'B\d{3,}', c['note'] or '') if x in self.breeders
               and self.breeders[x]['status'] != 'Terminated'][:n] if n == c['count'] else []
        procs = self.procedures_for(cage=cage, birth=c['birth'])
        for x in ids:
            procs += [p for p in self.procedures_for(animal=x) if p not in procs]
        procs.sort(key=lambda p: p['date'] or dt.date.min)
        self._term_row(wb, date, cage, ', '.join(ids), c['line'], c['sex'], n, c['birth'], c['sire'], c['dam'],
                       procs, outcome, reason, project, person, notes, gt=c['gt'])
        for x in ids:
            b = self.breeders[x]
            b['status'] = 'Terminated'
            b['notes'] = (b['notes'] + '; ' if b['notes'] else '') + f'Terminated {date} ({outcome})'
            self._write_breeder(wb['Breeders'], b)
        if n == c['count']:
            self._set_cage(wb['Cages'], cage, clear=True)
        else:
            self._set_cage(wb['Cages'], cage, count=c['count'] - n)
        self._save(wb)

    def terminate_breeders(self, ids, date, outcome, reason='', project='', person='', notes=''):
        """Terminate breeders (one history row each). A pair that loses its male or all females ends; animals
        left alive stay in the cage as a stock row."""
        wb = self._open_for_write()
        self._ensure_history_sheets(wb)
        ids = list(ids)
        for bid in ids:
            b = self.breeders[bid]
            self._term_row(wb, date, b['cage'], bid, b['line'], b['sex'], 1, b['birth'], b['sire'], b['dam'],
                           self.procedures_for(animal=bid), outcome, reason, project, person, notes, gt=b.get('gt'))
        for p in [p for p in self.pairs if p['status'] == 'Active' and set(self.pair_members(p)) & set(ids)]:
            left = [x for x in self.pair_members(p) if x not in ids]
            males = [x for x in left if self.breeders.get(x, {}).get('sex') == 'M']
            fems = [x for x in left if self.breeders.get(x, {}).get('sex') == 'F']
            ws = wb['Pairs']
            if males and fems:                       # e.g. one female of a trio
                ws.cell(p['row'], 3).value = males[0]
                ws.cell(p['row'], 4).value = fems[0]
                ws.cell(p['row'], 5).value = fems[1] if len(fems) > 1 else None
                self._set_cage(wb['Cages'], p['cage'], count=len(left))
            else:
                ws.cell(p['row'], 7).value = 'Ended'
                ws.cell(p['row'], 8).value = ((p['notes'] + '; ') if p['notes'] else '') + f'Ended {date}'
                if left:
                    self._pair_leftover_to_stock(wb, p, left, date)
                else:
                    self._set_cage(wb['Cages'], p['cage'], clear=True)
        for bid in ids:
            b = self.breeders[bid]
            b['status'] = 'Terminated'
            b['notes'] = (b['notes'] + '; ' if b['notes'] else '') + f'Terminated {date} ({outcome})'
            self._write_breeder(wb['Breeders'], b)
            for m in self.mice.values():
                if m['breeder'] == bid and m['status'] != 'Terminated':
                    m['status'], m['since'] = 'Terminated', date
                    self._write_mouse(wb['Mice'], m)
        self._save(wb)

    def _pair_leftover_to_stock(self, wb, p, left, date):
        bs = [self.breeders[x] for x in left]
        for b in bs:
            b['status'] = 'Retired'
            self._write_breeder(wb['Breeders'], b)
        one = bs[0] if len(bs) == 1 else None
        births = [b['birth'] for b in bs if b['birth']]
        self._set_cage(wb['Cages'], p['cage'], clear=True)
        ws = wb['Pairs']
        ws.cell(p['row'], 7).value = 'Ended'
        p['status'] = 'Ended'
        self._set_cage(wb['Cages'], p['cage'], line=bs[0]['line'] or p['line'], sex=bs[0]['sex'], count=len(bs),
                       birth=min(births) if births else '', sire=one['sire'] if one else '',
                       dam=one['dam'] if one else '',
                       note=f"Former breeders {', '.join(left)} (pair ended {date})",
                       gt=one['gt'] if one else '')
        for b in bs:
            linked = next((m for m in self.mice.values() if m['breeder'] == b['id']), None)
            if linked:
                linked.update(cage=p['cage'], status='Alive', since=date,
                              mark=self.free_marks(p['cage'], 1)[0])
                self._write_mouse(wb['Mice'], linked)
            else:
                self._add_mice(wb, p['cage'], 1, b['line'] or p['line'], b['sex'], b['birth'], b['sire'],
                               b['dam'], gt=b.get('gt'), breeders=[b['id']])
        self._sync_count(wb, p['cage'])

    def end_pair_keep(self, pair_cage, date=None):
        """End a pair without terminating: the animals stay in the cage as a retired stock row (same sex) or
        the cage keeps them as 'Retired pair'."""
        p = self.active_pair(pair_cage)
        if not p:
            return
        wb = self._open_for_write()
        date = date or self.today
        members = self.pair_members(p)
        wb['Pairs'].cell(p['row'], 7).value = 'Ended'
        sexes = {self.breeders[x]['sex'] for x in members if x in self.breeders}
        if len(sexes) == 1:
            self._pair_leftover_to_stock(wb, p, members, date)
        else:
            for x in members:
                b = self.breeders[x]
                b['status'] = 'Retired'
                self._write_breeder(wb['Breeders'], b)
            self._set_cage(wb['Cages'], pair_cage, clear=True)
        self._save(wb)

    # ---------------- writing ----------------
    def _open_for_write(self):
        # shared sheet: work on the local copy; it is uploaded in the background (and works offline)
        wb = openpyxl.load_workbook(self.path)
        self._ensure_history_sheets(wb)
        self._ensure_mice_sheet(wb)
        return wb

    def _ensure_mice_sheet(self, wb):
        from openpyxl.styles import Font, PatternFill, Alignment
        if 'Mice' in wb.sheetnames:
            ws = wb['Mice']
            if not ws.cell(1, 14).value:
                cell = ws.cell(1, 14, MICE_HEAD[13])
                cell.font = Font(name='Arial', bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='1F3A5F')
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                ws.column_dimensions['N'].width = MICE_W[13]
                ws.auto_filter.ref = f'A1:N{N * 5}'
            return
        idx = wb.sheetnames.index('Cages') + 1 if 'Cages' in wb.sheetnames else len(wb.sheetnames)
        ws = wb.create_sheet('Mice', index=idx)
        for c, (h, w) in enumerate(zip(MICE_HEAD, MICE_W), 1):
            cell = ws.cell(1, c, h)
            cell.font = Font(name='Arial', bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F3A5F')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = 'B2'
        ws.auto_filter.ref = f'A1:M{N * 5}'
        ws.sheet_properties.tabColor = '1F3A5F'
        self.mice = {}
        for c in self.cages:
            if c['sex'] not in ('M', 'F') or not c['count'] or self.active_pair(c['cage']):
                continue
            former = [x for x in re.findall(r'B\d{3,}', c['note'] or '') if x in self.breeders]
            for i in range(c['count']):
                bid = former[i] if i < len(former) else ''
                b = self.breeders.get(bid, {})
                m = dict(id=self.next_mouse_id(), cage=c['cage'], mark=EAR_MARKS[i] if i < len(EAR_MARKS)
                         else f'#{i + 1}', line=c['line'], sex=c['sex'], birth=b.get('birth') or c['birth'],
                         sire=b.get('sire', c['sire']) if bid else c['sire'],
                         dam=b.get('dam', c['dam']) if bid else c['dam'],
                         gt=dict(b.get('gt') or c['gt']), status='Alive', since=None, breeder=bid,
                         notes='Ear mark assigned automatically - check the ears')
                self._write_mouse(ws, m)

    # ---------------- cage-card import (JSON made from photos) ----------------
    @staticmethod
    def norm_cage(text):
        """'II/A/4/e', 'IIA4e', 'II A 4 e' -> 'IIA4e'"""
        t = re.sub(r'[\s/\\|.-]', '', str(text or ''))
        m = re.match(r'^(I{1,3}|IV|V|1{1,3})([A-Za-z])(\d+)([A-Za-z])$', t)
        if not m:
            return t
        rack = {'1': 'I', '11': 'II', '111': 'III'}.get(m.group(1), m.group(1).upper())
        return f'{rack}{m.group(2).upper()}{m.group(3)}{m.group(4).lower()}'

    @staticmethod
    def _card_date(v):
        if not v:
            return None
        v = str(v).strip().replace('.', '-').strip('-')
        for fmt in ('%Y-%m-%d', '%y-%m-%d'):
            try:
                return dt.datetime.strptime(v, fmt).date()
            except ValueError:
                pass
        return None

    def compare_cards(self, data):
        """Compare imported cage cards with the workbook. Returns a list of change dicts:
        {cage, what, app, card, kind, apply, value}. kind 'info' rows are only shown, never applied."""
        out = []
        def add(cage, what, app, card, kind, apply=True, value=None):
            out.append(dict(cage=cage, what=what, app='' if app is None else str(app),
                            card='' if card is None else str(card), kind=kind, apply=apply and kind != 'info',
                            value=value))
        for card in data.get('cards', []):
            cage = self.norm_cage(card.get('cage'))
            c = self.cage(cage)
            if not c:
                add(cage, 'Unknown cage', '', card.get('cage'), 'info')
                continue
            if card.get('unreadable'):
                add(cage, 'Card not readable', '', card.get('notes', ''), 'info')
                continue
            ctype = (card.get('type') or ('pair' if card.get('males') or card.get('females') else
                                          'empty' if not card.get('count') else 'stock')).lower()
            pair = self.active_pair(cage)
            mice = self.mice_in(cage)
            unsure = set(card.get('uncertain') or [])
            if ctype == 'empty':
                if mice or pair:
                    add(cage, 'Card says empty', f"{len(mice) or len(self.pair_members(pair))} animals", 'empty',
                        'info')
                continue
            if ctype == 'pair':
                if not pair:
                    add(cage, 'Card is a breeding pair, app has no pair here', c['line'] or 'empty', 'pair', 'info')
                    continue
                members = [self.breeders[x] for x in self.pair_members(pair) if x in self.breeders]
                for sex, key in (('M', 'males'), ('F', 'females')):
                    app_b = sorted(fmt_date(b['birth']) for b in members if b['sex'] == sex)
                    card_b = sorted(fmt_date(self._card_date(x.get('birth'))) for x in card.get(key) or [])
                    if card.get(key) is not None and app_b != card_b:
                        add(cage, f"{'Male' if sex == 'M' else 'Female'} birth date(s)", ', '.join(app_b),
                            ', '.join(card_b), 'info')
                continue
            if pair:
                add(cage, 'Card is a stock cage, app has a breeding pair', pair['line'], card.get('line'), 'info')
                continue
            line = card.get('line') or ''
            line_ok = line in self.lines
            sex = (card.get('sex') or '').upper()[:1]
            birth = self._card_date(card.get('birth'))
            n = card.get('count')
            if not mice:
                if not line_ok:
                    add(cage, 'New cage - line not recognised', '', f"{line} ({card.get('line_as_written', '')})",
                        'info')
                    continue
                sire, dam = card.get('sire') or '', card.get('dam') or ''
                p0 = self.active_pair(self.norm_cage(card.get('origin'))) if card.get('origin') else None
                if p0 and not sire:
                    fs = [f for f in (p0['f1'], p0['f2']) if f]
                    sire, dam = p0['m'], (fs[0] if len(fs) == 1 else '')
                add(cage, f'New cage: {n} {sex}', 'empty',
                    f'{line}, born {fmt_date(birth)}' + (f', parents {sire} x {dam or "?"}' if sire else ''),
                    'new_cage', value=dict(line=line, sex=sex, n=int(n or 0), birth=birth, sire=sire, dam=dam,
                                           note=card.get('notes') or ''))
                continue
            if line and line != c['line']:
                if line_ok:
                    add(cage, 'Line', c['line'], line, 'line', apply='line' not in unsure, value=line)
                else:
                    add(cage, 'Line (not a known line)', c['line'], f"{line} ({card.get('line_as_written', '')})",
                        'info')
            if sex and sex != c['sex']:
                add(cage, 'Sex', c['sex'], sex, 'sex', apply='sex' not in unsure, value=sex)
            if birth and birth != c['birth']:
                add(cage, 'Birth date', fmt_date(c['birth']), fmt_date(birth), 'birth', apply='birth' not in unsure,
                    value=birth)
            elif card.get('birth') and not birth:
                add(cage, 'Birth date unreadable', fmt_date(c['birth']), card.get('birth'), 'info')
            if isinstance(n, int) and n != len(mice):
                if n > len(mice):
                    k_ = n - len(mice)
                    add(cage, f"Count: add {k_} mouse" if k_ == 1 else f"Count: add {k_} mice", len(mice), n,
                        'add_mice',
                        apply='count' not in unsure, value=n - len(mice))
                else:
                    add(cage, 'Count lower on card - use Terminate / remove for the missing mice', len(mice), n,
                        'info')
            for k, label in (('sire', 'Sire'), ('dam', 'Dam')):
                v = card.get(k)
                if v and v != c[k]:
                    if v in self.breeders:
                        add(cage, label, c[k], v, k, apply=not c[k], value=v)
                    else:
                        add(cage, f'{label} (not a breeder ID)', c[k], v, 'info')
            if card.get('origin'):
                o = self.norm_cage(card['origin'])
                p0 = self.active_pair(o)
                if p0 and not c['sire'] and not card.get('sire'):
                    fs = [f for f in (p0['f1'], p0['f2']) if f]
                    add(cage, f'Parents from origin {self.card_name(o)}', '',
                        f"{p0['m']} x {fs[0] if len(fs) == 1 else '?'}", 'parents', apply=False,
                        value=(p0['m'], fs[0] if len(fs) == 1 else ''))
            note = (card.get('notes') or '').strip()
            if note and note.lower() not in (c['note'] or '').lower():
                add(cage, 'Note', c['note'], note, 'note', apply=False, value=note)
            for cm in card.get('mice') or []:
                mm = next((m for m in mice if m['mark'] == cm.get('mark')), None)
                g = {k: str(v).lower() for k, v in (cm.get('genotype') or {}).items()}
                if mm and g and any(mm['gt'].get(k) != v for k, v in g.items()):
                    add(cage, f"Genotype {mm['id']} (ear {mm['mark']})", format_gt(mm['gt']) or '?', format_gt(g),
                        'mouse_gt', value=(mm['id'], g))
                elif not mm and cm.get('mark'):
                    add(cage, f"No mouse with ear {cm.get('mark')}", '', format_gt(g), 'info')
        return out

    def apply_card_changes(self, changes, checked_on=None):
        wb = self._open_for_write()
        touched = set()
        for ch in changes:
            if not ch['apply'] or ch['kind'] == 'info':
                continue
            cage, v = ch['cage'], ch['value']
            c = self.cage(cage)
            k = ch['kind']
            if k == 'new_cage':
                self._set_cage(wb['Cages'], cage, line=v['line'], sex=v['sex'], birth=v['birth'] or '',
                               sire=v['sire'], dam=v['dam'], note=v['note'])
                c.update(line=v['line'], sex=v['sex'])
                self._add_mice(wb, cage, v['n'], v['line'], v['sex'], v['birth'], v['sire'], v['dam'])
            elif k == 'add_mice':
                self._add_mice(wb, cage, v, c['line'], c['sex'], c['birth'], c['sire'], c['dam'])
            elif k in ('line', 'sex', 'birth', 'sire', 'dam', 'parents'):
                upd = {'parents': {}}.get(k, {k: v})
                if k == 'parents':
                    upd = dict(sire=v[0], dam=v[1])
                self._set_cage(wb['Cages'], cage, **{kk: (vv if vv is not None else '') for kk, vv in upd.items()})
                for m in self.mice_in(cage):
                    m.update(upd)
                    self._write_mouse(wb['Mice'], m)
                c.update(upd)
            elif k == 'note':
                self._set_cage(wb['Cages'], cage, note=((c['note'] + '; ') if c['note'] else '') + v)
            elif k == 'mouse_gt':
                mid, g = v
                m = self.mice[mid]
                m['gt'].update({kk: vv for kk, vv in g.items() if vv != '?'})
                self._write_mouse(wb['Mice'], m)
            touched.add(cage)
        if checked_on and 'Settings' in wb.sheetnames:
            wb['Settings']['B7'] = dt.datetime.combine(checked_on, dt.time())
        self._save(wb)
        return len(touched)

    # ---------------- individual mice ----------------
    def next_mouse_id(self):
        nums = [int(x[1:]) for x in self.mice if x[:1] == 'M' and x[1:].isdigit()]
        return f'M{(max(nums) + 1) if nums else 1:04d}'

    @staticmethod
    def mark_key(mark):
        return EAR_MARKS.index(mark) if mark in EAR_MARKS else len(EAR_MARKS)

    def mice_in(self, cage):
        return sorted((m for m in self.mice.values() if m['cage'] == cage and m['status'] == 'Alive'),
                      key=lambda m: (self.mark_key(m['mark']), m['id']))

    def free_marks(self, cage, n):
        used = {m['mark'] for m in self.mice_in(cage)}
        free = [x for x in EAR_MARKS if x not in used]
        k = len(EAR_MARKS)
        while len(free) < n:
            k += 1
            free.append(f'#{k}')
        return free[:n]

    def _write_mouse(self, ws, m):
        r = m.get('row') or self._first_free_row(ws)
        from openpyxl.styles import Font
        vals = [m['id'], m['cage'], m['mark'], m['line'], m['sex'], m['birth'], m['sire'] or None,
                m['dam'] or None, format_gt(m['gt']) or None, m['status'], m['since'], m['breeder'] or None,
                m['notes'] or None, 'Y' if m.get('todo') else None]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c)
            cell.value = v
            cell.font = Font(name='Arial', size=10)
            if c in (6, 11):
                cell.number_format = DATE_FMT
        m['row'] = r
        self.mice[m['id']] = m

    def _add_mice(self, wb, cage, n, line, sex, birth, sire='', dam='', gt=None, notes='', breeders=(), flag=None):
        """Create n mice in a cage (cage row filled if it was empty). Returns the new mouse IDs.
        flag: True/False = mark 'needs genotyping'; None = mark only if the genotype is not certain."""
        c = self.cage(cage)
        if not self.mice_in(cage) and (not c or not c['line']):
            self._set_cage(wb['Cages'], cage, line=line, sex=sex, birth=birth or '', sire=sire or '',
                           dam=dam or '', note=notes or '')
            if c:
                c.update(line=line, sex=sex)
        ids = []
        for i, mark in enumerate(self.free_marks(cage, n)):
            bid = breeders[i] if i < len(breeders) else ''
            b = self.breeders.get(bid, {})
            m = dict(id=self.next_mouse_id(), cage=cage, mark=mark, line=line, sex=sex,
                     birth=b.get('birth') or birth, sire=b.get('sire', sire) if bid else sire,
                     dam=b.get('dam', dam) if bid else dam, gt=dict(b.get('gt') or gt or {}), status='Alive',
                     since=self.today, breeder=bid, notes='', todo=False)
            self.mice[m['id']] = m
            m['todo'] = (self.gt_unknown(m['id']) if flag is None else bool(flag)) and not bid
            self._write_mouse(wb['Mice'], m)
            ids.append(m['id'])
        self._sync_count(wb, cage)
        return ids

    def _sync_count(self, wb, cage):
        if self.active_pair(cage):
            return
        n = len(self.mice_in(cage))
        if n == 0:
            self._set_cage(wb['Cages'], cage, clear=True)
        else:
            self._set_cage(wb['Cages'], cage, count=n)
        c = self.cage(cage)
        if c:
            c['count'] = n

    def gt_unknown(self, mid):
        """True if any allele of the mouse's line is neither genotyped nor certain from the parents."""
        return any(v == '?' for v in self.mouse_gt(mid).values())

    def needs_genotyping(self, mid):
        m = self.mice[mid]
        return bool(m.get('todo')) and m['status'] == 'Alive'

    def mouse_gt(self, mid):
        """Effective genotype per locus (stored, or certain from the parents)."""
        m = self.mice[mid]
        if m['line'] == 'C57Bl6J':
            return {}
        out = {}
        for lc in self.loci_for(m['line']):
            v = self.gt_of('mouse:' + mid, lc)
            out[lc] = v + ('*' if v != '?' and not m['gt'].get(lc) else '')
        return out

    def mouse_gt_text(self, mid):
        m = self.mice[mid]
        if m['line'] == 'C57Bl6J':
            return 'wt'
        g = self.mouse_gt(mid)
        return '; '.join(f'{k}: {v}' for k, v in g.items() if v != '?') or '?'

    def cage_gt_text(self, cage):
        mice = self.mice_in(cage)
        if not mice:
            return ''
        if mice[0]['line'] == 'C57Bl6J':
            return 'wt'
        per = {}
        for m in mice:
            for k, v in self.mouse_gt(m['id']).items():
                per.setdefault(k, []).append(v.rstrip('*'))
        parts = []
        for k, vals in per.items():
            known = [v for v in vals if v != '?']
            if not known:
                continue
            if len(set(known)) == 1 and len(known) == len(vals):
                parts.append(f'{k}: {known[0]}')
            else:
                cnt = {v: vals.count(v) for v in dict.fromkeys(vals)}
                parts.append(f'{k}: ' + ' '.join(f'{n} {v}' for v, n in cnt.items()))
        return '; '.join(parts)

    def save_mouse(self, mid, mark=None, sex=None, gt=None, notes=None, cage=None, todo=None):
        wb = self._open_for_write()
        m = self.mice[mid]
        old = m['cage']
        if cage and cage != old:
            if self.active_pair(cage):
                raise ValueError('Move breeders with "Create pair" instead.')
            if not self.mice_in(cage):
                self._set_cage(wb['Cages'], cage, line=m['line'], sex=m['sex'], birth=m['birth'] or '',
                               sire=m['sire'] or '', dam=m['dam'] or '', note='')
            m['cage'] = cage
            if mark is None or mark in {x['mark'] for x in self.mice_in(cage) if x['id'] != mid}:
                mark = self.free_marks(cage, 1)[0]
        if mark is not None:
            m['mark'] = mark
        if sex:
            m['sex'] = sex
        if todo is not None:
            m['todo'] = bool(todo)
        if gt is not None:
            m['gt'] = {k: v for k, v in gt.items() if v and v != '?'}
            if m.get('todo'):
                m['todo'] = self.gt_unknown(mid)
        if notes is not None:
            m['notes'] = notes
        self._write_mouse(wb['Mice'], m)
        if m['breeder'] and m['breeder'] in self.breeders and gt is not None:
            b = self.breeders[m['breeder']]
            b['gt'] = dict(m['gt'])
            self._write_breeder(wb['Breeders'], b)
        self._sync_count(wb, old)
        if cage and cage != old:
            self._sync_count(wb, cage)
        self._save(wb)

    def record_genotypes(self, results, date, person='', notes='', log=True):
        """results: {mouse_id: {locus: code}} - stores genotypes and logs one Genotyping row per cage."""
        wb = self._open_for_write()
        by_cage = {}
        for mid, g in results.items():
            m = self.mice[mid]
            m['gt'].update({k: v for k, v in g.items() if v and v != '?'})
            for k, v in g.items():
                if v == '?':
                    m['gt'].pop(k, None)
            m['todo'] = self.gt_unknown(mid)
            self._write_mouse(wb['Mice'], m)
            if m['breeder'] in self.breeders:
                b = self.breeders[m['breeder']]
                b['gt'] = dict(m['gt'])
                self._write_breeder(wb['Breeders'], b)
            by_cage.setdefault(m['cage'], []).append(mid)
        if log:
            for cage, ids in by_cage.items():
                m0 = self.mice[ids[0]]
                det = ', '.join(f"{x} ({self.mice[x]['mark']}) {format_gt(self.mice[x]['gt']) or '?'}" for x in ids)
                self._append(wb['Procedures'], [date, cage, ', '.join(ids), m0['line'], m0['sex'], len(ids),
                                                m0['birth'], 'Genotyping (tail/ear)', det, person, notes],
                             (1, 7), 11)
        self._save(wb)

    def _save(self, wb):
        if self.remote:
            # local first: save on this computer now, upload in the background (GUI) or with sync_now()
            tmp = self.path + '.part.xlsx'
            wb.save(tmp)
            os.replace(tmp, self.path)
            self.pending += 1
            self._version += 1
            self._write_state()
            self._local_only = True
            try:
                self.load()
            finally:
                self._local_only = False
            self.make_backup()
            if self.on_change:
                self.on_change()
            return
        self.make_backup()
        wb.save(self.path)
        self.load()

    @staticmethod
    def _first_free_row(ws):
        r = 2
        while ws.cell(r, 1).value not in (None, ''):
            r += 1
        return r

    @staticmethod
    def _copy_style(ws, src_row, dst_row, ncols):
        if src_row < 2 or src_row == dst_row:
            return
        for c in range(1, ncols + 1):
            s, d = ws.cell(src_row, c), ws.cell(dst_row, c)
            d.font, d.border, d.fill = s.font.copy(), s.border.copy(), s.fill.copy()
            d.number_format = s.number_format

    def _write_breeder(self, ws, b):
        r = self.breeder_rows.get(b['id']) or self._first_free_row(ws)
        self._copy_style(ws, r - 1, r, 13)
        vals = [b['id'], b['line'], b['sex'], b['birth'], b['cage'], b['status'], b['sire'] or None,
                b['dam'] or None, b['notes'] or None] + breeder_formulas(r) + [format_gt(b.get('gt') or {}) or None]
        for c, v in enumerate(vals, 1):
            ws.cell(r, c).value = v
        ws.cell(r, 4).number_format = DATE_FMT
        ws.cell(r, 13).number_format = DATE_FMT
        self.breeder_rows[b['id']] = r

    def save_breeder(self, b, wb=None):
        own = wb is None
        wb = wb or self._open_for_write()
        self._write_breeder(wb['Breeders'], b)
        if own:
            self._save(wb)

    def _cage_row(self, ws, name):
        for r in range(2, ws.max_row + 1):
            if _s(ws.cell(r, 1).value) == name:
                return r
        r = self._first_free_row(ws)
        self._copy_style(ws, r - 1, r, 13)
        ws.cell(r, 1, name)
        for c, f in enumerate(cage_formulas(r), 7):
            ws.cell(r, c, f)
        return r

    def _set_cage(self, ws, name, line=None, sex=None, count=None, birth=None, sire=None, dam=None, note=None,
                  clear=False, gt=None):
        r = self._cage_row(ws, name)
        if clear:
            for c in (2, 3, 4, 5, 6, 10, 11, 12, 13, 14):
                ws.cell(r, c).value = None
            return
        if gt is not None:
            gt = format_gt(gt) if isinstance(gt, dict) else gt
        for c, v in ((2, line), (4, sex), (5, count), (6, birth), (10, sire), (11, dam), (12, note), (14, gt)):
            if v is not None:
                ws.cell(r, c).value = v if v != '' else None
        ws.cell(r, 6).number_format = DATE_FMT

    def take_from_cage(self, wb, cage_name, sex, n=1, mouse_ids=None):
        """Turn mice of a stock cage into breeder rows (the chosen ones, else the first n).
        Returns the new breeder IDs. The mouse rows stay, linked to their breeder ID."""
        c = self.cage(cage_name)
        ws = wb['Breeders']
        pool = [m for m in self.mice_in(cage_name) if not m['breeder']]
        chosen = [self.mice[x] for x in mouse_ids] if mouse_ids else pool[:n]
        ids = []
        for m in chosen:
            bid = self.next_id()
            gt = {k: v.rstrip('*') for k, v in self.mouse_gt(m['id']).items() if v != '?'}
            b = dict(id=bid, line=m['line'], sex=m['sex'] or sex, birth=m['birth'], cage='', status='Reserve',
                     sire=m['sire'], dam=m['dam'], gt=gt,
                     notes=f"Mouse {m['id']} (ear {m['mark']}) from cage {cage_name} on {self.today}")
            self._write_breeder(ws, b)
            self.breeders[bid] = b
            m['status'], m['since'], m['breeder'] = 'Breeder', self.today, bid
            self._write_mouse(wb['Mice'], m)
            ids.append(bid)
        self._sync_count(wb, cage_name)
        return ids

    def create_pair(self, cage, line, male, females, setup=None, notes=''):
        """male / females may be breeder IDs or 'cage:X' (one animal is taken from that cage)."""
        wb = self._open_for_write()
        def resolve(u, sex):
            if u.startswith('cage:'):
                return self.take_from_cage(wb, u[5:], sex)[0]
            if u.startswith('mouse:'):
                m = self.mice[u[6:]]
                if m['breeder']:
                    return m['breeder']
                return self.take_from_cage(wb, m['cage'], sex, mouse_ids=[m['id']])[0]
            return u
        m = resolve(male, 'M')
        fs = [resolve(f, 'F') for f in females if f]
        for bid in [m] + fs:
            b = self.breeders[bid]
            b['status'], b['cage'] = 'Breeding', cage
            self._write_breeder(wb['Breeders'], b)
        ws = wb['Pairs']
        r = self._first_free_row(ws)
        self._copy_style(ws, r - 1, r, 14)
        vals = [cage, line, m, fs[0] if fs else None, fs[1] if len(fs) > 1 else None,
                setup or self.today, 'Active', notes or None] + pair_formulas(r)
        for c, v in enumerate(vals, 1):
            ws.cell(r, c).value = v
        ws.cell(r, 6).number_format = DATE_FMT
        ws.cell(r, 14).number_format = DATE_FMT
        self._set_cage(wb['Cages'], cage, line=line, sex='TP', count=1 + len(fs), birth='',
                       note='Breeding pair - see Pairs')
        self._save(wb)
        return m, fs

    def end_pair(self, pair_cage, clear_cage=True):
        wb = self._open_for_write()
        ws = wb['Pairs']
        for p in self.pairs:
            if p['cage'] == pair_cage and p['status'] == 'Active':
                ws.cell(p['row'], 7).value = 'Ended'
                for bid in self.pair_members(p):
                    if bid in self.breeders:
                        b = self.breeders[bid]
                        b['status'] = 'Retired'
                        self._write_breeder(wb['Breeders'], b)
        if clear_cage:
            self._set_cage(wb['Cages'], pair_cage, clear=True)
        self._save(wb)

    # ---------------- animal-house log (events) ----------------
    @staticmethod
    def event_kind(ev):
        k = _s(ev.get('event')).lower()
        if k in ('removed', 'moved', 'weaned', 'died', 'other'):
            return k
        w = _s(ev.get('event_as_written') or ev.get('event')).lower().strip()
        return next((v for kk, v in EVENT_WORDS.items() if w.startswith(kk)), 'other')

    def event_key(self, ev):
        to = ','.join(self.norm_cage(t.get('cage')) for t in (ev.get('to') or []) if isinstance(t, dict))
        return '|'.join([_s(ev.get('date')), self.norm_cage(ev.get('cage')), self.event_kind(ev),
                         str(int(ev.get('males') or 0)), str(int(ev.get('females') or 0)), to])

    def _pick_mice(self, cage, males, females):
        """Mice of the requested sexes from a stock cage (first by ear mark). Returns (ids, problem)."""
        mice = self.mice_in(cage)
        want = {'M': males, 'F': females}
        ids, short = [], []
        for sx, n in want.items():
            if not n:
                continue
            have = [m for m in mice if (m['sex'] or '').upper() == sx]
            if len(have) < n:
                short.append(f"{n} {sx} wanted, {len(have)} in the app")
            ids += [m['id'] for m in have[:n]]
        return ids, '; '.join(short)

    def compare_events(self, data):
        """Animal-house log rows -> change dicts like compare_cards, plus 'date' and 'key'."""
        out = []
        for ev in data.get('events', []):
            cage = self.norm_cage(ev.get('cage'))
            kind = self.event_kind(ev)
            males, females = int(ev.get('males') or 0), int(ev.get('females') or 0)
            date = self._card_date(ev.get('date'))
            key = self.event_key(ev)
            who = ' '.join(x for x in (f'{males} M' if males else '', f'{females} F' if females else '') if x) or '?'
            written = _s(ev.get('event_as_written') or ev.get('event'))
            unsure = set(ev.get('uncertain') or [])
            base = dict(cage=cage, date=date, key=key, card=f'{written}: {who}' + (
                f" (uncertain: {', '.join(sorted(unsure))})" if unsure else ''), ev=ev)

            def add(what, app, kindc, apply=True, value=None):
                out.append(dict(base, what=what, app='' if app is None else str(app), kind=kindc,
                                apply=apply and kindc != 'info' and not unsure, value=value))
            if key in self.event_keys:
                add('Already imported', f'on {fmt_date(self.event_keys[key]) if isinstance(self.event_keys[key], dt.date) else self.event_keys[key]}', 'info')
                continue
            if not date:
                add('No readable date - not applied', '', 'info')
                continue
            if date > self.today:
                add(f'Date {fmt_date(date)} is in the future - check the reading', '', 'info')
                continue
            old = (self.today - date).days > 45
            if old:
                unsure = unsure | {'date'}
                base['card'] += f' - OLD ROW ({(self.today - date).days} days): check it is not already in the app'
            if kind == 'weaned':
                add('Leválasztás - done from the cage cards', '', 'info')
                continue
            if kind not in ('removed', 'moved', 'died'):
                add(f'"{written}" - not applied (do it by hand if needed)', '', 'info')
                continue
            c = self.cage(cage)
            if not c:
                add('Unknown cage', '', 'info')
                continue
            pair = self.active_pair(cage)
            if kind in ('removed', 'died'):
                outcome = 'Found dead' if kind == 'died' else 'Terminated - experiment'
                if pair:
                    mem = [(bid, self.breeders[bid]) for bid in self.pair_members(pair) if bid in self.breeders]
                    ids = [b for b, x in mem if x['sex'] == 'M'][:males] + [b for b, x in mem if x['sex'] == 'F'][:females]
                    if len(ids) != males + females:
                        add('Breeding cage: these animals are not all breeders - do it by hand', ', '.join(
                            self.pair_members(pair)), 'info')
                        continue
                    add(f'{"Found dead" if kind == "died" else "Taken out (kivitel)"}: breeders {", ".join(ids)}',
                        'breeding cage ' + ', '.join(self.pair_members(pair)), 'term_breeders',
                        value=dict(ids=ids, outcome=outcome))
                    continue
                if self.mice_in(cage):
                    ids, problem = self._pick_mice(cage, males, females)
                    if problem:
                        add(f'Not enough mice ({problem})', f'{len(self.mice_in(cage))} mice', 'info')
                        continue
                    lab = ', '.join(f"{i} ({self.mice[i]['mark']})" for i in ids)
                    add(f'{"Found dead" if kind == "died" else "Taken out (kivitel)"}: {lab}',
                        f"{len(self.mice_in(cage))} {c['sex']} {c['line']}", 'term_mice',
                        value=dict(ids=ids, outcome=outcome))
                    continue
                n = males + females
                if not c['count'] or n > c['count']:
                    add('Not enough animals in the app', f"{c['count'] or 0}", 'info')
                    continue
                add(f'{"Found dead" if kind == "died" else "Taken out (kivitel)"}: {n} of {c["count"]}',
                    f"{c['count']} {c['sex']} {c['line']}", 'term_cage', value=dict(n=n, outcome=outcome))
                continue
            # moved (áthelyezés)
            tos = [t for t in (ev.get('to') or []) if isinstance(t, dict) and _s(t.get('cage'))]
            if not tos:
                add('Moved - no target cage on the log', '', 'info')
                continue
            if pair:
                add('Breeders moved - change the pair by hand', ', '.join(self.pair_members(pair)), 'info')
                continue
            for t in tos:
                dest = self.norm_cage(t.get('cage'))
                tm = int(t.get('males') or 0) if ('males' in t or 'females' in t) else males
                tf = int(t.get('females') or 0) if ('males' in t or 'females' in t) else females
                if dest == cage:
                    add('Moved within the same cage - nothing to do', '', 'info')
                    continue
                if not self.cage(dest):
                    add(f'Unknown target cage {t.get("cage")}', '', 'info')
                    continue
                if self.active_pair(dest):
                    add(f'Target {self.card_name(dest)} is a breeding cage - set up the pair by hand', '', 'info')
                    continue
                ids, problem = self._pick_mice(cage, tm, tf)
                if problem or not ids:
                    add(f'Not enough mice to move ({problem or "none listed"})', f'{len(self.mice_in(cage))} mice',
                        'info')
                    continue
                there = self.mice_in(dest)
                add(f'Move to {self.card_name(dest)}: ' + ', '.join(f"{i} ({self.mice[i]['mark']})" for i in ids),
                    f"{len(self.mice_in(cage))} here" + (f'; {len(there)} already in {self.card_name(dest)}'
                                                          if there else ''), 'move',
                    value=dict(ids=ids, dest=dest))
        return out

    def move_mice(self, ids, dest):
        """Move stock mice to another (non-breeding) cage in one write; marks are re-assigned if taken."""
        if self.active_pair(dest):
            raise ValueError(f'{self.card_name(dest)} is a breeding cage.')
        wb = self._open_for_write()
        olds = set()
        for mid in ids:
            m = self.mice[mid]
            if m['cage'] == dest:
                continue
            olds.add(m['cage'])
            if not self.mice_in(dest):
                self._set_cage(wb['Cages'], dest, line=m['line'], sex=m['sex'], birth=m['birth'] or '',
                               sire=m['sire'] or '', dam=m['dam'] or '', note='')
            taken = {x['mark'] for x in self.mice_in(dest)}
            m['cage'] = dest
            if m['mark'] in taken:
                m['mark'] = next(x for x in self.free_marks(dest, len(taken) + 1) if x not in taken)
            self._write_mouse(wb['Mice'], m)
        for c in olds | {dest}:
            self._sync_count(wb, c)
        self._save(wb)

    def _log_events(self, rows):
        wb = self._open_for_write()
        if EVENTS_SHEET not in wb.sheetnames:
            ws = wb.create_sheet(EVENTS_SHEET)
            for i, h in enumerate(EVENTS_HEAD, 1):
                ws.cell(1, i).value = h
                ws.cell(1, i).font = openpyxl.styles.Font(name='Arial', size=10, bold=True)
        ws = wb[EVENTS_SHEET]
        today = self.today if isinstance(getattr(self, 'today', None), dt.date) else dt.date.today()
        for key, date, cage, kind, details in rows:
            r = self._first_free_row(ws)
            for i, v in enumerate([key, date, self.card_name(cage), kind, details, today, self.user or None], 1):
                ws.cell(r, i).value = v
                ws.cell(r, i).font = openpyxl.styles.Font(name='Arial', size=10)
            ws.cell(r, 2).number_format = DATE_FMT
            ws.cell(r, 6).number_format = DATE_FMT
        self._save(wb)

    def apply_events(self, changes, person=''):
        """Apply ticked event changes one by one; every applied log row is recorded so it is never applied twice.
        Returns the number applied."""
        done, logged = [], set()
        try:
            for ch in changes:
                if not ch['apply'] or ch['kind'] == 'info':
                    continue
                ev, v = ch['ev'], ch['value']
                who = _s(ev.get('person')) or person
                sev = _s(ev.get('severity'))
                kw = dict(reason=(f'severity {sev}' if sev else ''), person=who,
                          notes=f"animal-house log: {_s(ev.get('event_as_written') or ev.get('event'))}")
                if ch['kind'] == 'term_mice':
                    self.terminate_mice(v['ids'], ch['date'], v['outcome'], **kw)
                elif ch['kind'] == 'term_breeders':
                    self.terminate_breeders(v['ids'], ch['date'], v['outcome'], **kw)
                elif ch['kind'] == 'term_cage':
                    self.terminate_cage(ch['cage'], v['n'], ch['date'], v['outcome'], **kw)
                elif ch['kind'] == 'move':
                    self.move_mice(v['ids'], v['dest'])
                if ch['key'] not in logged:
                    done.append((ch['key'], ch['date'], ch['cage'], self.event_kind(ev), ch['what'][:200]))
                    logged.add(ch['key'])
        finally:
            if done:
                self._log_events(done)
        return len(done)

    def wean_litter(self, pair_cage, dam, born, weaned, groups, pups_born=None, notes='', unknown_gt=True):
        """groups: list of (sex, count, target_cage). unknown_gt: flag the pups 'needs genotyping' (unless their
        genotype is certain from the parents). Returns {target_cage: [new mouse IDs]}."""
        p = next((x for x in self.pairs if x['cage'] == pair_cage and x['status'] == 'Active'), None)
        if not p:
            raise ValueError(f'No active pair in {pair_cage}')
        weaned = _d(weaned) or weaned
        wb = self._open_for_write()
        wsc = wb['Cages']
        created = {}
        for sex, cnt, target in groups:
            c = self.cage(target)
            if c and (c['line'] or c['count']):
                raise ValueError(f'Cage {target} is not empty ({c["line"]}, {c["count"]} {c["sex"]})')
            created[target] = self._add_mice(wb, target, cnt, p['line'], sex, born, sire=p['m'], dam=dam,
                                             notes=f'Weaned from {pair_cage} on {weaned}',
                                             flag=None)
        ws = wb['Litters']
        r = self._first_free_row(ws)
        gm = next(((n, t) for s, n, t in groups if s == 'M'), (None, None))
        gf = next(((n, t) for s, n, t in groups if s == 'F'), (None, None))
        vals = [born, pair_cage, p['m'], dam, pups_born, weaned, gm[0], gf[0], gm[1], gf[1], notes or None]
        for c, v in enumerate(vals, 1):
            ws.cell(r, c).value = v
            ws.cell(r, c).font = openpyxl.styles.Font(name='Arial', size=10)
        ws.cell(r, 1).number_format = DATE_FMT
        ws.cell(r, 6).number_format = DATE_FMT
        self._save(wb)
        return created


# --------------------------------------------------------------------------
# GUI
# --------------------------------------------------------------------------
def run_gui(path=None):
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox

    COLORS = {'OVER AGE': '#F8CBAD', 'Replace soon': '#FFE699', 'Too young': '#DDEBF7', 'NO DATE': '#D9D9D9',
              'OK': '#E2F0D9', 'NO ACTIVE PAIR': '#F8CBAD', 'Pair over age': '#F8CBAD',
              'Plan replacement': '#FFE699', 'bad': '#F8CBAD', 'warn': '#FFE699', 'ok': '#E2F0D9', 'info': '#EDEDED',
              'Retired': '#EDEDED', 'pair': '#EDE7F6', 'empty': '#F7F7F7', 'Candidate': '#E2F0D9',
              'Stock': '#FFFFFF', 'Terminate': '#F8CBAD', 'Pair not on Pairs sheet': '#FFE699',
              'Terminated': '#E7D3D3', 'proc': '#EAF1FB'}

    import traceback
    dnd_files = None
    root = None
    if not os.environ.get('MOUSECOLONY_NO_DND'):
        try:                                   # optional: drag & drop of files onto the window
            from tkinterdnd2 import TkinterDnD, DND_FILES
            root = TkinterDnD.Tk()
            dnd_files = DND_FILES
        except Exception:
            root = None
    if root is None:
        root = tk.Tk()
    root.title(f'{APP_NAME} {VERSION}')
    root.geometry('1180x720')
    here = getattr(sys, '_MEIPASS', os.path.dirname(os.path.abspath(__file__)))
    try:
        if sys.platform.startswith('win') and os.path.exists(os.path.join(here, 'mousecolony.ico')):
            root.iconbitmap(default=os.path.join(here, 'mousecolony.ico'))
        elif os.path.exists(os.path.join(here, 'mousecolony.png')):
            root._icon = tk.PhotoImage(file=os.path.join(here, 'mousecolony.png'))
            root.iconphoto(True, root._icon)
    except tk.TclError:
        pass

    def tk_error(etype, value, tb):
        log_error(''.join(traceback.format_exception(etype, value, tb)))
        messagebox.showerror(APP_NAME, f'{value}\n\nDetails were written to\n{ERROR_LOG}')
    root.report_callback_exception = tk_error
    style = ttk.Style()
    try:
        style.theme_use('clam')
    except tk.TclError:
        pass
    style.configure('Treeview', rowheight=24)
    style.configure('Treeview.Heading', font=('TkDefaultFont', 9, 'bold'))

    state = {'col': None}

    def parse_date(s):
        s = s.strip().replace('.', '-').strip('-')
        return dt.datetime.strptime(s, '%Y-%m-%d') if s else None

    def make_tree(parent, cols, widths, height=12, show='headings'):
        frame = ttk.Frame(parent)
        tree = ttk.Treeview(frame, columns=cols, show=show, height=height)
        if show != 'headings':
            tree.column('#0', width=22, minwidth=22, stretch=False)
        for c, w in zip(cols, widths):
            tree.heading(c, text=c, command=lambda c=c: sort_tree(tree, c))
            tree.column(c, width=w, minwidth=30, anchor='w', stretch=False)
        sb = ttk.Scrollbar(frame, orient='vertical', command=tree.yview)
        hs = ttk.Scrollbar(frame, orient='horizontal', command=tree.xview)
        tree.configure(yscrollcommand=sb.set, xscrollcommand=hs.set)
        tree.grid(row=0, column=0, sticky='nsew')
        sb.grid(row=0, column=1, sticky='ns')
        hs.grid(row=1, column=0, sticky='ew')
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)
        for tag, col in COLORS.items():
            tree.tag_configure(tag, background=col)
        return frame, tree

    def sort_tree(tree, col):
        items = [(tree.set(k, col), k) for k in tree.get_children('')]
        rev = getattr(tree, '_rev', {}).get(col, False)
        def key(x):
            try:
                return (0, float(x[0].split()[0]))
            except (ValueError, IndexError):
                return (1, x[0])
        items.sort(key=key, reverse=rev)
        for i, (_, k) in enumerate(items):
            tree.move(k, '', i)
        tree._rev = {col: not rev}

    def fill(tree, rows):
        tree.delete(*tree.get_children())
        for vals, tag in rows:
            tree.insert('', 'end', values=vals, tags=(tag,) if tag else ())

    def guard(fn):
        def wrapped(*a, **k):
            if state['col'] is None:
                messagebox.showinfo('No workbook', 'Open the colony workbook first.')
                return
            try:
                return fn(*a, **k)
            except ConflictError as e:
                messagebox.showwarning('Changed by a colleague',
                                       f'{e}\n\nYour change was NOT saved, so nothing of theirs was overwritten. '
                                       f'The latest version is shown now - check it and do the change again.')
                try:
                    refresh_all()
                except Exception:
                    pass
            except ReadOnlyError as e:
                messagebox.showwarning('Read-only', str(e))
            except SyncError as e:
                messagebox.showerror('Shared sheet', str(e))
                try:
                    refresh_all()
                except Exception:
                    pass
            except PermissionError:
                messagebox.showerror('File is open', 'Could not save - close the workbook in Excel and try again.')
            except Exception as e:  # show anything else rather than crash
                log_error(traceback.format_exc())
                messagebox.showerror('Error', str(e))
        return wrapped

    # ---------- top bar ----------
    top = ttk.Frame(root, padding=6)
    top.pack(fill='x')
    path_var = tk.StringVar(value='No workbook loaded')
    ttk.Button(top, text='Shared database...', command=lambda: shared_dialog()).pack(side='left', padx=(0, 4))
    ttk.Button(top, text='Open workbook...', command=lambda: open_file()).pack(side='left')
    ttk.Button(top, text='Reload', command=lambda: reload()).pack(side='left', padx=4)
    ttk.Button(top, text='Import cage cards...', command=lambda: import_cards()).pack(side='left')
    ttk.Button(top, text='Read card photos...', command=lambda: read_photos()).pack(side='left', padx=(4, 0))
    ttk.Button(top, text='Paste cards', command=lambda: paste_cards()).pack(side='left', padx=(4, 0))
    ttk.Button(top, text='Enter cage card...', command=lambda: enter_cards()).pack(side='left', padx=4)
    ttk.Label(top, textvariable=path_var).pack(side='left', padx=10)
    sync_var = tk.StringVar()
    sync_lbl = tk.Label(top, textvariable=sync_var, font=('TkDefaultFont', 9, 'bold'))
    sync_lbl.pack(side='left', padx=(0, 10))
    info_var = tk.StringVar()
    ttk.Label(top, textvariable=info_var, foreground='#555').pack(side='right')

    banner_var = tk.StringVar()
    banner = tk.Label(root, textvariable=banner_var, bg='#C00000', fg='white', font=('TkDefaultFont', 10, 'bold'),
                      anchor='w', padx=8)

    def set_status():
        col = state['col']
        if col is None:
            return
        if col.remote:
            who = state.get('user') or 'name not set'
            path_var.set(f'Shared Google Sheet: {col.remote.title}   (you: {who})')
        else:
            p_ = col.path
            path_var.set(p_ if len(p_) <= 70 else '...' + p_[-67:])
        if col.remote:
            if state.get('syncing'):
                sync_var.set('Uploading...' if col.pending else 'Refreshing...')
                sync_lbl.config(fg='#1a5fb4')
            elif col.pending:
                sync_var.set(f'{col.pending} change{"" if col.pending == 1 else "s"} not uploaded yet')
                sync_lbl.config(fg='#B35C00')
            else:
                sync_var.set('All saved to Google')
                sync_lbl.config(fg='#2E7D32')
        else:
            sync_var.set('')
        if col.remote and col.readonly:
            banner_var.set(f'  OFFLINE - working on this computer\'s copy. '
                           + (f'{col.pending} change(s) will be uploaded when the shared sheet is reachable again. '
                              if col.pending else 'You can keep working; changes are uploaded later. ')
                           + f'({col.offline_reason[:80]})')
            banner.pack(fill='x', before=nb)
        else:
            banner.pack_forget()

    nb = ttk.Notebook(root)
    nb.pack(fill='both', expand=True, padx=6, pady=6)

    # ---------- Overview ----------
    ov = ttk.Frame(nb, padding=6)
    nb.add(ov, text='Overview')
    ttk.Label(ov, text='Lines', font=('TkDefaultFont', 10, 'bold')).pack(anchor='w')
    lf, lines_tree = make_tree(ov, ['Line', 'Must keep', 'Active pairs', 'Over age', 'Replace soon',
                                    'Candidate M', 'Candidate F', 'Status'],
                               [200, 70, 80, 70, 90, 90, 90, 160], height=9)
    lf.pack(fill='x')
    olb = ttk.Frame(ov)
    olb.pack(fill='x', pady=(2, 0))
    ttk.Button(olb, text='New line...', command=lambda: new_line_form(root)).pack(side='left')
    ttk.Button(olb, text='Combine lines (cross)...',
               command=lambda: line_chooser(root, '', lambda n: messagebox.showinfo(
                   'Line', f'Line "{n}" is ready to use.'))).pack(side='left', padx=4)
    ttk.Button(olb, text='Edit alleles of selected line...', command=lambda: edit_line_alleles()).pack(side='left')
    ttk.Button(olb, text='Merge selected line into...', command=lambda: merge_line_dialog()).pack(side='left',
                                                                                                padx=4)
    ttk.Label(ov, text='Breeding pairs', font=('TkDefaultFont', 10, 'bold')).pack(anchor='w', pady=(8, 0))
    pf, pairs_tree = make_tree(ov, ['Cage', 'Line', 'Male', 'Female(s)', 'Oldest', 'Age flag', 'Replace by',
                                    'Relatedness', 'Notes'], [60, 170, 60, 100, 70, 90, 90, 230, 200], height=12)
    pf.pack(fill='both', expand=True)
    ob = ttk.Frame(ov)
    ob.pack(fill='x', pady=4)
    ttk.Button(ob, text='Check selected pair', command=lambda: check_selected_pair()).pack(side='left')
    ttk.Button(ob, text='Suggest replacement', command=lambda: suggest_for_selected()).pack(side='left', padx=4)
    ttk.Button(ob, text='End selected pair (retire animals)', command=lambda: end_selected()).pack(side='left')

    # ---------- All mice ----------
    am = ttk.Frame(nb, padding=6)
    nb.add(am, text='All mice')
    amf = ttk.Frame(am)
    amf.pack(fill='x')
    ttk.Label(amf, text='Rack').pack(side='left')
    am_rack = ttk.Combobox(amf, width=7, state='readonly')
    am_rack.pack(side='left', padx=(2, 8))
    ttk.Label(amf, text='Line').pack(side='left')
    am_line = ttk.Combobox(amf, width=24, state='readonly')
    am_line.pack(side='left', padx=(2, 8))
    ttk.Label(amf, text='Show').pack(side='left')
    am_kind = ttk.Combobox(amf, width=16, state='readonly',
                           values=['(all)', 'Males', 'Females', 'Breeding pairs', 'Breeder candidates',
                                   'Needs genotyping', 'Needs attention'])
    am_kind.set('(all)')
    am_kind.pack(side='left', padx=(2, 8))
    ttk.Label(amf, text='Search').pack(side='left')
    am_search = tk.StringVar()
    ttk.Entry(amf, textvariable=am_search, width=16).pack(side='left', padx=(2, 8))
    am_empty = tk.BooleanVar(value=False)
    ttk.Checkbutton(amf, text='Empty cages', variable=am_empty,
                    command=lambda: refresh_all_mice()).pack(side='left')
    am_open = tk.BooleanVar(value=True)
    ttk.Checkbutton(amf, text='Show each mouse', variable=am_open,
                    command=lambda: refresh_all_mice()).pack(side='left', padx=(6, 0))
    for w in (am_rack, am_line, am_kind):
        w.bind('<<ComboboxSelected>>', lambda e: refresh_all_mice())
    am_search.trace_add('write', lambda *a: refresh_all_mice())
    am_total = tk.StringVar()
    ttk.Label(am, textvariable=am_total, font=('TkDefaultFont', 10, 'bold')).pack(anchor='w', pady=(6, 2))
    amt, am_tree = make_tree(am, ['Cage / mouse', 'Line', 'Sex', 'Mice', 'Born', 'Age', 'Status', 'Animals',
                                  'Parents', 'Genotype', 'Note'],
                             [88, 160, 38, 42, 94, 60, 175, 190, 100, 210, 200], height=12, show='tree headings')
    am_tree.tag_configure('mouse', background='#FFFFFF', foreground='#333333')
    am_tree.tag_configure('mouse_proc', background='#F3F7FC')
    am_tree.tag_configure('todo', background='#FFE699', foreground='#7F4F00')
    amt.pack(fill='both', expand=True)
    ttk.Label(am, text='Mice per line (filtered)', font=('TkDefaultFont', 10, 'bold')).pack(anchor='w',
                                                                                         pady=(6, 0))
    aml, am_lines_tree = make_tree(am, ['Line', 'Cages', 'Mice', 'Males', 'Females', 'In pairs',
                                        'Candidate cages'], [200, 60, 60, 60, 70, 70, 110], height=5)
    aml.pack(fill='x')
    amb = ttk.Frame(am)
    amb.pack(fill='x', pady=4)
    ttk.Button(amb, text='Log procedure...', command=lambda: proc_for_selected_cage()).pack(side='left')
    ttk.Button(amb, text='Genotyping results...', command=lambda: genotyping_for_selected()).pack(side='left',
                                                                                                padx=(4, 0))
    ttk.Button(amb, text='Terminate / remove...', command=lambda: term_for_selected_cage()).pack(side='left',
                                                                                              padx=4)
    ttk.Button(amb, text='Cage history', command=lambda: history_of_selected_cage()).pack(side='left')
    ttk.Separator(amb, orient='vertical').pack(side='left', fill='y', padx=8)
    ttk.Button(amb, text='Edit...', command=lambda: edit_selected_item()).pack(side='left')
    ttk.Button(amb, text='Add to empty cage...', command=lambda: cage_dialog(None)).pack(side='left', padx=4)
    ttk.Button(amb, text='Clear (data fix)...', command=lambda: empty_selected_cage()).pack(side='left')
    am_tree.bind('<Double-1>', lambda e: edit_selected_item())
    am_rows = {}

    # ---------- Breeders ----------
    bt = ttk.Frame(nb, padding=6)
    nb.add(bt, text='Breeders')
    bfil = ttk.Frame(bt)
    bfil.pack(fill='x')
    ttk.Label(bfil, text='Line').pack(side='left')
    b_line = ttk.Combobox(bfil, width=26, state='readonly')
    b_line.pack(side='left', padx=4)
    ttk.Label(bfil, text='Status').pack(side='left')
    b_status = ttk.Combobox(bfil, width=12, state='readonly', values=['(all)', 'Breeding', 'Reserve', 'Retired',
                                                                       'Terminated', '(alive)'])
    b_status.set('(alive)')
    b_status.pack(side='left', padx=4)
    b_line.bind('<<ComboboxSelected>>', lambda e: refresh_breeders())
    b_status.bind('<<ComboboxSelected>>', lambda e: refresh_breeders())
    bf, br_tree = make_tree(bt, ['ID', 'Line', 'Sex', 'Birth', 'Age', 'Age flag', 'Cage', 'Status', 'Sire', 'Dam',
                                 'Inbreeding F', 'Genotype', 'Notes'],
                            [55, 160, 38, 88, 62, 90, 58, 78, 52, 52, 78, 190, 200], height=20)
    bf.pack(fill='both', expand=True, pady=4)
    bb = ttk.Frame(bt)
    bb.pack(fill='x')
    ttk.Button(bb, text='Add breeder...', command=lambda: breeder_dialog(None)).pack(side='left')
    ttk.Button(bb, text='Edit selected...', command=lambda: edit_selected_breeder()).pack(side='left', padx=4)
    ttk.Button(bb, text='Take breeders from a stock cage...', command=lambda: take_dialog()).pack(side='left')
    ttk.Separator(bb, orient='vertical').pack(side='left', fill='y', padx=8)
    ttk.Button(bb, text='Log procedure...', command=lambda: proc_for_selected_breeder()).pack(side='left')
    ttk.Button(bb, text='Terminate...', command=lambda: term_for_selected_breeder()).pack(side='left', padx=4)

    # ---------- History ----------
    hi = ttk.Frame(nb, padding=6)
    nb.add(hi, text='History')
    hif = ttk.Frame(hi)
    hif.pack(fill='x')
    ttk.Label(hif, text='Line').pack(side='left')
    h_line = ttk.Combobox(hif, width=24, state='readonly')
    h_line.pack(side='left', padx=(2, 8))
    ttk.Label(hif, text='From').pack(side='left')
    h_from = ttk.Entry(hif, width=11)
    h_from.pack(side='left', padx=(2, 4))
    ttk.Label(hif, text='to').pack(side='left')
    h_to = ttk.Entry(hif, width=11)
    h_to.pack(side='left', padx=(2, 8))
    ttk.Label(hif, text='Search').pack(side='left')
    h_search = tk.StringVar()
    ttk.Entry(hif, textvariable=h_search, width=18).pack(side='left', padx=(2, 8))
    ttk.Button(hif, text='Apply', command=lambda: refresh_history()).pack(side='left')
    ttk.Button(hif, text='Clear filters', command=lambda: clear_history_filters()).pack(side='left', padx=4)
    h_line.bind('<<ComboboxSelected>>', lambda e: refresh_history())
    h_search.trace_add('write', lambda *a: refresh_history())
    h_total = tk.StringVar()
    ttk.Label(hi, textvariable=h_total, font=('TkDefaultFont', 10, 'bold')).pack(anchor='w', pady=(6, 2))
    htf, term_tree = make_tree(hi, ['Date', 'Cage', 'Animal', 'Line', 'Sex', 'Mice', 'Age', 'Outcome',
                                    'Reason / use', 'Project', 'Person', 'Genotype', 'Procedures done'],
                               [85, 62, 70, 150, 36, 40, 58, 160, 140, 90, 80, 150, 330], height=9)
    htf.pack(fill='both', expand=True)
    ttk.Label(hi, text='Procedure log', font=('TkDefaultFont', 10, 'bold')).pack(anchor='w', pady=(8, 0))
    hpf, proc_tree = make_tree(hi, ['Date', 'Cage', 'Animal', 'Line', 'Mice', 'Procedure', 'Details', 'Person',
                                    'Notes', 'Still alive?'],
                               [85, 62, 70, 150, 40, 150, 230, 90, 200, 90], height=7)
    hpf.pack(fill='both', expand=True)
    hb = ttk.Frame(hi)
    hb.pack(fill='x', pady=4)
    ttk.Button(hb, text='Delete selected procedure (mistake)', command=lambda: delete_selected_proc()).pack(
        side='left')
    ttk.Label(hb, text='  Terminations are permanent records - correct them in the Terminated sheet if needed.',
              foreground='#555').pack(side='left')
    proc_rows = {}

    # ---------- Pairing check ----------
    pc = ttk.Frame(nb, padding=10)
    nb.add(pc, text='Pairing check')
    ttk.Label(pc, text='Pick breeders (IDs) or stock cages. Leave Female 2 empty for a pair.').grid(
        row=0, column=0, columnspan=4, sticky='w')
    pc_m = ttk.Combobox(pc, width=45)
    pc_f1 = ttk.Combobox(pc, width=45)
    pc_f2 = ttk.Combobox(pc, width=45)
    for i, (lab, w) in enumerate([('Male', pc_m), ('Female 1', pc_f1), ('Female 2', pc_f2)], 1):
        ttk.Label(pc, text=lab).grid(row=i, column=0, sticky='w', pady=3)
        w.grid(row=i, column=1, sticky='w', pady=3)
    pc_all = tk.BooleanVar(value=False)
    ttk.Checkbutton(pc, text='List all animals (not only candidates)', variable=pc_all,
                    command=lambda: refresh_pc_lists()).grid(row=4, column=1, sticky='w')
    ttk.Button(pc, text='Check', command=lambda: run_check()).grid(row=5, column=1, sticky='w', pady=6)
    ttk.Button(pc, text='Create this pair...', command=lambda: create_from_check()).grid(row=5, column=1, padx=80,
                                                                                         sticky='w')
    cf, check_tree = make_tree(pc, ['Result'], [1000], height=14)
    cf.grid(row=6, column=0, columnspan=4, sticky='nsew')
    pc.rowconfigure(6, weight=1)
    pc.columnconfigure(3, weight=1)

    # ---------- Suggestions ----------
    sg = ttk.Frame(nb, padding=6)
    nb.add(sg, text='Suggest pairs')
    sgt = ttk.Frame(sg)
    sgt.pack(fill='x')
    ttk.Label(sgt, text='Line').pack(side='left')
    sg_line = ttk.Combobox(sgt, width=28, state='readonly')
    sg_line.pack(side='left', padx=4)
    ttk.Button(sgt, text='Suggest', command=lambda: run_suggest()).pack(side='left')
    ttk.Label(sgt, text='  Candidates: Reserve breeders + stock cages aged MinAge..StartMax, from the line\'s '
                        'allowed partner lines. Least related first, then longest breeding time left.',
              foreground='#555').pack(side='left')
    sf, sug_tree = make_tree(sg, ['Male', 'Female', 'Offspring F', 'Risk', 'Relationship', 'Breeding days left',
                                  'Expected pups'],
                             [230, 230, 80, 170, 160, 100, 300], height=20)
    sf.pack(fill='both', expand=True, pady=4)
    ttk.Button(sg, text='Check / create selected suggestion', command=lambda: suggestion_to_check()).pack(anchor='w')

    # ---------- Due soon ----------
    du = ttk.Frame(nb, padding=6)
    nb.add(du, text='Due soon')
    dut = ttk.Frame(du)
    dut.pack(fill='x')
    ttk.Label(dut, text='Pairs reaching the age limit within').pack(side='left')
    du_days = ttk.Spinbox(dut, from_=0, to=365, width=5)
    du_days.set(90)
    du_days.pack(side='left', padx=4)
    ttk.Label(dut, text='days').pack(side='left')
    ttk.Button(dut, text='Show', command=lambda: refresh_due()).pack(side='left', padx=6)
    df, due_tree = make_tree(du, ['Cage', 'Line', 'Animals', 'Oldest', 'Replace by', 'Candidate M', 'Candidate F'],
                             [60, 180, 140, 70, 90, 300, 300], height=20)
    df.pack(fill='both', expand=True, pady=4)

    # ---------- Wean litter ----------
    wl = ttk.Frame(nb, padding=10)
    nb.add(wl, text='Wean litter')
    ttk.Label(wl, text='Puts weaned pups into empty cages with their parents recorded, and logs the litter.',
              foreground='#555').grid(row=0, column=0, columnspan=4, sticky='w', pady=(0, 6))
    wl_pair = ttk.Combobox(wl, width=30, state='readonly')
    wl_dam = ttk.Combobox(wl, width=30, state='readonly')
    wl_born, wl_weaned, wl_pups = ttk.Entry(wl, width=14), ttk.Entry(wl, width=14), ttk.Entry(wl, width=6)
    wl_mn, wl_mc = ttk.Entry(wl, width=6), ttk.Combobox(wl, width=10)
    wl_fn, wl_fc = ttk.Entry(wl, width=6), ttk.Combobox(wl, width=10)
    wl_notes = ttk.Entry(wl, width=50)
    rows = [('Breeding pair', wl_pair), ('Mother', wl_dam), ('Born (yyyy-mm-dd)', wl_born),
            ('Weaned (yyyy-mm-dd)', wl_weaned), ('Pups born', wl_pups)]
    for i, (lab, w) in enumerate(rows, 1):
        ttk.Label(wl, text=lab).grid(row=i, column=0, sticky='w', pady=3)
        w.grid(row=i, column=1, sticky='w', pady=3)
    ttk.Label(wl, text='Males: count / cage').grid(row=6, column=0, sticky='w', pady=3)
    wl_mn.grid(row=6, column=1, sticky='w')
    wl_mc.grid(row=6, column=1, padx=60, sticky='w')
    ttk.Label(wl, text='Females: count / cage').grid(row=7, column=0, sticky='w', pady=3)
    wl_fn.grid(row=7, column=1, sticky='w')
    wl_fc.grid(row=7, column=1, padx=60, sticky='w')
    ttk.Label(wl, text='Notes').grid(row=8, column=0, sticky='w', pady=3)
    wl_notes.grid(row=8, column=1, sticky='w')
    ttk.Label(wl, text='Pup genotype').grid(row=9, column=0, sticky='nw', pady=3)
    wl_unknown = tk.BooleanVar(value=True)
    ttk.Checkbutton(wl, text='Unknown - flag the pups "needs genotyping" (untick to enter results right away)',
                    variable=wl_unknown).grid(row=9, column=1, sticky='w')
    wl_exp = tk.StringVar()
    ttk.Label(wl, textvariable=wl_exp, foreground='#1F4E78', wraplength=620, justify='left').grid(
        row=10, column=1, sticky='w', pady=(2, 0))
    wl_weaned.insert(0, dt.date.today().isoformat())
    ttk.Button(wl, text='Save litter', command=lambda: save_litter()).grid(row=11, column=1, sticky='w', pady=8)

    def on_pair_pick(_=None):
        col = state['col']
        cage = wl_pair.get().split(' ')[0]
        p = next((x for x in col.pairs if x['cage'] == cage and x['status'] == 'Active'), None)
        fs = [f for f in (p['f1'], p['f2']) if f] if p else []
        wl_dam['values'] = fs
        wl_dam.set(fs[0] if fs else '')
        show_expected()

    def show_expected(_=None):
        col = state['col']
        cage = wl_pair.get().split(' ')[0]
        p = col.active_pair(cage) if col and cage else None
        if not p or not wl_dam.get():
            wl_exp.set('')
            return
        txt = col.pups_text(p['m'], wl_dam.get())
        certain = all(len(v or {}) == 1 for v in col.expected_pups(p['m'], wl_dam.get()).values() if v is not None) \
            and None not in col.expected_pups(p['m'], wl_dam.get()).values()
        wl_exp.set(('Expected pups: ' + txt if txt else 'No alleles set for this line.')
                   + ('\nAll pups have a certain genotype - they will not be flagged.' if txt and certain else ''))
    wl_pair.bind('<<ComboboxSelected>>', on_pair_pick)
    wl_dam.bind('<<ComboboxSelected>>', show_expected)

    # ---------- refreshers ----------
    def unit_label(u):
        inf = state['col'].unit_info(u)
        return f'{u}  -  {inf["label"]}' if inf and not u.startswith('cage:') else (u + '  -  ' + inf['label'] if inf else u)

    def unit_from_label(s):
        return s.split('  -  ')[0].strip()

    def refresh_all():
        col = state['col']
        info_var.set(f"Today {col.today}  |  max {col.settings['MaxAge']} d, warn {col.settings['WarnDays']} d, "
                     f"min {col.settings['MinAge']} d")
        fill(lines_tree, [((r['line'], 'Y' if r['keep'] else '', r['active'], r['over'], r['soon'], r['cand_m'],
                            r['cand_f'], r['status']), r['status'] or None) for r in col.line_status()])
        prow = []
        for p in col.pairs:
            if p['status'] != 'Active':
                continue
            fs = [f for f in (p['f1'], p['f2']) if f]
            ks = [col.kinship(p['m'], f) for f in fs] if p['m'] else []
            kmax = max(ks) if ks else 0
            rel = col.risk(kmax) if kmax else 'none recorded'
            prow.append(((p['cage'], p['line'], p['m'], ', '.join(fs), fmt_age(col.pair_oldest(p)),
                          col.pair_flag(p), fmt_date(col.pair_replace_by(p)), rel, p['notes']), col.pair_flag(p)))
        fill(pairs_tree, prow)
        lines = sorted(col.lines)
        b_line['values'] = ['(all)'] + lines
        if not b_line.get():
            b_line.set('(all)')
        sg_line['values'] = lines
        racks = sorted({col.rack(c['cage']) for c in col.cages})
        am_rack['values'] = ['(all)'] + racks
        am_line['values'] = ['(all)'] + lines
        for w in (am_rack, am_line):
            if not w.get():
                w.set('(all)')
        refresh_all_mice()
        h_line['values'] = ['(all)'] + sorted(set(lines) | {t['line'] for t in col.terminated if t['line']})
        if not h_line.get():
            h_line.set('(all)')
        refresh_history()
        refresh_breeders()
        refresh_pc_lists()
        refresh_due()
        wl_pair['values'] = [f"{p['cage']}  ({p['line']})" for p in col.pairs if p['status'] == 'Active']
        empties = col.empty_cages()
        wl_mc['values'] = empties
        wl_fc['values'] = empties

    MORE = '＋ Combine lines / new line…'

    def enhance_line_combo(cb, on_change=None):
        """Line combobox: lists all lines (fresh every time it opens) plus an entry that opens the chooser."""
        cb.configure(postcommand=lambda: cb.configure(values=sorted(state['col'].lines) + [MORE]))
        cb.configure(values=sorted(state['col'].lines) + [MORE])

        def picked(_=None):
            if cb.get() == MORE:
                prev = getattr(cb, '_prev', '')
                cb.set(prev)
                line_chooser(cb.winfo_toplevel(), prev, lambda name: (cb.set(name), setattr(cb, '_prev', name),
                                                                     cb.event_generate('<<ComboboxSelected>>')))
            else:
                cb._prev = cb.get()
                if on_change:
                    on_change(cb.get())
        cb.bind('<<ComboboxSelected>>', picked, add='+')
        cb._prev = cb.get()
        return cb

    def line_chooser(parent, current, on_pick):
        """Pick one line, tick several to combine them into a mixed line, or define a new line."""
        col = state['col']
        d = tk.Toplevel(parent)
        d.title('Choose line(s)')
        d.transient(parent)
        d.grab_set()
        ttk.Label(d, text='Tick one line, or several to make a cross (mixed line).',
                  font=('TkDefaultFont', 10, 'bold')).grid(row=0, column=0, columnspan=2, sticky='w', padx=10,
                                                           pady=(10, 4))
        box = ttk.Frame(d)
        box.grid(row=1, column=0, sticky='nw', padx=10)
        vars_ = {}

        def fill_boxes():
            for w_ in box.winfo_children():
                w_.destroy()
            for i, ln in enumerate(sorted(col.lines)):
                v = vars_.get(ln) or tk.BooleanVar(value=(ln == current))
                vars_[ln] = v
                loci = ', '.join(col.lines[ln]['loci']) or ('wild type' if ln == 'C57Bl6J' else 'no alleles set')
                ttk.Checkbutton(box, variable=v, text=f'{ln}   ({loci})', command=update).grid(
                    row=i, column=0, sticky='w')
        result = tk.StringVar()
        ttk.Label(d, textvariable=result, foreground='#1F4E78', wraplength=420, justify='left').grid(
            row=2, column=0, columnspan=2, sticky='w', padx=10, pady=6)

        def chosen():
            return [ln for ln in sorted(col.lines) if vars_.get(ln) and vars_[ln].get()]

        def update():
            ch = chosen()
            if not ch:
                result.set('Nothing ticked.')
            elif len(ch) == 1:
                result.set(f'Use: {ch[0]}')
            else:
                name = col.combined_name(ch)
                new = '' if name in col.lines else '  (new line - will be added to the Lines sheet)'
                result.set(f'Cross: {name}{new}\nAlleles: {", ".join(col.combined_loci(ch)) or "none"}')
        fill_boxes()
        update()

        @guard
        def use():
            ch = chosen()
            if not ch:
                return
            name = col.ensure_combined_line(ch)
            d.destroy()
            on_pick(name)

        def new_line():
            def done(name):
                vars_[name] = tk.BooleanVar(value=True)
                for k, v in vars_.items():
                    if k != name:
                        v.set(False)
                fill_boxes()
                update()
            new_line_form(d, done)
        bf = ttk.Frame(d)
        bf.grid(row=3, column=0, columnspan=2, sticky='e', padx=10, pady=10)
        ttk.Button(bf, text='Define a new line...', command=new_line).pack(side='left')
        ttk.Button(bf, text='Use', command=use).pack(side='left', padx=6)
        ttk.Button(bf, text='Cancel', command=d.destroy).pack(side='left')

    @guard
    def merge_line_dialog():
        v = selected(lines_tree)
        if not v:
            messagebox.showinfo('Merge', 'Select the line to merge away in the Lines table first.')
            return
        col = state['col']
        old = v[0]
        d = tk.Toplevel(root)
        d.title('Merge line')
        d.grab_set()
        ttk.Label(d, text=f'Merge "{old}" into:').grid(row=0, column=0, sticky='w', padx=10, pady=(10, 3))
        cb = ttk.Combobox(d, values=[x for x in sorted(col.lines) if x != old], width=30, state='readonly')
        cb.grid(row=0, column=1, padx=10, pady=(10, 3))
        ttk.Label(d, text='All breeders, pairs, cages, mice and history rows of the first line get the second '
                          'name, and the first line is removed from the Lines sheet.', foreground='#555',
                  wraplength=420).grid(row=1, column=0, columnspan=2, sticky='w', padx=10)

        @guard
        def ok():
            if not cb.get():
                return
            if messagebox.askyesno('Merge', f'Merge "{old}" into "{cb.get()}"?', parent=d):
                col.merge_lines(old, cb.get())
                d.destroy()
                refresh_all()
        ttk.Button(d, text='Merge', command=ok).grid(row=2, column=1, sticky='e', padx=10, pady=10)

    @guard
    def edit_line_alleles():
        v = selected(lines_tree)
        if not v:
            messagebox.showinfo('Line', 'Select a line in the Lines table first.')
            return
        col = state['col']
        ln = v[0]
        d = tk.Toplevel(root)
        d.title(f'Line {ln}')
        d.grab_set()
        ttk.Label(d, text='Alleles tracked (; separated)').grid(row=0, column=0, sticky='w', padx=10, pady=(10, 3))
        e = ttk.Entry(d, width=40)
        e.insert(0, '; '.join(col.lines[ln]['loci']))
        e.grid(row=0, column=1, padx=10, pady=(10, 3))
        ttk.Label(d, text='Renaming an allele does not rename genotypes already entered.', foreground='#555').grid(
            row=1, column=0, columnspan=2, sticky='w', padx=10)

        @guard
        def ok():
            col.update_line(ln, loci=[x.strip() for x in e.get().split(';') if x.strip()])
            d.destroy()
            refresh_all()
        ttk.Button(d, text='Save', command=ok).grid(row=2, column=1, sticky='e', padx=10, pady=10)

    def new_line_form(parent, on_done=None, name=''):
        col = state['col']
        f = tk.Toplevel(parent)
        f.title('New line')
        f.transient(parent)
        f.grab_set()
        ent = {}
        rows_ = [('Name (as used in the colony)', 'name', 30), ('Alleles tracked (; separated)', 'loci', 36),
                 ('Genetic background', 'bg', 24), ('Notes', 'notes', 36)]
        for i, (lab_, k, wdt) in enumerate(rows_):
            ttk.Label(f, text=lab_).grid(row=i, column=0, sticky='w', padx=10, pady=3)
            e = ttk.Entry(f, width=wdt)
            e.grid(row=i, column=1, sticky='w', padx=10)
            ent[k] = e
        ent['name'].insert(0, name)
        r = len(rows_)
        keep_v = tk.BooleanVar(value=False)
        ttk.Checkbutton(f, text='Must always keep a breeding pair of this line', variable=keep_v).grid(
            row=r, column=1, sticky='w', padx=10)
        r += 1
        ttk.Label(f, text='May be paired with').grid(row=r, column=0, sticky='nw', padx=10, pady=3)
        pf = ttk.Frame(f)
        pf.grid(row=r, column=1, sticky='w', padx=10)
        pvars = {}
        for i, ln in enumerate(sorted(col.lines)):
            pvars[ln] = tk.BooleanVar(value=False)
            ttk.Checkbutton(pf, text=ln, variable=pvars[ln]).grid(row=i // 2, column=i % 2, sticky='w',
                                                                  padx=(0, 10))
        r += 1
        ttk.Label(f, text='Alleles are the things you genotype, e.g. "Cx36-flox; PV-Cre". The line itself is '
                          'always allowed as a partner.', foreground='#555', wraplength=440).grid(
            row=r, column=0, columnspan=2, sticky='w', padx=10, pady=4)
        r += 1

        @guard
        def ok():
            nm = ent['name'].get().strip()
            loci = [x.strip() for x in ent['loci'].get().split(';') if x.strip()]
            col.add_line(nm, loci=loci, partners=[nm] + [k for k, v in pvars.items() if v.get()],
                         keep=keep_v.get(), background=ent['bg'].get().strip(), notes=ent['notes'].get().strip())
            f.destroy()
            refresh_all()
            if on_done:
                on_done(nm)
        ttk.Button(f, text='Add line', command=ok).grid(row=r, column=1, sticky='e', padx=10, pady=10)

    def gt_editor(parent, line, current):
        """One hom/het/wt/? picker per locus of the line. Returns (frame, rebuild(line), get())."""
        col = state['col']
        frame = ttk.Frame(parent)
        boxes = {}

        def build(ln):
            old = {k: gt_code(w.get()) for k, w in boxes.items()}
            for w in frame.winfo_children():
                w.destroy()
            boxes.clear()
            if ln == 'C57Bl6J':
                ttk.Label(frame, text='wild type for everything', foreground='#555').grid(row=0, column=0)
                return
            loci = col.loci_for(ln)
            loci += [k for k in current if k not in loci]
            if not loci:
                ttk.Label(frame, text='No alleles set for this line (Lines sheet, last column)',
                          foreground='#555').grid(row=0, column=0)
                return
            for i, lc in enumerate(loci):
                ttk.Label(frame, text=lc).grid(row=i // 2, column=(i % 2) * 2, sticky='w', padx=(0, 4))
                cb = ttk.Combobox(frame, values=[GT_LABEL[c] for c in GT_CODES], width=14, state='readonly')
                cb.set(GT_LABEL.get(old.get(lc) or current.get(lc) or '?', '? (not tested)'))
                cb.grid(row=i // 2, column=(i % 2) * 2 + 1, sticky='w', padx=(0, 14), pady=1)
                boxes[lc] = cb
        build(line)
        return frame, build, lambda: {k: gt_code(w.get()) for k, w in boxes.items()}

    def refresh_all_mice():
        col = state['col']
        if col is None:
            return
        rack, line, kind = am_rack.get(), am_line.get(), am_kind.get()
        q = am_search.get().strip().lower()
        rows_, per_line = [], {}
        am_rows.clear()
        n_cages = n_mice = n_m = n_f = n_pairs = 0
        for r in col.all_mice():
            if r['tag'] == 'empty' and not am_empty.get():
                continue
            if rack not in ('', '(all)') and r['rack'] != rack:
                continue
            if line not in ('', '(all)') and r['line'] != line:
                continue
            is_pair = r['pair'] is not None or r['sex'] == 'TP'
            if kind == 'Males' and not (r['males'] and not is_pair):
                continue
            if kind == 'Females' and not (r['females'] and not is_pair):
                continue
            if kind == 'Breeding pairs' and not is_pair:
                continue
            if kind == 'Breeder candidates' and r['tag'] != 'Candidate':
                continue
            if kind == 'Needs attention' and r['tag'] not in ('OVER AGE', 'Replace soon', 'Terminate', 'NO DATE') \
                    and 'related' not in r['status'] and 'not on Pairs' not in r['status']:
                continue
            parents = ' x '.join(x for x in (r['sire'], r['dam']) if x)
            n_todo_c = sum(1 for m in r['mice'] if col.needs_genotyping(m['id']))
            if kind == 'Needs genotyping' and not n_todo_c:
                continue
            status_txt = r['status'] + (f' | {n_todo_c} to genotype' if n_todo_c else '')
            vals = (r['card'], r['line'], 'pair' if is_pair else r['sex'], r['count'] or '',
                    fmt_date(r['birth']) if not is_pair else '', fmt_age(r['age']), status_txt, r['animals'],
                    parents, gt_pretty(r['gt']), r['note'])
            mrows = []
            for m in r['mice']:
                procs = col.procedures_for_mouse(m['id'])
                todo = col.needs_genotyping(m['id'])
                if kind == 'Needs genotyping' and not todo:
                    continue
                st = 'ex-breeder ' + m['breeder'] if m['breeder'] else 'alive'
                if todo:
                    st = 'NEEDS GENOTYPING'
                if procs:
                    st += f' | {len(procs)} procedure' + ('s' if len(procs) > 1 else '')
                mv = (m['id'], m['line'], m['sex'], '', fmt_date(m['birth']), fmt_age(col.age(m['birth'])), st,
                      f"ear {m['mark']}", ' x '.join(x for x in col.parents('mouse:' + m['id']) if x),
                      gt_pretty(col.mouse_gt_text(m['id'])), m['notes'])
                mrows.append((m['id'], mv, 'todo' if todo else 'mouse_proc' if procs else 'mouse'))
            if q:
                cage_hit = q in ' '.join(str(v) for v in vals).lower() + ' ' + r['cage'].lower()
                hits = [x for x in mrows if q in ' '.join(str(v) for v in x[1]).lower()]
                if not cage_hit and not hits:
                    continue
                if not cage_hit:
                    mrows = hits
            tag = r['tag']
            if 'related' in r['status']:
                tag = 'bad'
            rows_.append((vals, tag, r['cage'], mrows))
            if r['tag'] != 'empty':
                n_cages += 1
                n_mice += r['count']
                n_m += r['males']
                n_f += r['females']
                n_pairs += 1 if is_pair else 0
                pl = per_line.setdefault(r['line'] or '(no line)', [0, 0, 0, 0, 0, 0])
                pl[0] += 1
                pl[1] += r['count']
                pl[2] += r['males']
                pl[3] += r['females']
                pl[4] += r['count'] if is_pair else 0
                pl[5] += 1 if r['tag'] == 'Candidate' else 0
        am_tree.delete(*am_tree.get_children())
        for vals, tag, cage, mrows in rows_:
            iid = am_tree.insert('', 'end', values=vals, tags=(tag,) if tag else (), open=am_open.get() or bool(q))
            am_rows[iid] = ('cage', cage)
            for mid, mv, mtag in mrows:
                ciid = am_tree.insert(iid, 'end', values=mv, tags=(mtag,))
                am_rows[ciid] = ('mouse', mid)
        n_todo = sum(1 for m in col.mice if col.needs_genotyping(m))
        am_total.set(f"{n_mice} mice in {n_cages} cage{'' if n_cages == 1 else 's'}   ({n_m} ♂, {n_f} ♀; "
                     f"{n_pairs} breeding cage{'' if n_pairs == 1 else 's'})"
                     + (f'   |   {n_todo} mice need genotyping' if n_todo else ''))
        fill(am_lines_tree, [((ln_, *v), None) for ln_, v in sorted(per_line.items(), key=lambda x: -x[1][1])])

    def selected_cage():
        sel = am_tree.selection()
        if not sel or sel[0] not in am_rows:
            return None
        kind, key = am_rows[sel[0]]
        col = state['col']
        cage = key if kind == 'cage' else col.mice[key]['cage']
        return col.cage(cage)

    def selected_mice():
        """Mouse IDs selected in the All mice tree (only mouse rows)."""
        return [am_rows[i][1] for i in am_tree.selection() if i in am_rows and am_rows[i][0] == 'mouse']

    @guard
    def edit_selected_item():
        ms = selected_mice()
        if ms:
            mouse_dialog(ms[0])
        else:
            edit_selected_cage()

    @guard
    def edit_selected_cage():
        c = selected_cage()
        if not c:
            messagebox.showinfo('Edit cage', 'Select a cage first.')
            return
        col = state['col']
        if c['sex'] == 'TP' or any(p['cage'] == c['cage'] and p['status'] == 'Active' for p in col.pairs):
            messagebox.showinfo('Breeding cage', f'{col.card_name(c["cage"])} is a breeding cage. Edit its animals '
                                                 f'on the Breeders tab, or end the pair on the Overview tab.')
            return
        cage_dialog(c)

    @guard
    def empty_selected_cage():
        c = selected_cage()
        col = state['col']
        if not c:
            return
        if c['sex'] == 'TP' or any(p['cage'] == c['cage'] and p['status'] == 'Active' for p in col.pairs):
            messagebox.showinfo('Breeding cage', 'End the pair on the Overview tab instead.')
            return
        if messagebox.askyesno('Empty cage', f'Clear cage {col.card_name(c["cage"])} '
                                             f'({c["count"]} {c["sex"]}, {c["line"]})? Use this when the animals '
                                             f'were used or terminated.'):
            col.clear_cage(c['cage'])
            refresh_all()

    @guard
    def cage_dialog(c):
        col = state['col']
        new = c is None
        d = tk.Toplevel(root)
        d.title('Add animals to a cage' if new else f'Edit cage {col.card_name(c["cage"])}')
        d.grab_set()
        ids = [''] + sorted(col.breeders)
        widgets = {}
        fields = [('Cage', 'cage'), ('Line', 'line'), ('Sex', 'sex'),
                  ('Number of mice' if new else f'Add mice (has {len(col.mice_in(c["cage"]))})', 'count'),
                  ('Born (yyyy-mm-dd)', 'birth'), ('Sire ID', 'sire'), ('Dam ID', 'dam'), ('Note', 'note')]
        for i, (lab, k) in enumerate(fields):
            ttk.Label(d, text=lab).grid(row=i, column=0, sticky='w', padx=10, pady=3)
            if k == 'cage':
                w = ttk.Combobox(d, width=12, values=[col.card_name(x) for x in col.empty_cages()],
                                 state='readonly' if new else 'disabled')
                w.set(col.card_name(c['cage']) if c else '')
            elif k == 'line':
                w = ttk.Combobox(d, width=32)
                w.set(c['line'] if c else '')
                enhance_line_combo(w)
            elif k == 'sex':
                w = ttk.Combobox(d, width=6, values=['M', 'F'], state='readonly')
                w.set(c['sex'] if c else '')
            elif k in ('sire', 'dam'):
                w = ttk.Combobox(d, width=10, values=ids)
                w.set(c[k] if c else '')
            else:
                w = ttk.Entry(d, width=40 if k == 'note' else 14)
                v = (fmt_date(c['birth']) if k == 'birth' else c[k]) if c else ''
                if k == 'count' and not new:
                    v = 0
                w.insert(0, '' if v is None else str(v))
            w.grid(row=i, column=1, sticky='w', padx=10)
            widgets[k] = w
        ttk.Label(d, text='Genotype of\nnew mice').grid(row=len(fields), column=0, sticky='nw', padx=10, pady=3)
        gfr, gbuild, gget = gt_editor(d, c['line'] if c else '', c['gt'] if c else {})
        gfr.grid(row=len(fields), column=1, sticky='w', padx=10, pady=3)
        exp = col.expected_gt_text(c['sire'], c['dam']) if c else ''
        ttk.Label(d, text=(f'From the parents: {exp[5:]}. ' if exp else '') + GT_HELP +
                  '. Each mouse gets its own ID and the next free ear mark; genotype individual mice with '
                  '"Genotyping results".', foreground='#555',
                  wraplength=360).grid(row=len(fields) + 1, column=1, sticky='w', padx=10)
        widgets['line'].bind('<<ComboboxSelected>>', lambda e: gbuild(widgets['line'].get()), add='+')

        @guard
        def ok():
            g = {k: widgets[k].get().strip() for k, _ in [(f[1], 0) for f in fields]}
            if new:
                cage = next((x for x in col.empty_cages() if col.card_name(x) == g['cage']), None)
                if not cage:
                    messagebox.showerror('Cage', 'Pick an empty cage.', parent=d)
                    return
            else:
                cage = c['cage']
            if g['sex'] not in ('M', 'F'):
                messagebox.showerror('Sex', 'Sex must be M or F (breeding cages are made with "Create pair").',
                                     parent=d)
                return
            try:
                cnt = int(g['count'] or 0)
                assert cnt > 0 or (not new and cnt == 0)
            except (ValueError, AssertionError):
                messagebox.showerror('Number', 'Enter a whole number of mice (use "Terminate / remove" to take '
                                               'mice out).', parent=d)
                return
            try:
                bd = parse_date(g['birth'])
            except ValueError:
                messagebox.showerror('Date', 'Birth date must look like 2026-09-21.', parent=d)
                return
            for k in ('sire', 'dam'):
                if g[k] and g[k] not in col.breeders:
                    messagebox.showerror('Parent', f'{g[k]} is not a breeder ID.', parent=d)
                    return
            col.save_cage(cage, g['line'], g['sex'], cnt, bd if bd else '', g['sire'], g['dam'], g['note'],
                          gt=gget())
            d.destroy()
            refresh_all()
        ttk.Button(d, text='Save', command=ok).grid(row=len(fields) + 2, column=1, sticky='e', padx=10, pady=10)

    # ---------- procedures / termination / history ----------
    def cage_label(cage):
        col = state['col']
        c = col.cage(cage)
        p = col.active_pair(cage)
        if p:
            return f"{col.card_name(cage)}  -  breeding cage, {p['line']}: {', '.join(col.pair_members(p))}"
        return f"{col.card_name(cage)}  -  {c['count'] or 0} {c['sex']}, {c['line']}, born {fmt_date(c['birth'])}"

    def person_combo(parent):
        col = state['col']
        names = col.people()
        me = state.get('user') or ''
        if me and me not in names:
            names = [me] + names
        cb = ttk.Combobox(parent, values=names, width=20)
        cb.set(me)
        return cb

    def mouse_checklist(parent, cage, preselect=()):
        """Checkbox per mouse of a stock cage. Returns (frame, get_selected_ids)."""
        col = state['col']
        fr = ttk.Frame(parent)
        vars_ = {}
        mice = col.mice_in(cage)
        for i, m in enumerate(mice):
            v = tk.BooleanVar(value=(m['id'] in preselect) if preselect else True)
            ttk.Checkbutton(fr, variable=v, text=f"{m['id']}  ear {m['mark']:<5} {m['sex']}  "
                                                 f"{gt_pretty(col.mouse_gt_text(m['id']))}"
                                                 + ('  (needs genotyping)' if col.needs_genotyping(m['id']) else '')).grid(row=i // 2, column=i % 2,
                                                                                       sticky='w', padx=(0, 16))
            vars_[m['id']] = v
        btns = ttk.Frame(fr)
        btns.grid(row=(len(mice) + 1) // 2, column=0, columnspan=2, sticky='w', pady=(2, 0))
        ttk.Button(btns, text='All', width=5, command=lambda: [v.set(True) for v in vars_.values()]).pack(
            side='left')
        ttk.Button(btns, text='None', width=5, command=lambda: [v.set(False) for v in vars_.values()]).pack(
            side='left', padx=4)
        return fr, lambda: [k for k, v in vars_.items() if v.get()]

    def procedure_dialog(cage=None, animal=None, preselect=()):
        col = state['col']
        d = tk.Toplevel(root)
        d.title('Log procedure')
        d.grab_set()
        if animal:
            b = col.breeders[animal]
            target = f"{animal}  -  {b['sex']}, {b['line']}, cage {col.card_name(b['cage'])}"
        else:
            target = cage_label(cage)
        ttk.Label(d, text=target, font=('TkDefaultFont', 10, 'bold')).grid(row=0, column=0, columnspan=2,
                                                                           sticky='w', padx=10, pady=(10, 6))
        r = 1
        pair = col.active_pair(cage) if cage else None
        who = None
        if pair:
            ttk.Label(d, text='Done to').grid(row=r, column=0, sticky='w', padx=10, pady=3)
            who = ttk.Combobox(d, state='readonly', width=24,
                               values=['(whole breeding cage)'] + col.pair_members(pair))
            who.set('(whole breeding cage)')
            who.grid(row=r, column=1, sticky='w', padx=10)
            r += 1
        n_w = get_mice = None
        if cage and not pair:
            c = col.cage(cage)
            if col.mice_in(cage):
                ttk.Label(d, text='Mice').grid(row=r, column=0, sticky='nw', padx=10, pady=3)
                fr, get_mice = mouse_checklist(d, cage, preselect)
                fr.grid(row=r, column=1, sticky='w', padx=10, pady=3)
            else:
                ttk.Label(d, text='Number of mice').grid(row=r, column=0, sticky='w', padx=10, pady=3)
                n_w = ttk.Spinbox(d, from_=1, to=max(1, c['count'] or 1), width=5)
                n_w.set(c['count'] or 1)
                n_w.grid(row=r, column=1, sticky='w', padx=10)
            r += 1
        ent = {}
        for lab, k in (('Date (yyyy-mm-dd)', 'date'), ('Procedure', 'proc'), ('Details (virus, coordinates...)',
                                                                                'details'),
                       ('Person', 'person'), ('Notes', 'notes')):
            ttk.Label(d, text=lab).grid(row=r, column=0, sticky='w', padx=10, pady=3)
            if k == 'proc':
                w = ttk.Combobox(d, values=col.procedure_types(), width=30)
            elif k == 'person':
                w = person_combo(d)
            else:
                w = ttk.Entry(d, width=44 if k in ('details', 'notes') else 14)
            if k == 'date':
                w.insert(0, dt.date.today().isoformat())
            w.grid(row=r, column=1, sticky='w', padx=10)
            ent[k] = w
            r += 1

        @guard
        def ok():
            g = {k: w.get().strip() for k, w in ent.items()}
            if not g['proc']:
                messagebox.showerror('Procedure', 'Enter what was done.', parent=d)
                return
            try:
                day = parse_date(g['date'])
            except ValueError:
                messagebox.showerror('Date', 'Date must look like 2026-10-25.', parent=d)
                return
            day = day.date() if day else dt.date.today()
            an = animal
            if who is not None and who.get() != '(whole breeding cage)':
                an = who.get()
            if get_mice:
                chosen = get_mice()
                if not chosen:
                    messagebox.showerror('Mice', 'Tick at least one mouse.', parent=d)
                    return
                col.log_procedure(day, g['proc'], g['details'], g['person'], g['notes'], mice=chosen)
            else:
                col.log_procedure(day, g['proc'], g['details'], g['person'], g['notes'],
                                  cage=None if an else cage, animal=an, n=int(n_w.get()) if n_w else None)
            d.destroy()
            refresh_all()
        ttk.Button(d, text='Save', command=ok).grid(row=r, column=1, sticky='e', padx=10, pady=10)

    def terminate_dialog(cage=None, ids=None, preselect=()):
        """Stock cage (tick the mice) or breeders (pair members / selected breeders)."""
        col = state['col']
        d = tk.Toplevel(root)
        d.title('Terminate / remove animals')
        d.grab_set()
        pair = col.active_pair(cage) if cage else None
        if pair:
            ids = col.pair_members(pair)
        ttk.Label(d, text=cage_label(cage) if cage else ', '.join(ids),
                  font=('TkDefaultFont', 10, 'bold')).grid(row=0, column=0, columnspan=2, sticky='w', padx=10,
                                                           pady=(10, 6))
        r = 1
        n_w, checks, get_mice = None, {}, None
        if ids:
            ttk.Label(d, text='Animals').grid(row=r, column=0, sticky='nw', padx=10, pady=3)
            fr = ttk.Frame(d)
            fr.grid(row=r, column=1, sticky='w', padx=10)
            for i, x in enumerate(ids):
                b = col.breeders[x]
                v = tk.BooleanVar(value=True)
                ttk.Checkbutton(fr, variable=v, text=f"{x}  {b['sex']}  {b['line']}  "
                                                    f"{fmt_age(col.age(b['birth']))}").grid(row=i, column=0,
                                                                                            sticky='w')
                checks[x] = v
            r += 1
        elif col.mice_in(cage):
            ttk.Label(d, text='Mice').grid(row=r, column=0, sticky='nw', padx=10, pady=3)
            fr, get_mice = mouse_checklist(d, cage, preselect)
            fr.grid(row=r, column=1, sticky='w', padx=10, pady=3)
            r += 1
        else:
            c = col.cage(cage)
            ttk.Label(d, text='How many').grid(row=r, column=0, sticky='w', padx=10, pady=3)
            n_w = ttk.Spinbox(d, from_=1, to=c['count'], width=5)
            n_w.set(c['count'])
            n_w.grid(row=r, column=1, sticky='w', padx=10)
            ttk.Label(d, text=f'of {c["count"]} (all = cage becomes empty)', foreground='#555').grid(
                row=r, column=1, sticky='w', padx=70)
            r += 1
        ent = {}
        for lab, k in (('Date (yyyy-mm-dd)', 'date'), ('Outcome', 'outcome'), ('Reason / use', 'reason'),
                       ('Project', 'project'), ('Person', 'person'), ('Notes', 'notes')):
            ttk.Label(d, text=lab).grid(row=r, column=0, sticky='w', padx=10, pady=3)
            if k == 'outcome':
                w = ttk.Combobox(d, values=OUTCOMES, width=30)
                w.set(OUTCOMES[0])
            elif k == 'person':
                w = person_combo(d)
            elif k == 'project':
                w = ttk.Combobox(d, values=sorted({t['project'] for t in col.terminated if t['project']}),
                                 width=20)
            else:
                w = ttk.Entry(d, width=44 if k in ('reason', 'notes') else 14)
            if k == 'date':
                w.insert(0, dt.date.today().isoformat())
            w.grid(row=r, column=1, sticky='w', padx=10)
            ent[k] = w
            r += 1
        procs = []
        if ids:
            for x in ids:
                procs += [p for p in col.procedures_for(animal=x) if p not in procs]
        elif get_mice:
            for m in col.mice_in(cage):
                procs += [p for p in col.procedures_for_mouse(m['id']) if p not in procs]
        else:
            procs = col.procedures_for(cage=cage, birth=col.cage(cage)['birth'])
        ttk.Label(d, text='Procedures that go into the history (per mouse):\n' + (col.proc_summary(procs).replace('; ', '\n')
                                                                     or 'none logged'),
                  foreground='#555', justify='left').grid(row=r, column=0, columnspan=2, sticky='w', padx=10,
                                                          pady=6)
        r += 1

        @guard
        def ok():
            g = {k: w.get().strip() for k, w in ent.items()}
            try:
                day = parse_date(g['date'])
            except ValueError:
                messagebox.showerror('Date', 'Date must look like 2026-10-25.', parent=d)
                return
            day = day.date() if day else dt.date.today()
            if not g['outcome']:
                messagebox.showerror('Outcome', 'Pick an outcome.', parent=d)
                return
            kw = dict(reason=g['reason'], project=g['project'], person=g['person'], notes=g['notes'])
            if ids:
                chosen = [x for x, v in checks.items() if v.get()]
                if not chosen:
                    return
                if not messagebox.askyesno('Confirm', f"Record {len(chosen)} animal(s) as '{g['outcome']}' on "
                                                      f"{day}? This moves them to the history.", parent=d):
                    return
                col.terminate_breeders(chosen, day, g['outcome'], **kw)
            elif get_mice:
                chosen = get_mice()
                if not chosen:
                    return
                marks = ', '.join(f"{x} ({col.mice[x]['mark']})" for x in chosen)
                if not messagebox.askyesno('Confirm', f"Record {marks} from {col.card_name(cage)} as "
                                                      f"'{g['outcome']}' on {day}?", parent=d):
                    return
                col.terminate_mice(chosen, day, g['outcome'], **kw)
            else:
                n = int(n_w.get())
                if not messagebox.askyesno('Confirm', f"Record {n} mice from {col.card_name(cage)} as "
                                                      f"'{g['outcome']}' on {day}?", parent=d):
                    return
                col.terminate_cage(cage, n, day, g['outcome'], **kw)
            d.destroy()
            refresh_all()
        ttk.Button(d, text='Record', command=ok).grid(row=r, column=1, sticky='e', padx=10, pady=10)

    def _sel_cage_or_msg():
        c = selected_cage()
        if not c or (not c['count'] and not state['col'].active_pair(c['cage'])):
            messagebox.showinfo('Cage', 'Select a cage with animals first.')
            return None
        return c

    @guard
    def proc_for_selected_cage():
        c = _sel_cage_or_msg()
        if c:
            procedure_dialog(cage=c['cage'], preselect=selected_mice())

    @guard
    def term_for_selected_cage():
        c = _sel_cage_or_msg()
        if c:
            terminate_dialog(cage=c['cage'], preselect=selected_mice())

    @guard
    def genotyping_for_selected():
        c = _sel_cage_or_msg()
        if not c:
            return
        col = state['col']
        if col.active_pair(c['cage']):
            messagebox.showinfo('Genotyping', 'For breeders, set the genotype in the breeder dialog (Breeders tab).')
            return
        genotyping_dialog(c['cage'], selected_mice())

    def genotyping_dialog(cage, preselect=()):
        """Grid: one row per mouse, one hom/het/wt/? picker per allele."""
        col = state['col']
        mice = col.mice_in(cage)
        if not mice:
            messagebox.showinfo('Genotyping', 'This cage has no mouse IDs.')
            return
        if preselect:
            mice = [m for m in mice if m['id'] in preselect] or mice
        elif any(col.needs_genotyping(m['id']) for m in mice):
            mice = [m for m in mice if col.needs_genotyping(m['id'])]
        loci = col.loci_for(mice[0]['line'])
        if not loci:
            messagebox.showinfo('Genotyping', f"No alleles are set for {mice[0]['line']} (Lines sheet, last column).")
            return
        d = tk.Toplevel(root)
        d.title(f'Genotyping results - {col.card_name(cage)}')
        d.grab_set()
        ttk.Label(d, text=cage_label(cage), font=('TkDefaultFont', 10, 'bold')).grid(
            row=0, column=0, columnspan=3 + len(loci), sticky='w', padx=10, pady=(10, 6))
        for j, h in enumerate(['Mouse', 'Ear', 'Sex'] + loci):
            ttk.Label(d, text=h, font=('TkDefaultFont', 9, 'bold')).grid(row=1, column=j, sticky='w', padx=6)
        boxes = {}
        for i, m in enumerate(mice, 2):
            for j, t in enumerate((m['id'], m['mark'], m['sex'])):
                ttk.Label(d, text=t).grid(row=i, column=j, sticky='w', padx=6, pady=1)
            boxes[m['id']] = {}
            exp = col.mouse_gt(m['id'])
            for j, lc in enumerate(loci, 3):
                cb = ttk.Combobox(d, values=[GT_LABEL[c] for c in GT_CODES if c != 'mixed'], width=14,
                                  state='readonly')
                cb.set(GT_LABEL.get(m['gt'].get(lc, '?'), '? (not tested)'))
                cb.grid(row=i, column=j, sticky='w', padx=6, pady=1)
                boxes[m['id']][lc] = cb
            hint = ', '.join(f'{k} {GT_SHORT.get(v.rstrip("*"), v)}' for k, v in exp.items() if v.endswith('*'))
            if hint:
                ttk.Label(d, text=f'(certain from parents: {hint})', foreground='#888').grid(
                    row=i, column=3 + len(loci), sticky='w')
            elif col.needs_genotyping(m['id']):
                ttk.Label(d, text='needs genotyping', foreground='#B26B00').grid(row=i, column=3 + len(loci),
                                                                                 sticky='w')
        r = len(mice) + 2
        ent = {}
        for lab, k in (('Date', 'date'), ('Person', 'person'), ('Notes', 'notes')):
            ttk.Label(d, text=lab).grid(row=r, column=0, sticky='w', padx=10, pady=3)
            w = person_combo(d) if k == 'person' else \
                ttk.Entry(d, width=36 if k == 'notes' else 12)
            if k == 'date':
                w.insert(0, dt.date.today().isoformat())
            w.grid(row=r, column=1, columnspan=3, sticky='w', padx=6)
            ent[k] = w
            r += 1
        log_v = tk.BooleanVar(value=True)
        ttk.Checkbutton(d, text='Also log a "Genotyping" procedure for these mice', variable=log_v).grid(
            row=r, column=0, columnspan=4, sticky='w', padx=10)
        r += 1

        @guard
        def ok():
            try:
                day = parse_date(ent['date'].get())
            except ValueError:
                messagebox.showerror('Date', 'Date must look like 2026-10-25.', parent=d)
                return
            res = {mid: {lc: gt_code(cb.get()) for lc, cb in bx.items()} for mid, bx in boxes.items()}
            col.record_genotypes(res, day.date() if day else dt.date.today(), ent['person'].get().strip(),
                                 ent['notes'].get().strip(), log=log_v.get())
            d.destroy()
            refresh_all()
        ttk.Button(d, text='Save results', command=ok).grid(row=r, column=3, columnspan=3, sticky='e', padx=10,
                                                            pady=10)

    def mouse_dialog(mid):
        col = state['col']
        m = col.mice[mid]
        d = tk.Toplevel(root)
        d.title(f'Mouse {mid}')
        d.grab_set()
        info = (f"{mid}  -  {m['line']}, born {fmt_date(m['birth'])} ({fmt_age(col.age(m['birth']))}), "
                f"parents {' x '.join(x for x in col.parents('mouse:' + mid) if x) or 'unknown'}"
                + (f", ex-breeder {m['breeder']}" if m['breeder'] else ''))
        ttk.Label(d, text=info, font=('TkDefaultFont', 10, 'bold')).grid(row=0, column=0, columnspan=2, sticky='w',
                                                                         padx=10, pady=(10, 6))
        ttk.Label(d, text='Cage').grid(row=1, column=0, sticky='w', padx=10, pady=3)
        stock = [c['cage'] for c in col.cages if not col.active_pair(c['cage']) and
                 (c['cage'] == m['cage'] or not c['line'] or (c['line'] == m['line'] and c['sex'] == m['sex']))]
        cage_cb = ttk.Combobox(d, values=[col.card_name(x) for x in stock], width=12, state='readonly')
        cage_cb.set(col.card_name(m['cage']))
        cage_cb.grid(row=1, column=1, sticky='w', padx=10)
        ttk.Label(d, text='Ear mark').grid(row=2, column=0, sticky='w', padx=10, pady=3)
        mark_cb = ttk.Combobox(d, values=EAR_MARKS, width=8)
        mark_cb.set(m['mark'])
        mark_cb.grid(row=2, column=1, sticky='w', padx=10)
        ttk.Label(d, text='Sex').grid(row=3, column=0, sticky='w', padx=10, pady=3)
        sex_cb = ttk.Combobox(d, values=['M', 'F'], width=4, state='readonly')
        sex_cb.set(m['sex'])
        sex_cb.grid(row=3, column=1, sticky='w', padx=10)
        ttk.Label(d, text='Genotype').grid(row=4, column=0, sticky='nw', padx=10, pady=3)
        gfr, _, gget = gt_editor(d, m['line'], m['gt'])
        gfr.grid(row=4, column=1, sticky='w', padx=10)
        todo_v = tk.BooleanVar(value=bool(m.get('todo')))
        ttk.Checkbutton(d, text='Needs genotyping (cleared automatically when all alleles have a result)',
                        variable=todo_v).grid(row=8, column=1, sticky='w', padx=10)
        ttk.Label(d, text='Notes').grid(row=5, column=0, sticky='w', padx=10, pady=3)
        notes_e = ttk.Entry(d, width=46)
        notes_e.insert(0, m['notes'])
        notes_e.grid(row=5, column=1, sticky='w', padx=10)
        procs = col.procedures_for_mouse(mid)
        ttk.Label(d, text='Procedures:\n' + (col.proc_summary(procs).replace('; ', '\n') or 'none'),
                  foreground='#555', justify='left').grid(row=6, column=0, columnspan=2, sticky='w', padx=10,
                                                          pady=6)

        @guard
        def ok():
            new_cage = next(x for x in stock if col.card_name(x) == cage_cb.get())
            others = {x['mark'] for x in col.mice_in(new_cage) if x['id'] != mid}
            mark = mark_cb.get().strip() or 'none'
            if new_cage == m['cage'] and mark in others:
                messagebox.showerror('Ear mark', f'Another mouse in this cage already has ear mark {mark}.',
                                     parent=d)
                return
            col.save_mouse(mid, mark=mark, sex=sex_cb.get(), gt=gget(), notes=notes_e.get().strip(),
                           cage=new_cage, todo=todo_v.get())
            d.destroy()
            refresh_all()
        ttk.Button(d, text='Save', command=ok).grid(row=9, column=1, sticky='e', padx=10, pady=10)

    @guard
    def history_of_selected_cage():
        c = selected_cage()
        if not c:
            return
        clear_history_filters(refresh=False)
        h_search.set(c['cage'])
        nb.select(hi)

    @guard
    def proc_for_selected_breeder():
        v = selected(br_tree)
        if v:
            procedure_dialog(animal=v[0])

    @guard
    def term_for_selected_breeder():
        sel = [br_tree.item(i, 'values')[0] for i in br_tree.selection()]
        col = state['col']
        sel = [x for x in sel if col.breeders[x]['status'] != 'Terminated']
        if sel:
            terminate_dialog(ids=sel)

    def clear_history_filters(refresh=True):
        h_line.set('(all)')
        h_from.delete(0, 'end')
        h_to.delete(0, 'end')
        h_search.set('')
        if refresh:
            refresh_history()

    def refresh_history():
        col = state['col']
        if col is None:
            return
        line = h_line.get()
        try:
            d0 = parse_date(h_from.get()) if h_from.get().strip() else None
            d1 = parse_date(h_to.get()) if h_to.get().strip() else None
        except ValueError:
            h_total.set('Dates must look like 2026-01-01')
            return
        d0, d1 = (d0.date() if d0 else None), (d1.date() if d1 else None)
        q = h_search.get().strip().lower()

        def keep(x, vals):
            if line not in ('', '(all)') and x['line'] != line:
                return False
            if d0 and (not x['date'] or x['date'] < d0):
                return False
            if d1 and (not x['date'] or x['date'] > d1):
                return False
            hay = ' '.join(str(v) for v in vals).lower() + ' ' + x['cage'].lower()
            return not q or q in hay
        trows, by_outcome, n_mice = [], {}, 0
        for t in sorted(col.terminated, key=lambda x: x['date'] or dt.date.min, reverse=True):
            vals = (fmt_date(t['date']), col.card_name(t['cage']), t['animal'], t['line'], t['sex'], t['n'],
                    fmt_age(t['age']) if isinstance(t['age'], (int, float)) else '', t['outcome'], t['reason'],
                    t['project'], t['person'], t['gt'], t['procs'])
            if keep(t, vals):
                trows.append((vals, 'Terminated'))
                n_mice += t['n'] or 0
                by_outcome[t['outcome']] = by_outcome.get(t['outcome'], 0) + (t['n'] or 0)
        fill(term_tree, trows)
        prows = []
        proc_rows.clear()
        alive_cohorts = {(c['cage'], c['birth']) for c in col.cages if c['count']}
        for p in sorted(col.procedures, key=lambda x: x['date'] or dt.date.min, reverse=True):
            if p['animal']:
                alive = 'yes' if col.breeders.get(p['animal'], {}).get('status') != 'Terminated' else 'no'
            elif p['sex'] == 'pair':
                alive = 'yes' if col.active_pair(p['cage']) else '-'
            else:
                alive = 'yes' if (p['cage'], p['birth']) in alive_cohorts else 'no'
            vals = (fmt_date(p['date']), col.card_name(p['cage']), p['animal'], p['line'], p['n'] or '', p['proc'],
                    p['details'], p['person'], p['notes'], alive)
            if keep(p, vals):
                prows.append((vals, 'proc', p['row']))
        fill(proc_tree, [(v, t) for v, t, _ in prows])
        for iid, (_, _, row) in zip(proc_tree.get_children(), prows):
            proc_rows[iid] = row
        h_total.set(f'{n_mice} mice in {len(trows)} termination records' +
                    (('   (' + ', '.join(f'{k}: {v}' for k, v in sorted(by_outcome.items(), key=lambda x: -x[1]))
                      + ')') if by_outcome else '') + f'   |   {len(prows)} procedures')

    @guard
    def delete_selected_proc():
        sel = proc_tree.selection()
        if not sel or sel[0] not in proc_rows:
            return
        v = proc_tree.item(sel[0], 'values')
        if not messagebox.askyesno('Delete', f'Delete procedure "{v[5]}" on {v[1]} {v[2]} ({v[0]})?'):
            return
        col = state['col']
        wb = col._open_for_write()
        wb['Procedures'].delete_rows(proc_rows[sel[0]])
        col._save(wb)
        refresh_all()

    def refresh_breeders():
        col = state['col']
        lf_, sf_ = b_line.get(), b_status.get()
        rows_ = []
        for b in sorted(col.breeders.values(), key=lambda x: x['id']):
            if lf_ not in ('', '(all)') and b['line'] != lf_:
                continue
            if sf_ == '(alive)':
                if b['status'] == 'Terminated':
                    continue
            elif sf_ not in ('', '(all)') and b['status'] != sf_:
                continue
            flag = col.age_flag(b['birth'], b['status'])
            fco = col.inbreeding(b['id'])
            gtxt = 'wt' if b['line'] == 'C57Bl6J' else gt_pretty(format_gt(b.get('gt') or {}))
            rows_.append(((b['id'], b['line'], b['sex'], fmt_date(b['birth']), fmt_age(col.age(b['birth'])), flag,
                           b['cage'], b['status'], b['sire'], b['dam'], f'{fco:.3f}' if fco else '', gtxt,
                           b['notes']),
                          b['status'] if b['status'] in GONE else flag))
        fill(br_tree, rows_)

    def refresh_pc_lists():
        col = state['col']
        if col is None:
            return
        if pc_all.get():
            alive = [m for m in col.mice.values() if m['status'] == 'Alive' and not m['breeder']]
            ms = [b['id'] for b in col.breeders.values() if b['sex'] == 'M' and b['status'] != 'Terminated'] + \
                 ['mouse:' + m['id'] for m in alive if m['sex'] == 'M']
            fs = [b['id'] for b in col.breeders.values() if b['sex'] == 'F' and b['status'] != 'Terminated'] + \
                 ['mouse:' + m['id'] for m in alive if m['sex'] == 'F']
        else:
            ms, fs = col.candidate_units('M', include_breeding=True), col.candidate_units('F', include_breeding=True)
        pc_m['values'] = [unit_label(u) for u in ms]
        pc_f1['values'] = [unit_label(u) for u in fs]
        pc_f2['values'] = [''] + [unit_label(u) for u in fs]

    def refresh_due():
        col = state['col']
        if col is None:
            return
        try:
            h = int(du_days.get())
        except ValueError:
            h = 90
        rows_ = []
        for p in col.due(h):
            partners = col.partner_lines(p['line'])
            cm = col.candidate_units('M', partners)
            cf_ = col.candidate_units('F', partners)
            def nice(u):
                if u.startswith('mouse:'):
                    m = col.mice[u[6:]]
                    return f"{col.card_name(m['cage'])} {m['mark']}"
                return u.replace('cage:', '')
            short = lambda us: ', '.join(nice(u) for u in us[:6]) + (f' ... (+{len(us) - 6})' if len(us) > 6
                                                                       else '') if us else 'NONE - set animals aside'
            rows_.append(((p['cage'], p['line'], ', '.join(col.pair_members(p)), fmt_age(col.pair_oldest(p)),
                           fmt_date(col.pair_replace_by(p)), short(cm), short(cf_)), col.pair_flag(p)))
        fill(due_tree, rows_)

    # ---------- actions ----------
    def open_file(p=None):
        p = p or filedialog.askopenfilename(filetypes=[('Excel', '*.xlsx')])
        if not p:
            return
        cfg = load_config()
        try:
            state['col'] = Colony(p, backup_dir=cfg.get('backup_dir') or None, user=cfg.get('user_name', ''))
        except Exception as e:
            messagebox.showerror('Could not open', f'{p}\n\n{e}')
            return
        cfg['last_workbook'] = os.path.abspath(p)
        cfg['mode'] = 'local'
        save_config(cfg)
        state['user'] = cfg.get('user_name', '')
        set_status()
        refresh_all()

    def google_token_for(client, parent=None, force=False):
        """The saved Google sign-in for this login ID, or sign in now (opens the browser)."""
        cfg = load_config()
        tok = cfg.get('google_token') or {}
        cid = client['installed']['client_id']
        if tok and tok.get('client_id') == cid and not force:
            return tok
        w = tk.Toplevel(parent or root)
        w.title('Google sign-in')
        ttk.Label(w, text='A browser window opened: sign in with the Google account that the sheet is shared '
                          'with and click Allow.\n\nIf Google says "Google hasn\'t verified this app": click '
                          'Advanced -> Go to Mouse Colony (it is your lab\'s own app).\n\nWaiting up to 5 '
                          'minutes...', wraplength=420, justify='left').pack(padx=16, pady=16)
        w.update()
        root.config(cursor='watch')
        root.update()
        try:
            tok = google_sign_in(client)
        finally:
            root.config(cursor='')
            w.destroy()
        cfg = load_config()
        cfg['google_token'] = tok
        if not cfg.get('user_name') and tok.get('account'):
            cfg['user_name'] = tok['account'].split('@')[0]
        save_config(cfg)
        return tok

    def make_backend(conn, parent=None, sign_in=True):
        if conn.get('key'):
            return GoogleBackend(conn['key'], conn['sheet_id'], conn.get('title', ''))
        cfg = load_config()
        tok = cfg.get('google_token') or {}
        if not tok or tok.get('client_id') != conn['client']['installed']['client_id']:
            tok = google_token_for(conn['client'], parent) if sign_in else None
        return GoogleBackend(None, conn['sheet_id'], conn.get('title', ''), token=tok)

    def run_busy(fn, msg=''):
        """Run a slow network call in a background thread while the window stays responsive."""
        import threading
        box = {}

        def work():
            try:
                box['result'] = fn()
            except BaseException as e:      # handed back to the caller below
                box['error'] = e
        t = threading.Thread(target=work, daemon=True)
        w = None
        try:
            w = tk.Toplevel(root)
            w.title('Please wait')
            w.resizable(False, False)
            ttk.Label(w, text=msg or 'Working...', padding=(18, 14)).pack()
            pb = ttk.Progressbar(w, mode='indeterminate', length=260)
            pb.pack(padx=18, pady=(0, 16))
            pb.start(12)
            w.transient(root)
            w.update_idletasks()
            try:
                w.grab_set()
            except tk.TclError:
                pass
        except tk.TclError:
            w = None
        t.start()
        import time
        while t.is_alive():
            try:
                root.update()
            except tk.TclError:
                break
            time.sleep(0.03)
        t.join()
        if w is not None:
            try:
                w.grab_release()
                w.destroy()
            except tk.TclError:
                pass
        if 'error' in box:
            raise box['error']
        return box.get('result')

    def open_shared(conn, quiet=False):
        """conn: {'sheet_id', 'title', 'client': Google login ID} or {..., 'key': service account}.
        Returns True when open (maybe read-only)."""
        cfg = load_config()
        try:
            remote = make_backend(conn)
            col = Colony(remote=remote, backup_dir=cfg.get('backup_dir') or None, user=cfg.get('user_name', ''),
                         io_hook=run_busy)
        except Exception as e:
            if not quiet:
                messagebox.showerror('Shared sheet', f'Could not open the shared sheet:\n\n{e}')
            return False
        state['col'] = col
        state['user'] = load_config().get('user_name', '')
        col.user = state['user']
        col.on_change = lambda: (set_status(), schedule_sync(5000))
        set_status()
        refresh_all()
        if col.pending:
            schedule_sync(1500)              # changes left from last time (or made offline)
        return True

    def schedule_sync(delay=5000):
        """Upload a few seconds after the last change (changes made meanwhile go in the same upload)."""
        if state.get('closing'):
            return
        if state.get('sync_timer'):
            try:
                root.after_cancel(state['sync_timer'])
            except tk.TclError:
                pass
        state['sync_timer'] = root.after(delay, lambda: start_sync())

    def start_sync(refresh_only_if_idle=False):
        """Background upload + refresh of the shared sheet; the window stays usable."""
        import threading
        state['sync_timer'] = None
        col = state['col']
        if col is None or not col.remote or state.get('closing'):
            return
        if state.get('syncing'):
            schedule_sync(3000)
            return
        if refresh_only_if_idle and not col.pending and any(isinstance(w, tk.Toplevel)
                                                            for w in root.winfo_children()):
            return                               # don't change the data under an open dialog
        job = col.prepare_sync()
        box = {}

        def work():
            try:
                box['res'] = col.do_sync(job)
            except BaseException as e:
                box['err'] = e
        th = threading.Thread(target=work, daemon=True)
        state['syncing'] = th
        set_status()
        th.start()

        def poll():
            if th.is_alive():
                root.after(200, poll)
                return
            state['syncing'] = None
            if 'err' in box:
                col.sync_failed(box['err'])
                set_status()
                if isinstance(box['err'], ConflictError):
                    conflict_dialog(box['err'])
                else:
                    log_error(f"sync: {box['err'].__class__.__name__}: {box['err']}")
                    if col.pending:
                        schedule_sync(60 * 1000)     # offline: try again in a minute
                return
            try:
                changed = col.finish_sync(box['res'])
            except Exception:
                log_error(traceback.format_exc())
                changed = False
            set_status()
            if changed:
                refresh_all()
            if col.pending:
                schedule_sync(1000)
        root.after(200, poll)

    def conflict_dialog(err):
        col = state['col']
        ans = messagebox.askyesnocancel(
            'Changed by a colleague',
            f'A colleague changed the same cells in the shared sheet before your changes were uploaded.\n'
            f'({err})\n\n'
            f'YES - upload MY {col.pending} change(s) anyway (their values in those cells are replaced)\n'
            f'NO - use THEIR version (my changes are kept in a file in the backup folder)\n'
            f'CANCEL - decide later (my changes stay on this computer)')
        if ans is None:
            return
        if ans:
            try:
                col.sync_now(force=True)
            except Exception as e:
                messagebox.showerror('Upload', f'Could not upload: {e}')
        else:
            try:
                keep = col.take_theirs()
                messagebox.showinfo('Their version', f'Your unsent changes are kept in\n{keep}')
            except Exception as e:
                messagebox.showerror('Shared sheet', str(e))
        set_status()
        refresh_all()

    def auto_refresh():
        """Shared mode: every 2 minutes upload anything waiting and pick up colleagues' changes."""
        try:
            start_sync(refresh_only_if_idle=True)
        except Exception:
            log_error(traceback.format_exc())
        root.after(2 * 60 * 1000, auto_refresh)

    def on_close():
        col = state['col']
        if col is not None and col.remote:
            import time
            state['closing'] = True
            if state.get('sync_timer'):
                try:
                    root.after_cancel(state['sync_timer'])
                except tk.TclError:
                    pass
                state['sync_timer'] = None
            t0 = time.time()
            while state.get('syncing') and time.time() - t0 < 120:   # let a running upload finish
                root.update()
                time.sleep(0.05)
            state['syncing'] = None
            if col.pending:
                try:
                    col.sync_now()
                except ConflictError as e:
                    conflict_dialog(e)
                except Exception as e:
                    if not messagebox.askyesno(
                            'Not uploaded',
                            f'{col.pending} change(s) could not be uploaded ({str(e)[:120]}).\n\nThey are saved on '
                            f'this computer and will be uploaded the next time you open Mouse Colony.\n\n'
                            f'Quit anyway?'):
                        state['closing'] = False
                        schedule_sync(60 * 1000)
                        return
                if col.pending and not messagebox.askyesno(
                        'Not uploaded', f'{col.pending} change(s) are still only on this computer. They will be '
                                        f'uploaded the next time you open Mouse Colony.\n\nQuit anyway?'):
                    state['closing'] = False
                    schedule_sync(60 * 1000)
                    return
        root.destroy()

    def shared_dialog():
        cfg = load_config()
        col = state['col']
        d = tk.Toplevel(root)
        d.title('Shared database and backups')
        d.grab_set()
        r = 0
        status = ('Now: shared Google Sheet "' + col.remote.title + '"' + (' - OFFLINE (working on this computer)' if col.readonly
                                                                          else ' - connected')
                  if col is not None and col.remote else
                  'Now: local workbook ' + (col.path if col is not None else '(none)'))
        ttk.Label(d, text=status, font=('TkDefaultFont', 10, 'bold')).grid(row=r, column=0, columnspan=3, sticky='w',
                                                                           padx=10, pady=(10, 6))
        r += 1
        ttk.Label(d, text='Your name').grid(row=r, column=0, sticky='w', padx=10)
        name_e = ttk.Entry(d, width=30)
        name_e.insert(0, cfg.get('user_name', ''))
        name_e.grid(row=r, column=1, sticky='w', padx=10, pady=3)
        ttk.Label(d, text='(default "Person" in logs; recorded in the change log)', foreground='#555').grid(
            row=r, column=2, sticky='w')
        r += 1
        ttk.Label(d, text='Local backup folder').grid(row=r, column=0, sticky='w', padx=10)
        bk_e = ttk.Entry(d, width=50)
        bk_e.insert(0, cfg.get('backup_dir', '') or (col.backup_folder() if col is not None else ''))
        bk_e.grid(row=r, column=1, sticky='w', padx=10, pady=3)
        ttk.Button(d, text='Browse...', command=lambda: (lambda p_: p_ and (bk_e.delete(0, 'end'), bk_e.insert(0, p_)))(
            filedialog.askdirectory(parent=d))).grid(row=r, column=2, sticky='w')
        r += 1
        lf = ttk.LabelFrame(d, text=' Connect to the shared Google Sheet ', padding=8)
        lf.grid(row=r, column=0, columnspan=3, sticky='ew', padx=10, pady=8)
        g = cfg.get('google') or {}
        pick = {'client': g.get('client'), 'key': g.get('key')}
        ttk.Label(lf, text='Google Sheet link').grid(row=0, column=0, sticky='w')
        link_e = ttk.Entry(lf, width=70)
        if g.get('sheet_id'):
            link_e.insert(0, f"https://docs.google.com/spreadsheets/d/{g['sheet_id']}")
        link_e.grid(row=0, column=1, columnspan=2, sticky='w', padx=6, pady=2)
        ttk.Label(lf, text='Login ID file (.json)').grid(row=1, column=0, sticky='w')
        id_var = tk.StringVar()
        ttk.Label(lf, textvariable=id_var, foreground='#333').grid(row=1, column=1, sticky='w', padx=6, pady=2)
        acct_var = tk.StringVar()
        ttk.Label(lf, text='Google account').grid(row=2, column=0, sticky='w')
        ttk.Label(lf, textvariable=acct_var, foreground='#333').grid(row=2, column=1, sticky='w', padx=6, pady=2)

        def show_pick():
            if pick['key']:
                id_var.set(f"service-account key ({pick['key'].get('client_email', '')})")
                acct_var.set('(not needed with a service-account key)')
                return
            id_var.set(('loaded: ' + pick['client']['installed']['client_id'][:24] + '...') if pick['client']
                       else '(none - from GOOGLE_SETUP.md step 4, or use a connection file)')
            tok = load_config().get('google_token') or {}
            ok = pick['client'] and tok.get('client_id') == pick['client']['installed']['client_id']
            acct_var.set(('signed in as ' + (tok.get('account') or '?')) if ok else 'not signed in yet')

        def choose_id():
            import json
            p_ = filedialog.askopenfilename(parent=d, filetypes=[('Login ID / key', '*.json')])
            if not p_:
                return
            try:
                with open(p_, encoding='utf-8') as fh:
                    data = json.load(fh)
            except Exception:
                data = {}
            if data.get('format') == 'mouse-colony-connection':
                save_settings(); d.destroy(); use_connection_file(p_)
                return
            if client_config_ok(data):
                pick.update(client=data, key=None)
            elif data.get('type') == 'service_account':
                pick.update(client=None, key=data)
            else:
                messagebox.showerror('Login ID', 'Choose the "Desktop app" client file downloaded from Google '
                                                 'Cloud (GOOGLE_SETUP.md, step 4).', parent=d)
                return
            show_pick()

        def sign_out():
            c2 = load_config()
            c2.pop('google_token', None)
            save_config(c2)
            show_pick()

        ttk.Button(lf, text='Browse...', command=choose_id).grid(row=1, column=2, sticky='w')
        ttk.Button(lf, text='Sign out', command=sign_out).grid(row=2, column=2, sticky='w')
        show_pick()
        bar1 = ttk.Frame(lf)
        bar1.grid(row=3, column=0, columnspan=3, sticky='w', pady=(6, 0))

        def save_settings():
            c2 = load_config()
            c2['user_name'] = name_e.get().strip()
            c2['backup_dir'] = bk_e.get().strip()
            if col is not None and bk_e.get().strip() == col.backup_folder() and not cfg.get('backup_dir'):
                c2['backup_dir'] = ''
            save_config(c2)
            state['user'] = c2['user_name']
            if state['col'] is not None:
                state['col'].user = c2['user_name']
                state['col'].backup_dir = c2['backup_dir'] or None
            return c2

        def connect():
            save_settings()
            sid = GoogleBackend.sheet_id_from(link_e.get())
            if not sid:
                messagebox.showerror('Connect', 'Paste the link of the Google Sheet.', parent=d)
                return
            if not pick['client'] and not pick['key']:
                messagebox.showerror('Connect', 'Choose the login ID file (GOOGLE_SETUP.md, step 4), or use a '
                                                'connection file from a colleague.', parent=d)
                return
            conn = {'sheet_id': sid, 'title': ''}
            conn.update({'key': pick['key']} if pick['key'] else {'client': pick['client']})
            try:
                be = make_backend(conn, d)
                show_pick()
                d.config(cursor='watch')
                d.update()
                title = be.test()
                data = be.export_xlsx()
                import io
                names = openpyxl.load_workbook(io.BytesIO(data), read_only=True).sheetnames
                if not {'Breeders', 'Cages', 'Pairs'} <= set(names):
                    raise SyncError('This Google Sheet is not the colony workbook (no Breeders/Cages/Pairs '
                                    'sheets). Upload the workbook to Google Drive and open it as a Google Sheet '
                                    'first (GOOGLE_SETUP.md, step 5).')
            except SyncError as e:
                d.config(cursor='')
                messagebox.showerror('Connect', str(e), parent=d)
                return
            except Exception as e:
                d.config(cursor='')
                messagebox.showerror('Connect', f'Could not connect: {e}', parent=d)
                return
            d.config(cursor='')
            conn['title'] = title
            c2 = load_config()
            c2['google'] = conn
            c2['mode'] = 'google'
            save_config(c2)
            d.destroy()
            if open_shared(conn):
                messagebox.showinfo('Connected', f'Connected to "{title}".\n\nA local copy is kept in\n'
                                                 f"{state['col'].backup_folder()}\n\nUse \"Make a connection file\" "
                                                 f'to connect your colleagues.')
        ttk.Button(bar1, text='Sign in with Google and connect', command=connect).pack(side='left')
        ttk.Button(bar1, text='Use a connection file from a colleague...',
                   command=lambda: (save_settings(), d.destroy(), use_connection_file())).pack(side='left', padx=6)
        ttk.Button(bar1, text='Setup guide', command=open_guide).pack(side='left')
        r += 1
        bar = ttk.Frame(d)
        bar.grid(row=r, column=0, columnspan=3, sticky='ew', padx=10, pady=(0, 10))
        shared_now = col is not None and col.remote is not None
        ttk.Button(bar, text='Make a connection file for colleagues...', state='normal' if shared_now else 'disabled',
                   command=lambda: make_connection_file(d)).pack(side='left')
        ttk.Button(bar, text='Restore from a backup...', command=lambda: (save_settings(), restore_dialog(d))).pack(
            side='left', padx=6)
        ttk.Button(bar, text='Use a local workbook instead...',
                   command=lambda: (save_settings(), d.destroy(), open_file())).pack(side='left')
        ttk.Button(bar, text='Card photo settings...',
                   command=lambda: (save_settings(), d.destroy(), photo_settings())).pack(side='left', padx=6)
        ttk.Button(bar, text='Save and close', command=lambda: (save_settings(), d.destroy(), set_status())).pack(
            side='right')

    def open_guide():
        here = getattr(sys, '_MEIPASS', os.path.dirname(os.path.abspath(__file__)))
        p_ = os.path.join(here, 'GOOGLE_SETUP.md')
        import webbrowser
        if os.path.exists(p_):
            webbrowser.open('file://' + os.path.abspath(p_))
        else:
            messagebox.showinfo('Guide', 'GOOGLE_SETUP.md is in the folder you installed from.')

    @guard
    def make_connection_file(parent):
        import json
        g = load_config().get('google') or {}
        if not (g.get('key') or g.get('client')):
            return
        p_ = filedialog.asksaveasfilename(parent=parent, defaultextension='.json',
                                          initialfile='mouse_colony_connection.json',
                                          filetypes=[('Connection file', '*.json')])
        if not p_:
            return
        out = {'format': 'mouse-colony-connection', 'version': 2, 'sheet_id': g['sheet_id'],
               'title': g.get('title', '')}
        if g.get('key'):
            out['service_account'] = g['key']
            note = 'It contains the lab key: share it only inside the lab (not by public links).'
        else:
            out['client'] = g['client']
            note = ('Each colleague signs in with their own Google account, so also share the Google Sheet with '
                    'their address (Share -> Editor). The file holds no password.')
        with open(p_, 'w', encoding='utf-8') as fh:
            json.dump(out, fh, indent=1)
        messagebox.showinfo('Connection file', f'Saved {os.path.basename(p_)}.\n\nColleagues drop it onto the app '
                                               f'(or Shared database -> Use a connection file). {note}',
                            parent=parent)

    def use_connection_file(path_=None):
        import json
        path_ = path_ or filedialog.askopenfilename(filetypes=[('Connection file', '*.json')])
        if not path_:
            return
        try:
            with open(path_, encoding='utf-8') as fh:
                data = json.load(fh)
            assert data.get('format') == 'mouse-colony-connection'
            assert data.get('service_account') or client_config_ok(data.get('client'))
        except Exception:
            messagebox.showerror('Connection file', 'This is not a Mouse Colony connection file.')
            return
        cfg = load_config()
        conn = {'sheet_id': data['sheet_id'], 'title': data.get('title', '')}
        conn.update({'key': data['service_account']} if data.get('service_account') else {'client': data['client']})
        cfg['google'] = conn
        cfg['mode'] = 'google'
        save_config(cfg)
        if open_shared(conn):
            cfg = load_config()
            if not cfg.get('user_name') or (cfg.get('google_token') or {}).get('account', '').startswith(
                    cfg.get('user_name', '') + '@'):
                from tkinter import simpledialog
                nm = simpledialog.askstring('Your name', 'Your name (used as "Person" and in the change log):',
                                            initialvalue=cfg.get('user_name', ''), parent=root)
                if nm and nm.strip():
                    cfg['user_name'] = nm.strip()
                    save_config(cfg)
                    state['user'] = state['col'].user = nm.strip()
            messagebox.showinfo('Connected', f'Connected to the shared sheet "{state["col"].remote.title}".')

    @guard
    def restore_dialog(parent):
        col = state['col']
        items = col.list_backups()
        d = tk.Toplevel(parent)
        d.title('Restore from a backup')
        d.grab_set()
        ttk.Label(d, text=f'Backups in {col.backup_folder()}', font=('TkDefaultFont', 10, 'bold')).pack(
            anchor='w', padx=10, pady=(10, 4))
        fr, tree = make_tree(d, ['When', 'File', 'Size'], [150, 360, 80], height=14)
        fr.pack(fill='both', expand=True, padx=10)
        rows_ = {}
        for when, p_ in items:
            iid = tree.insert('', 'end', values=(when.strftime('%Y-%m-%d %H:%M:%S'), os.path.basename(p_),
                                                 f'{os.path.getsize(p_) // 1024} kB'))
            rows_[iid] = p_
        ttk.Label(d, text=('Restoring makes the chosen copy the current version' +
                           (' of the SHARED sheet for everyone' if col.remote else '') +
                           '. The current version is backed up first, so this can be undone.'),
                  foreground='#555', wraplength=600).pack(anchor='w', padx=10, pady=6)

        @guard
        def ok():
            sel = tree.selection()
            if not sel:
                return
            p_ = rows_[sel[0]]
            if not messagebox.askyesno('Restore', f'Restore {os.path.basename(p_)}?', parent=d):
                return
            col.restore_backup(p_)
            d.destroy()
            set_status()
            refresh_all()
            messagebox.showinfo('Restore', 'Restored.')
        bf = ttk.Frame(d)
        bf.pack(fill='x', padx=10, pady=(0, 10))
        ttk.Button(bf, text='Restore selected', command=ok).pack(side='right')
        ttk.Button(bf, text='Open backup folder', command=lambda: __import__('webbrowser').open(
            'file://' + os.path.abspath(col.backup_folder()))).pack(side='left')

    def reload():
        col = state['col']
        if col and col.remote:
            start_sync()
        elif col:
            col.load()
            set_status()
            refresh_all()

    @guard
    def import_cards(path=None):
        import json
        path = path or filedialog.askopenfilename(title='Cage cards file',
                                                  filetypes=[('Cage cards', '*.json *.txt'), ('All', '*.*')])
        if not path:
            return
        with open(path, encoding='utf-8-sig') as fh:
            text = fh.read()
        try:
            data = json.loads(text)
        except Exception:
            data = None
        if not isinstance(data, dict) or not ('cards' in data or 'events' in data):
            try:
                data = cards_from_text(text)
            except PhotoReadError:
                messagebox.showerror('Import', 'This is not a cage-card or log file (no "cards" or "events").')
                return
        handle_records(data, os.path.basename(path))

    @guard
    def paste_cards():
        try:
            text = root.clipboard_get()
        except tk.TclError:
            text = ''
        if not text.strip():
            messagebox.showinfo('Paste cards', 'Copy the answer from the Gemini chat first (the JSON with the '
                                               'cards), then click Paste cards.')
            return
        try:
            data = cards_from_text(text)
        except PhotoReadError as e:
            messagebox.showerror('Paste cards', str(e))
            return
        handle_records(data, 'Pasted')

    def save_chat_instructions(parent=None):
        col = state['col']
        if col is None:
            return
        p_ = filedialog.asksaveasfilename(parent=parent, defaultextension='.txt',
                                          initialfile='GEMINI_CARD_INSTRUCTIONS.txt',
                                          filetypes=[('Text', '*.txt')])
        if not p_:
            return
        with open(p_, 'w', encoding='utf-8') as fh:
            fh.write(chat_instructions(col))
        messagebox.showinfo('Instructions', f'Saved {os.path.basename(p_)}.\n\nIn gemini.google.com: attach this '
                                            f'file and the card photos, send. Copy the answer and click "Paste '
                                            f'cards" in Mouse Colony.\n\nMake a new file after adding lines or '
                                            f'cages, so the lists in it stay current.', parent=parent)

    def photo_settings(parent=None, then=None):
        """Gemini API key (free, aistudio.google.com) and model for reading card photos."""
        cfg = load_config()
        d = tk.Toplevel(parent or root)
        d.title('Card photo settings')
        d.grab_set()
        ttk.Label(d, text='Reading cage-card photos with Google Gemini', font=('TkDefaultFont', 10, 'bold')).grid(
            row=0, column=0, columnspan=3, sticky='w', padx=10, pady=(10, 2))
        ttk.Label(d, text='Get a free API key at aistudio.google.com -> API keys (see CAGE_CARDS_FORMAT.md). The key '
                          'stays on this computer. On the free plan Google may use the photos to improve its '
                          'products, so photograph only the cards.', wraplength=520, foreground='#555',
                  justify='left').grid(row=1, column=0, columnspan=3, sticky='w', padx=10, pady=(0, 8))
        ttk.Label(d, text='Gemini API key').grid(row=2, column=0, sticky='w', padx=10)
        key_e = ttk.Entry(d, width=50, show='*')
        key_e.insert(0, cfg.get('gemini_key', ''))
        key_e.grid(row=2, column=1, sticky='w', pady=3)
        show_v = tk.BooleanVar()
        ttk.Checkbutton(d, text='show', variable=show_v,
                        command=lambda: key_e.config(show='' if show_v.get() else '*')).grid(row=2, column=2,
                                                                                            sticky='w', padx=6)
        ttk.Label(d, text='Model').grid(row=3, column=0, sticky='w', padx=10)
        model_e = ttk.Entry(d, width=30)
        model_e.insert(0, cfg.get('gemini_model', '') or GEMINI_DEFAULT_MODEL)
        model_e.grid(row=3, column=1, sticky='w', pady=3)
        bar = ttk.Frame(d)
        bar.grid(row=4, column=0, columnspan=3, sticky='ew', padx=10, pady=10)

        def save():
            c2 = load_config()
            c2['gemini_key'] = key_e.get().strip()
            c2['gemini_model'] = model_e.get().strip() or GEMINI_DEFAULT_MODEL
            save_config(c2)
            d.destroy()
            if then and c2['gemini_key']:
                then()
        import webbrowser
        ttk.Button(bar, text='Open AI Studio', command=lambda: webbrowser.open(
            'https://aistudio.google.com/apikey')).pack(side='left')
        ttk.Button(bar, text='Instructions file for the Gemini chat...',
                   command=lambda: save_chat_instructions(d)).pack(side='left', padx=6)
        ttk.Button(bar, text='Save', command=save).pack(side='right')
        ttk.Button(bar, text='Cancel', command=d.destroy).pack(side='right', padx=6)
        key_e.focus_set()

    def read_photos(paths=None):
        """Read cage-card photos with Gemini in the background, then show the usual review screen."""
        import threading, queue
        col = state['col']
        if col is None:
            messagebox.showinfo('Card photos', 'Open the colony workbook first.')
            return
        cfg = load_config()
        if not cfg.get('gemini_key'):
            photo_settings(then=lambda: read_photos(paths))
            return
        paths = list(paths or filedialog.askopenfilenames(
            title='Cage-card photos', filetypes=[('Photos', ' '.join('*' + e for e in PHOTO_EXT)), ('All', '*.*')]))
        if not paths:
            return
        prompt = card_prompt(col)
        key, model = cfg['gemini_key'], cfg.get('gemini_model') or GEMINI_DEFAULT_MODEL
        w = tk.Toplevel(root)
        w.title('Reading card photos')
        w.grab_set()
        msg_v = tk.StringVar(value=f'Reading 1 of {len(paths)}: {os.path.basename(paths[0])}')
        ttk.Label(w, textvariable=msg_v, width=60).pack(padx=14, pady=(14, 6), anchor='w')
        pb = ttk.Progressbar(w, maximum=len(paths), length=420)
        pb.pack(padx=14, pady=(0, 6))
        stop = {'flag': False}
        ttk.Button(w, text='Stop', command=lambda: stop.update(flag=True)).pack(pady=(0, 12))
        q = queue.Queue()

        def work():
            for i, p_ in enumerate(paths):
                if stop['flag']:
                    break
                q.put(('start', i, p_))
                try:
                    q.put(('ok', i, read_card_photo(p_, key, model, prompt=prompt)))
                except PhotoReadError as e:
                    q.put(('err', i, f'{os.path.basename(p_)}: {e}'))
                    if 'limit' in str(e) or 'API key' in str(e) or 'internet' in str(e):
                        break
                except Exception as e:
                    log_error(traceback.format_exc())
                    q.put(('err', i, f'{os.path.basename(p_)}: {e.__class__.__name__}: {e}'))
            q.put(('done', None, None))
        cards, events, errors = [], [], []

        def poll():
            try:
                while True:
                    kind, i, val = q.get_nowait()
                    if kind == 'start':
                        msg_v.set(f'Reading {i + 1} of {len(paths)}: {os.path.basename(val)}')
                    elif kind == 'ok':
                        cards.extend(val['cards'])
                        events.extend(val['events'])
                        pb['value'] = i + 1
                    elif kind == 'err':
                        errors.append(val)
                        pb['value'] = i + 1
                    else:
                        w.destroy()
                        finish()
                        return
            except queue.Empty:
                pass
            root.after(200, poll)

        def finish():
            if any('API key' in e_ for e_ in errors):
                messagebox.showerror('Card photos', errors[0])
                photo_settings()
                return
            if errors:
                messagebox.showwarning('Card photos', 'Some photos could not be read:\n\n' + '\n'.join(errors[:8]) +
                                       ('\n...' if len(errors) > 8 else ''))
            if not cards and not events:
                if not errors:
                    messagebox.showinfo('Card photos', 'Gemini found no cage cards or log rows in the photos.')
                return
            data = {'format': 'mouse-colony-cards', 'version': 1, 'read_on': dt.date.today().isoformat(),
                    'source': f'{len(paths)} photo(s) read by Gemini ({model})', 'cards': cards, 'events': events}
            handle_records(data, 'Gemini read')
        threading.Thread(target=work, daemon=True).start()
        root.after(200, poll)

    @guard
    def enter_cards():
        """Type cage cards by hand; they go through the same review as an imported file."""
        data = {'format': 'mouse-colony-cards', 'version': 1, 'read_on': dt.date.today().isoformat(),
                'source': 'typed in', 'cards': []}
        rv = review_cards(data, 'Cards typed in')
        rv['add']()

    def review_cards(data, label):
        col = state['col']
        changes = []
        d = tk.Toplevel(root)
        d.title('Cage cards - review')
        d.geometry('1080x580')
        d.grab_set()
        head_v, sub_v = tk.StringVar(), tk.StringVar()
        ttk.Label(d, textvariable=head_v, font=('TkDefaultFont', 10, 'bold')).pack(anchor='w', padx=10,
                                                                                    pady=(10, 0))
        ttk.Label(d, textvariable=sub_v, foreground='#555').pack(anchor='w', padx=10, pady=(2, 6))
        fr, tree = make_tree(d, ['Apply', 'Cage', 'What', 'In the app', 'On the card'],
                             [60, 70, 330, 230, 300], height=16)
        fr.pack(fill='both', expand=True, padx=10)
        tree.tag_configure('on', background='#E2F0D9')
        tree.tag_configure('off', background='#FFFFFF')
        tree.tag_configure('info', background='#EDEDED', foreground='#555')
        iids = {}

        def recompute():
            old = {(ch['cage'], ch['what']): ch['apply'] for ch in changes}
            changes[:] = col.compare_cards(data)
            for ch in changes:
                if (ch['cage'], ch['what']) in old and ch['kind'] != 'info':
                    ch['apply'] = old[(ch['cage'], ch['what'])]
            n = len(data['cards'])
            head_v.set(f"{label}: {n} card{'' if n == 1 else 's'}"
                       + (f", read {data.get('read_on')}" if data.get('read_on') else '')
                       + (f" from {data.get('source')}" if data.get('source') and data.get('source') != 'typed in'
                          else ''))
            n_ok = len({cg for cg in (col.norm_cage(c.get('cage')) for c in data['cards'])}
                       - {ch['cage'] for ch in changes})
            sub_v.set((f"{n_ok} card{'' if n_ok == 1 else 's'} match{'es' if n_ok == 1 else ''} the workbook. "
                       if n else 'No cards yet - add one. ') +
                      'Click a row to switch "Apply" on or off; grey rows are information only.')
            draw()

        def draw():
            tree.delete(*tree.get_children())
            iids.clear()
            for i, ch in enumerate(changes):
                ap = 'info' if ch['kind'] == 'info' else ('yes' if ch['apply'] else 'no')
                iid = tree.insert('', 'end', values=(ap, col.card_name(ch['cage']), ch['what'], ch['app'],
                                                     ch['card']),
                                  tags=('info' if ch['kind'] == 'info' else 'on' if ch['apply'] else 'off',))
                iids[iid] = i
        draw()

        def toggle(_=None):
            sel = tree.selection()
            if not sel:
                return
            ch = changes[iids[sel[0]]]
            if ch['kind'] != 'info':
                ch['apply'] = not ch['apply']
                idx = tree.index(sel[0])
                draw()
                tree.selection_set(tree.get_children()[idx])
        tree.bind('<ButtonRelease-1>', toggle)
        bar = ttk.Frame(d)
        bar.pack(fill='x', padx=10, pady=8)

        def add_card(card):
            data['cards'] = [c for c in data['cards'] if col.norm_cage(c.get('cage')) != col.norm_cage(card['cage'])]
            data['cards'].append(card)
            recompute()
        ttk.Button(bar, text='Add a card by hand...', command=lambda: card_form(d, add_card)).pack(side='left')
        ttk.Button(bar, text='Save cards as .json...', command=lambda: save_cards_json(data)).pack(side='left',
                                                                                              padx=6)

        @guard
        def ok():
            when = col._card_date(data.get('read_on'))
            todo = [ch for ch in changes if ch['apply'] and ch['kind'] != 'info']
            if not todo:
                d.destroy()
                return
            n = col.apply_card_changes(todo, checked_on=when)
            d.destroy()
            refresh_all()
            messagebox.showinfo('Import', f'{len(todo)} changes applied to {n} cages.')
        ttk.Button(bar, text='Apply ticked changes', command=ok).pack(side='right')
        ttk.Button(bar, text='Cancel', command=d.destroy).pack(side='right', padx=6)
        recompute()
        return {'add': lambda: card_form(d, add_card), 'win': d}

    def handle_records(data, label):
        """Cards and/or log events from a file, a paste or photos: card review first, then the events review."""
        cards, events = data.get('cards') or [], data.get('events') or []
        if not cards and not events:
            messagebox.showinfo('Import', 'No cage cards or log rows in it.')
            return
        if cards:
            rv = review_cards(dict(data, cards=cards),
                              f'{label}: check the values against the cards' if label == 'Gemini read' else label)
            if events:
                root.wait_window(rv['win'])
        if events:
            review_events(dict(data, events=events), label)

    def review_events(data, label):
        """Animal-house log rows: kivitel -> terminate, áthelyezés -> move. Applied rows are remembered."""
        col = state['col']
        changes = col.compare_events(data)
        d = tk.Toplevel(root)
        d.title('Animal-house log - review')
        d.geometry('1120x560')
        d.grab_set()
        n_ev = len(data['events'])
        ttk.Label(d, text=f"{label}: {n_ev} log row{'' if n_ev == 1 else 's'}" +
                          (f" from {data.get('source')}" if data.get('source') else ''),
                  font=('TkDefaultFont', 10, 'bold')).pack(anchor='w', padx=10, pady=(10, 0))
        ttk.Label(d, text='Kivitel = taken out (goes to the Terminated history), áthelyezés = moved to another cage. '
                          'Check every row against the log page. Click a row to switch "Apply" on or off; grey rows '
                          'are information only. Rows with unclear readings start off. Applied rows are remembered '
                          'and never applied again.', foreground='#555', wraplength=1080, justify='left').pack(
            anchor='w', padx=10, pady=(2, 6))
        fr, tree = make_tree(d, ['Apply', 'Date', 'Cage', 'What happens', 'In the app', 'On the log'],
                             [60, 90, 75, 420, 200, 250], height=15)
        fr.pack(fill='both', expand=True, padx=10)
        tree.tag_configure('on', background='#F8CBAD')
        tree.tag_configure('on_move', background='#E2F0D9')
        tree.tag_configure('off', background='#FFFFFF')
        tree.tag_configure('info', background='#EDEDED', foreground='#555')
        iids = {}

        def draw():
            tree.delete(*tree.get_children())
            iids.clear()
            for i, ch in enumerate(changes):
                info = ch['kind'] == 'info'
                tag = 'info' if info else ('off' if not ch['apply'] else 'on_move' if ch['kind'] == 'move' else 'on')
                iid = tree.insert('', 'end', values=('info' if info else 'yes' if ch['apply'] else 'no',
                                                     fmt_date(ch['date']) if ch['date'] else '?',
                                                     col.card_name(ch['cage']), ch['what'], ch['app'], ch['card']),
                                  tags=(tag,))
                iids[iid] = i

        def toggle(_=None):
            sel = tree.selection()
            if not sel:
                return
            ch = changes[iids[sel[0]]]
            if ch['kind'] != 'info':
                ch['apply'] = not ch['apply']
                idx = tree.index(sel[0])
                draw()
                tree.selection_set(tree.get_children()[idx])
        tree.bind('<ButtonRelease-1>', toggle)
        draw()
        bar = ttk.Frame(d)
        bar.pack(fill='x', padx=10, pady=8)

        def save_json():
            import json
            p_ = filedialog.asksaveasfilename(parent=d, defaultextension='.json', filetypes=[('Log', '*.json')],
                                              initialfile=f"animal_log_{data.get('read_on', '')}.json")
            if p_:
                with open(p_, 'w', encoding='utf-8') as fh:
                    json.dump({k: v for k, v in data.items() if k != 'cards'}, fh, indent=2, ensure_ascii=False,
                              default=str)
        ttk.Button(bar, text='Save log rows as .json...', command=save_json).pack(side='left')

        @guard
        def ok():
            todo = [ch for ch in changes if ch['apply'] and ch['kind'] != 'info']
            if not todo:
                d.destroy()
                return
            n_t = sum(1 for ch in todo if ch['kind'] != 'move')
            if not messagebox.askyesno('Apply', f'Apply {len(todo)} change(s)? {n_t} of them terminate animals.',
                                       parent=d):
                return
            try:
                n = col.apply_events(todo, person=state.get('user', ''))
            finally:
                d.destroy()
                refresh_all()
            messagebox.showinfo('Animal-house log', f'{n} log row(s) applied.')
        ttk.Button(bar, text='Apply ticked changes', command=ok).pack(side='right')
        ttk.Button(bar, text='Cancel', command=d.destroy).pack(side='right', padx=6)

    def save_cards_json(data):
        import json
        p = filedialog.asksaveasfilename(defaultextension='.json', filetypes=[('Cage cards', '*.json')],
                                         initialfile=f"cage_cards_{data.get('read_on', '')}.json")
        if p:
            with open(p, 'w', encoding='utf-8') as fh:
                json.dump(data, fh, indent=2, ensure_ascii=False)

    def card_form(parent, on_add):
        """Type one cage card. Picking a cage fills in what the workbook has; change what the card says."""
        col = state['col']
        f = tk.Toplevel(parent)
        f.title('Enter a cage card')
        f.transient(parent)
        f.grab_set()
        names = [col.card_name(c['cage']) for c in col.cages]
        w = {}
        row = [0]

        def lab(text, widget, **kw):
            ttk.Label(f, text=text).grid(row=row[0], column=0, sticky='w', padx=10, pady=3)
            widget.grid(row=row[0], column=1, sticky='w', padx=10, **kw)
            row[0] += 1
            return widget
        w['cage'] = lab('Ketrec kód (cage)', ttk.Combobox(f, values=names, width=12))
        w['type'] = lab('Card type', ttk.Combobox(f, values=['stock', 'pair', 'empty'], width=8, state='readonly'))
        w['line'] = lab('Line', enhance_line_combo(ttk.Combobox(f, width=32)))
        w['as'] = lab('Törzs as written', ttk.Entry(f, width=30))
        w['sex'] = lab('Nem (sex)', ttk.Combobox(f, values=['M', 'F'], width=4, state='readonly'))
        w['count'] = lab('Egerek száma (count)', ttk.Entry(f, width=6))
        w['birth'] = lab('Születés ideje (born)', ttk.Entry(f, width=14))
        w['males'] = lab('Pair: male birth date(s)', ttk.Entry(f, width=30))
        w['females'] = lab('Pair: female birth date(s)', ttk.Entry(f, width=30))
        w['origin'] = lab('Származási hely (origin)', ttk.Combobox(f, values=names, width=12))
        w['arrived'] = lab('Behozatal ideje (arrived)', ttk.Entry(f, width=14))
        w['notes'] = lab('Notes / procedure', ttk.Entry(f, width=40))
        unsure_fr = ttk.Frame(f)
        unsure = {k: tk.BooleanVar() for k in ('line', 'sex', 'count', 'birth')}
        for k, v in unsure.items():
            ttk.Checkbutton(unsure_fr, text=k, variable=v).pack(side='left', padx=(0, 6))
        lab('Hard to read', unsure_fr)
        ttk.Label(f, text='Dates as on the card are fine (26.09.04 or 2026.09.04). Several pair dates: separate '
                          'with commas.', foreground='#555').grid(row=row[0], column=0, columnspan=2, sticky='w',
                                                                  padx=10)
        row[0] += 1
        status = tk.StringVar()
        ttk.Label(f, textvariable=status, foreground='#1F6F3F').grid(row=row[0], column=0, columnspan=2,
                                                                     sticky='w', padx=10)
        row[0] += 1

        def setv(k, v):
            x = w[k]
            if isinstance(x, ttk.Combobox):
                x.set(v or '')
            else:
                x.delete(0, 'end')
                x.insert(0, v or '')

        def prefill(_=None):
            cage = col.norm_cage(w['cage'].get())
            c = col.cage(cage)
            for k in ('line', 'as', 'sex', 'count', 'birth', 'males', 'females', 'origin', 'arrived', 'notes'):
                setv(k, '')
            if not c:
                return
            p = col.active_pair(cage)
            if p:
                setv('type', 'pair')
                setv('line', p['line'])
                ms = [col.breeders[x] for x in col.pair_members(p) if x in col.breeders]
                setv('males', ', '.join(fmt_date(b['birth']) for b in ms if b['sex'] == 'M'))
                setv('females', ', '.join(fmt_date(b['birth']) for b in ms if b['sex'] == 'F'))
            elif col.mice_in(cage) or c['count']:
                setv('type', 'stock')
                setv('line', c['line'])
                setv('sex', c['sex'])
                setv('count', str(len(col.mice_in(cage)) or c['count'] or ''))
                setv('birth', fmt_date(c['birth']))
            else:
                setv('type', 'stock')
            for v in unsure.values():
                v.set(False)
        w['cage'].bind('<<ComboboxSelected>>', prefill)
        w['cage'].bind('<FocusOut>', lambda e: prefill() if not w['line'].get() else None)

        def build():
            cage = col.norm_cage(w['cage'].get())
            if not col.cage(cage):
                messagebox.showerror('Cage', f'"{w["cage"].get()}" is not a cage in the workbook.', parent=f)
                return None
            card = {'cage': col.card_name(cage), 'type': w['type'].get() or 'stock'}
            def date_or_err(txt, what):
                if not txt.strip():
                    return ''
                d0 = col._card_date(txt)
                if not d0:
                    raise ValueError(f'{what}: "{txt}" is not a date')
                return d0.isoformat()
            try:
                if card['type'] == 'pair':
                    card['males'] = [{'birth': date_or_err(x, 'Male birth')} for x in
                                     w['males'].get().split(',') if x.strip()]
                    card['females'] = [{'birth': date_or_err(x, 'Female birth')} for x in
                                       w['females'].get().split(',') if x.strip()]
                elif card['type'] == 'stock':
                    cnt = w['count'].get().strip()
                    if cnt and not cnt.isdigit():
                        raise ValueError('Count must be a number')
                    card.update(line=w['line'].get().strip(), line_as_written=w['as'].get().strip(),
                                sex=w['sex'].get(), count=int(cnt) if cnt else None,
                                birth=date_or_err(w['birth'].get(), 'Birth'))
                    if w['origin'].get().strip():
                        card['origin'] = w['origin'].get().strip()
                    if w['arrived'].get().strip():
                        card['arrived'] = date_or_err(w['arrived'].get(), 'Arrived')
                    card['uncertain'] = [k for k, v in unsure.items() if v.get()]
            except ValueError as e:
                messagebox.showerror('Card', str(e), parent=f)
                return None
            if w['notes'].get().strip():
                card['notes'] = w['notes'].get().strip()
            return card

        def add(close):
            card = build()
            if not card:
                return
            on_add(card)
            if close:
                f.destroy()
            else:
                status.set(f"Added {card['cage']}. Next card:")
                setv('cage', '')
                prefill()
                w['cage'].focus_set()
        bf = ttk.Frame(f)
        bf.grid(row=row[0], column=0, columnspan=2, sticky='e', padx=10, pady=10)
        ttk.Button(bf, text='Add and enter next', command=lambda: add(False)).pack(side='left')
        ttk.Button(bf, text='Add and close', command=lambda: add(True)).pack(side='left', padx=6)
        ttk.Button(bf, text='Close', command=f.destroy).pack(side='left')
        w['cage'].focus_set()

    def on_drop(event):
        files = root.tk.splitlist(event.data)
        for f in files:
            if f.lower().endswith(('.json', '.txt')):
                try:
                    import json
                    with open(f, encoding='utf-8') as fh:
                        kind = json.load(fh).get('format')
                except Exception:
                    kind = None
                if kind == 'mouse-colony-connection':
                    use_connection_file(f)
                else:
                    import_cards(f)
            elif f.lower().endswith('.xlsx'):
                open_file(f)
        photos = [f for f in files if f.lower().endswith(PHOTO_EXT)]
        if photos:
            read_photos(photos)

    def selected(tree):
        sel = tree.selection()
        return tree.item(sel[0], 'values') if sel else None

    @guard
    def check_selected_pair():
        v = selected(pairs_tree)
        if not v:
            return
        pc_m.set(unit_label(v[2]))
        fs = [f.strip() for f in v[3].split(',') if f.strip()]
        pc_f1.set(unit_label(fs[0]) if fs else '')
        pc_f2.set(unit_label(fs[1]) if len(fs) > 1 else '')
        nb.select(pc)
        run_check()

    @guard
    def suggest_for_selected():
        v = selected(pairs_tree)
        if not v:
            return
        sg_line.set(v[1])
        nb.select(sg)
        run_suggest()

    @guard
    def end_selected():
        v = selected(pairs_tree)
        if not v:
            return
        ans = messagebox.askyesnocancel(
            'End pair', f'End pair {v[0]}.\n\nWere the animals terminated?\n\n'
                        f'Yes - record them in the termination history\n'
                        f'No - they stay alive (kept in the cage as retired breeders)')
        if ans is None:
            return
        if ans:
            terminate_dialog(cage=v[0])
        else:
            state['col'].end_pair_keep(v[0])
            refresh_all()

    @guard
    def run_check():
        col = state['col']
        m = unit_from_label(pc_m.get())
        fs = [unit_from_label(x.get()) for x in (pc_f1, pc_f2) if x.get().strip()]
        if not m or not fs:
            messagebox.showinfo('Pairing check', 'Pick a male and at least one female.')
            return
        fill(check_tree, [((msg,), lvl) for lvl, msg in col.check_pair(m, fs)])

    @guard
    def create_from_check():
        col = state['col']
        m = unit_from_label(pc_m.get())
        fs = [unit_from_label(x.get()) for x in (pc_f1, pc_f2) if x.get().strip()]
        if not m or not fs:
            return
        res = col.check_pair(m, fs)
        bad = [msg for lvl, msg in res if lvl == 'bad']
        if bad and not messagebox.askyesno('Problems found', 'This pair has problems:\n\n' + '\n'.join(bad) +
                                           '\n\nCreate it anyway?'):
            return
        lines = [col.unit_info(u)['line'] for u in [m] + fs]
        non_b6 = list(dict.fromkeys(l for l in lines if l and l != 'C57Bl6J'))
        if len(non_b6) > 1:
            default_line = col.combined_name(non_b6)
        else:
            default_line = non_b6[0] if non_b6 else lines[0]
        pair_dialog(m, fs, default_line)

    def pair_dialog(m, fs, line):
        col = state['col']
        d = tk.Toplevel(root)
        d.title('Create pair')
        d.grab_set()
        ttk.Label(d, text=f'Male: {m}    Females: {", ".join(fs)}').grid(row=0, column=0, columnspan=2, padx=10,
                                                                         pady=8, sticky='w')
        ttk.Label(d, text='Pair cage').grid(row=1, column=0, sticky='w', padx=10)
        cage_cb = ttk.Combobox(d, values=col.empty_cages(), width=12)
        cage_cb.grid(row=1, column=1, sticky='w')
        if col.empty_cages():
            cage_cb.set(col.empty_cages()[0])
        ttk.Label(d, text='Line').grid(row=2, column=0, sticky='w', padx=10)
        line_cb = ttk.Combobox(d, width=32)
        line_cb.set(line)
        enhance_line_combo(line_cb)
        line_cb.grid(row=2, column=1, sticky='w')
        if line not in col.lines:
            ttk.Label(d, text=f'"{line}" is a new mixed line - it will be added to the Lines sheet.',
                      foreground='#1F4E78').grid(row=6, column=0, columnspan=2, sticky='w', padx=10)
        ttk.Label(d, text='Set up (yyyy-mm-dd)').grid(row=3, column=0, sticky='w', padx=10)
        setup_e = ttk.Entry(d, width=14)
        setup_e.insert(0, dt.date.today().isoformat())
        setup_e.grid(row=3, column=1, sticky='w')
        ttk.Label(d, text='Animals from stock cages get new IDs (B###) and the cage count goes down by one.',
                  foreground='#555').grid(row=4, column=0, columnspan=2, padx=10, pady=6, sticky='w')

        @guard
        def ok():
            if not cage_cb.get().strip():
                return
            sd = parse_date(setup_e.get())
            ln = line_cb.get().strip()
            if ln and ln not in col.lines:
                parts = col.line_parts(ln)
                if all(x in col.lines for x in parts) and len(parts) > 1:
                    ln = col.ensure_combined_line(parts)
                elif not messagebox.askyesno('New line', f'"{ln}" is not on the Lines sheet. Add it as a new '
                                                         f'line (you can set its alleles later)?', parent=d):
                    return
                else:
                    col.add_line(ln, partners=[ln])
            nm, nfs = col.create_pair(cage_cb.get().strip(), ln, m, fs, sd.date() if sd else None)
            d.destroy()
            refresh_all()
            messagebox.showinfo('Pair created', f'Pair in {cage_cb.get()}: male {nm}, female(s) {", ".join(nfs)}.\n'
                                                f'Write the new IDs on the cage card / ear tags.')
        ttk.Button(d, text='Create', command=ok).grid(row=5, column=1, sticky='e', padx=10, pady=10)

    @guard
    def run_suggest():
        col = state['col']
        line = sg_line.get()
        if not line:
            return
        rows_ = []
        for s in col.suggest(line):
            lvl = 'bad' if s['kin'] >= 0.125 else 'warn' if s['kin'] > 0 else 'ok'
            rows_.append(((unit_label(s['m']), unit_label(s['f']), f"{s['kin']:.3f}", col.risk(s['kin']), s['rel'],
                           s['left'], s['pups']), lvl))
        if not rows_:
            rows_ = [(('No candidates - no Reserve breeders or young stock cages of the partner lines', '', '', '', '',
                       ''), 'bad')]
        fill(sug_tree, rows_)

    @guard
    def suggestion_to_check():
        v = selected(sug_tree)
        if not v or not v[1]:
            return
        pc_all.set(True)
        refresh_pc_lists()
        pc_m.set(v[0])
        pc_f1.set(v[1])
        pc_f2.set('')
        nb.select(pc)
        run_check()

    def edit_selected_breeder():
        v = selected(br_tree)
        if v:
            breeder_dialog(v[0])

    @guard
    def breeder_dialog(bid):
        col = state['col']
        b = dict(col.breeders[bid]) if bid else dict(id=col.next_id(), line='', sex='', birth=None, cage='',
                                                         status='Reserve', sire='', dam='', notes='')
        d = tk.Toplevel(root)
        d.title('Edit breeder' if bid else 'Add breeder')
        d.grab_set()
        ids = sorted(col.breeders)
        fields = [('ID', 'entry', None), ('Line', 'combo', sorted(col.lines)), ('Sex', 'combo', ['M', 'F']),
                  ('Birth (yyyy-mm-dd)', 'entry', None), ('Cage', 'entry', None),
                  ('Status', 'combo', ['Breeding', 'Reserve', 'Retired', 'Terminated']),
                  ('Sire ID', 'combo', [''] + ids), ('Dam ID', 'combo', [''] + ids), ('Notes', 'entry', None)]
        keys = ['id', 'line', 'sex', 'birth', 'cage', 'status', 'sire', 'dam', 'notes']
        widgets = {}
        for i, ((lab, kind, vals), k) in enumerate(zip(fields, keys)):
            ttk.Label(d, text=lab).grid(row=i, column=0, sticky='w', padx=10, pady=3)
            w = ttk.Combobox(d, values=vals, width=30) if kind == 'combo' else ttk.Entry(d, width=33)
            val = fmt_date(b[k]) if k == 'birth' else b[k]
            if kind == 'combo':
                w.set(val)
            else:
                w.insert(0, val)
            w.grid(row=i, column=1, sticky='w', padx=10)
            widgets[k] = w
        if bid:
            widgets['id'].configure(state='disabled')
        enhance_line_combo(widgets['line'])
        ttk.Label(d, text='Genotype').grid(row=len(fields), column=0, sticky='nw', padx=10, pady=3)
        gfr, gbuild, gget = gt_editor(d, b['line'], b.get('gt') or {})
        gfr.grid(row=len(fields), column=1, sticky='w', padx=10, pady=3)
        ttk.Label(d, text=GT_HELP, foreground='#555', wraplength=330).grid(row=len(fields) + 1, column=1,
                                                                          sticky='w', padx=10)
        widgets['line'].bind('<<ComboboxSelected>>', lambda e: gbuild(widgets['line'].get()), add='+')

        @guard
        def ok():
            nb_ = {k: widgets[k].get().strip() for k in keys}
            nb_['gt'] = gget()
            nb_['id'] = b['id']
            nb_['sex'] = nb_['sex'].upper()
            bd = parse_date(nb_['birth'])
            nb_['birth'] = bd.date() if bd else None
            if nb_['sex'] not in ('M', 'F'):
                messagebox.showerror('Sex', 'Sex must be M or F', parent=d)
                return
            for k in ('sire', 'dam'):
                if nb_[k] and nb_[k] not in col.breeders:
                    messagebox.showerror('Parent', f'{nb_[k]} is not a breeder ID. Add the parent first '
                                                   f'(status Retired is fine).', parent=d)
                    return
                if nb_[k] == nb_['id']:
                    messagebox.showerror('Parent', 'An animal cannot be its own parent.', parent=d)
                    return
            col.save_breeder(nb_)
            d.destroy()
            refresh_all()
        ttk.Button(d, text='Save', command=ok).grid(row=len(fields) + 2, column=1, sticky='e', padx=10, pady=10)

    @guard
    def take_dialog():
        col = state['col']
        d = tk.Toplevel(root)
        d.title('Take breeders from a stock cage')
        d.grab_set()
        stock = [f"{c['cage']}  -  {c['line']}, {c['count']} {c['sex']}, born {fmt_date(c['birth'])}"
                 for c in col.cages if c['sex'] in ('M', 'F') and c['count']]
        ttk.Label(d, text='Stock cage').grid(row=0, column=0, padx=10, pady=4, sticky='w')
        cb = ttk.Combobox(d, values=stock, width=60, state='readonly')
        cb.grid(row=0, column=1, padx=10)
        ttk.Label(d, text='Mice').grid(row=1, column=0, padx=10, pady=4, sticky='nw')
        holder = ttk.Frame(d)
        holder.grid(row=1, column=1, sticky='w', padx=10)
        getter = {'f': lambda: []}

        def on_pick(_=None):
            for w in holder.winfo_children():
                w.destroy()
            name = cb.get().split('  -  ')[0]
            fr, get = mouse_checklist(holder, name, preselect=('-',))   # nothing ticked by default
            fr.pack(anchor='w')
            getter['f'] = get
        cb.bind('<<ComboboxSelected>>', on_pick)
        ttk.Label(d, text='Ticked mice become Reserve breeders with B-IDs; parents and genotype come with them.',
                  foreground='#555').grid(row=2, column=0, columnspan=2, padx=10, pady=4, sticky='w')

        @guard
        def ok():
            if not cb.get():
                return
            name = cb.get().split('  -  ')[0]
            c = col.cage(name)
            chosen = getter['f']()
            if not chosen:
                messagebox.showerror('Mice', 'Tick the mice to take.', parent=d)
                return
            wb = col._open_for_write()
            ids_ = col.take_from_cage(wb, name, c['sex'], mouse_ids=chosen)
            col._save(wb)
            d.destroy()
            refresh_all()
            messagebox.showinfo('Done', f'New Reserve breeders: {", ".join(ids_)}')
        ttk.Button(d, text='Take', command=ok).grid(row=3, column=1, sticky='e', padx=10, pady=10)

    @guard
    def save_litter():
        col = state['col']
        cage = wl_pair.get().split(' ')[0]
        if not cage or not wl_dam.get():
            messagebox.showinfo('Litter', 'Pick the pair and the mother.')
            return
        born, weaned = parse_date(wl_born.get()), parse_date(wl_weaned.get())
        if not born:
            messagebox.showinfo('Litter', 'Enter the birth date.')
            return
        groups = []
        for sex, n_w, c_w in (('M', wl_mn, wl_mc), ('F', wl_fn, wl_fc)):
            if n_w.get().strip():
                if not c_w.get().strip():
                    messagebox.showinfo('Litter', f'Pick a cage for the {sex} pups.')
                    return
                groups.append((sex, int(n_w.get()), c_w.get().strip()))
        if not groups:
            messagebox.showinfo('Litter', 'Enter at least one group of pups.')
            return
        pups = int(wl_pups.get()) if wl_pups.get().strip() else None
        created = col.wean_litter(cage, wl_dam.get(), born, weaned or dt.datetime.now(), groups, pups,
                                  wl_notes.get(), unknown_gt=wl_unknown.get())
        for w in (wl_born, wl_pups, wl_mn, wl_fn, wl_notes):
            w.delete(0, 'end')
        wl_mc.set('')
        wl_fc.set('')
        refresh_all()
        ids = [x for v in created.values() for x in v]
        n_flag = sum(1 for x in ids if col.needs_genotyping(x))
        if not wl_unknown.get() and n_flag:
            for target, new_ids in created.items():
                genotyping_dialog(target, new_ids)
        else:
            messagebox.showinfo('Litter saved', f'{len(ids)} pups got IDs and ear marks; litter logged.'
                                + (f'\n{n_flag} are flagged "needs genotyping" - filter All mice by '
                                   f'"Needs genotyping" to find them.' if n_flag else ''))

    if dnd_files:
        try:
            root.drop_target_register(dnd_files)
            root.dnd_bind('<<Drop>>', on_drop)
            info_var.set('Tip: drop card photos, a cage-card .json file or a workbook onto this window')
        except Exception:
            pass
    cfg0 = load_config()
    state['user'] = cfg0.get('user_name', '')
    g0 = cfg0.get('google') or {}
    if not path and cfg0.get('mode') == 'google' and (g0.get('key') or client_config_ok(g0.get('client'))):
        if not open_shared(cfg0['google']):
            start = default_workbook()
            if start:
                open_file(start)
    else:
        start = path or default_workbook()
        if start:
            open_file(start)
        else:
            path_var.set('No workbook yet - click "Open workbook..." or "Shared database..."')
    root.after(2 * 60 * 1000, auto_refresh)
    root.protocol('WM_DELETE_WINDOW', on_close)
    root.mainloop()


if __name__ == '__main__':
    try:
        run_gui(sys.argv[1] if len(sys.argv) > 1 else None)
    except Exception:
        import traceback
        log_error(traceback.format_exc())
        raise
